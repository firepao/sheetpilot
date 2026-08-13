#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


SHEET_NAME = "SheetPilotReviews"
COLUMNS = [
    "测试平台", "模型", "版本标签", "测试时间", "评估时间", "运行ID", "Runtime版本", "Git提交", "Skill哈希", "Bundle哈希",
    "测试提示词", "测试场景", "任务结果", "质量分", "诊断分",
    "接口合规", "执行效率", "恢复行为", "证据与报告", "输出可用性", "总用时", "总用时(ms)", "时间来源", "时间证据完整",
    "工具耗时合计(ms)", "Agent响应耗时合计(ms)", "Runtime耗时(ms)", "步骤数", "步骤平均耗时(ms)", "步骤P95耗时(ms)", "Agent操作数",
    "CLI调用数", "文件系统探索数", "Critical Violation数", "语义复核", "失败分类", "严重问题", "评分说明",
]
SCORE_COLUMNS = ["质量分", "诊断分", "接口合规", "执行效率", "恢复行为", "证据与报告", "输出可用性"]


def _number(value: Any, default: float = 0.0) -> float:
    try: return float(value)
    except (TypeError, ValueError): return default


def _optional_number(value: Any) -> float | str:
    return "" if value is None else _number(value)


def _human_duration(value: Any) -> str:
    if value is None: return ""
    seconds = int(round(_number(value) / 1000)); hours, remainder = divmod(seconds, 3600); minutes, seconds = divmod(remainder, 60)
    return f"{hours}:{minutes:02d}:{seconds:02d}" if hours else f"{minutes}:{seconds:02d}"


def normalize(score: dict[str, Any]) -> dict[str, Any]:
    breakdown, metrics = score.get("score_breakdown", {}), score.get("metrics", {})
    critical = sum(_number(value) for value in score.get("critical_violations", {}).values())
    diagnostic = sum(_number(breakdown.get(key)) for key in ("interface_compliance", "execution_efficiency", "recovery_behavior", "evidence_and_report", "output_usability"))
    passed = score.get("task_result") == "PASS"
    quality = diagnostic if passed else ""
    skill_hash = str(score.get("skill_sha256", ""))
    commit = str(score.get("git_commit", ""))
    version_label = score.get("skill_version") or (f"{commit[:8]}-{skill_hash[:8]}" if commit and skill_hash else commit[:8] or skill_hash[:12] or "未知")
    return {
        "测试平台": score.get("platform", "未知"), "模型": score.get("model", "未知"), "版本标签": version_label,
        "测试时间": score.get("tested_at", ""), "评估时间": score.get("evaluated_at", ""), "运行ID": score.get("run_id", ""),
        "Runtime版本": score.get("runtime_version", "未知"), "Git提交": commit, "Skill哈希": skill_hash, "Bundle哈希": score.get("bundle_sha256", ""),
        "测试提示词": score.get("prompt", ""), "测试场景": score.get("scenario_id", ""), "任务结果": score.get("task_result", "FAIL"),
        "质量分": quality, "诊断分": diagnostic, "接口合规": _number(breakdown.get("interface_compliance")),
        "执行效率": _number(breakdown.get("execution_efficiency")), "恢复行为": _number(breakdown.get("recovery_behavior")),
        "证据与报告": _number(breakdown.get("evidence_and_report")), "输出可用性": _number(breakdown.get("output_usability")),
        "总用时": _human_duration(metrics.get("wall_clock_duration_ms", metrics.get("elapsed_ms"))),
        "总用时(ms)": _optional_number(metrics.get("wall_clock_duration_ms", metrics.get("elapsed_ms"))),
        "时间来源": metrics.get("timing_source", "unavailable"), "时间证据完整": "是" if metrics.get("timing_complete") else "否",
        "工具耗时合计(ms)": _optional_number(metrics.get("tool_duration_sum_ms")),
        "Agent响应耗时合计(ms)": _optional_number(metrics.get("agent_response_duration_sum_ms")),
        "Runtime耗时(ms)": _optional_number(metrics.get("runtime_duration_ms")), "步骤数": _optional_number(metrics.get("step_count")),
        "步骤平均耗时(ms)": _optional_number(metrics.get("step_avg_ms")), "步骤P95耗时(ms)": _optional_number(metrics.get("step_p95_ms")),
        "Agent操作数": _number(metrics.get("agent_operations")),
        "CLI调用数": _number(metrics.get("sheetpilot_cli_calls")), "文件系统探索数": _number(metrics.get("filesystem_exploration")),
        "Critical Violation数": critical, "语义复核": score.get("semantic_review", {}).get("status", "not_assessed"),
        "失败分类": score.get("failure_class", ""), "严重问题": "否" if passed and critical == 0 else "是",
        "评分说明": score.get("review_notes", ""),
    }


def ensure_workbook(path: Path, append_only: bool):
    if path.exists(): workbook = load_workbook(path); created = False
    elif append_only: raise FileNotFoundError(f"append-only workbook does not exist: {path}")
    else: path.parent.mkdir(parents=True, exist_ok=True); workbook = Workbook(); created = True
    if SHEET_NAME in workbook.sheetnames: sheet = workbook[SHEET_NAME]
    elif created: sheet = workbook.active; sheet.title = SHEET_NAME
    else: sheet = workbook.create_sheet(SHEET_NAME)
    if sheet.max_row == 1 and sheet["A1"].value is None:
        for column, name in enumerate(COLUMNS, 1): sheet.cell(1, column, name)
    existing = [sheet.cell(1, index).value for index in range(1, sheet.max_column + 1)]
    for name in COLUMNS:
        if name not in existing: sheet.cell(1, len(existing) + 1, name); existing.append(name)
    return workbook, sheet, created, {name: existing.index(name) + 1 for name in COLUMNS}


def prompt_key(value: Any) -> str:
    return " ".join(str(value or "").casefold().split())


def groups(sheet, columns: dict[str, int]) -> list[tuple[int, int]]:
    if sheet.max_row < 2: return []
    result, start, current = [], 2, prompt_key(sheet.cell(2, columns["测试提示词"]).value)
    for row in range(3, sheet.max_row + 2):
        key = prompt_key(sheet.cell(row, columns["测试提示词"]).value) if row <= sheet.max_row else None
        if key != current: result.append((start, row - 1)); start, current = row, key
    return result


def format_sheet(sheet, columns: dict[str, int]) -> None:
    navy, white, yellow = PatternFill("solid", fgColor="1F4E78"), Font(color="FFFFFF", bold=True), PatternFill("solid", fgColor="FFF2CC")
    for cell in sheet[1]: cell.fill = navy; cell.font = white; cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    for index, name in enumerate(COLUMNS, 1): sheet.column_dimensions[sheet.cell(1, index).column_letter].width = 48 if name in {"测试提示词", "评分说明"} else (24 if name in {"测试时间", "评估时间", "Skill哈希", "Bundle哈希"} else 16)
    for row in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 42
        for cell in sheet[row]: cell.alignment = Alignment(horizontal="left" if cell.column in {columns["测试提示词"], columns["评分说明"]} else "center", vertical="center", wrap_text=True); cell.border = Border(); cell.fill = PatternFill(fill_type=None)
    thin = Side(style="thin", color="7A7A7A")
    for start, end in groups(sheet, columns):
        for row in range(start, end + 1):
            for column in range(1, len(COLUMNS) + 1): sheet.cell(row, column).border = Border(left=thin if column == 1 else None, right=thin if column == len(COLUMNS) else None, top=thin if row == start else None, bottom=thin if row == end else None)
        if end <= start: continue
        for name in SCORE_COLUMNS:
            column = columns[name]; values = [(row, _number(sheet.cell(row, column).value, float("nan"))) for row in range(start, end + 1)]
            values = [(row, value) for row, value in values if value == value]
            if not values: continue
            best = max(value for _, value in values)
            for row, value in values:
                if value == best: sheet.cell(row, column).fill = yellow; sheet.cell(row, column).font = Font(bold=True)


def main() -> int:
    parser = argparse.ArgumentParser(description="Append SheetPilot Agent score to an Excel history")
    parser.add_argument("excel_path", type=Path); parser.add_argument("--score-json", type=Path); parser.add_argument("--append-only", action="store_true")
    parser.add_argument("--check-prompt"); parser.add_argument("--connect", type=int)
    args = parser.parse_args()
    try:
        workbook, sheet, created, columns = ensure_workbook(args.excel_path, args.append_only)
        if args.check_prompt is not None:
            matches = [(row, str(sheet.cell(row, columns["测试提示词"]).value or "")) for row in range(2, sheet.max_row + 1) if args.check_prompt.casefold() in str(sheet.cell(row, columns["测试提示词"]).value or "").casefold()]
            print(json.dumps(matches, ensure_ascii=False)); workbook.close(); return 0
        if not args.score_json: raise ValueError("--score-json is required")
        record = normalize(json.loads(args.score_json.read_text(encoding="utf-8")))
        if args.connect is not None:
            if args.connect < 2 or args.connect > sheet.max_row: raise ValueError("--connect row is outside the data range")
            row = args.connect + 1; sheet.insert_rows(row)
        else: row = sheet.max_row + 1
        for name, value in record.items(): sheet.cell(row, columns[name], value)
        format_sheet(sheet, columns); group = next((item for item in groups(sheet, columns) if item[0] <= row <= item[1]), (row, row))
        workbook.save(args.excel_path); workbook.close()
        print(json.dumps({"status": "ok", "excel_path": str(args.excel_path.resolve()), "sheet": SHEET_NAME, "row": row, "created": created, "archive_card": {"row": row, "connect": args.connect, "group_range": f"{group[0]}-{group[1]}", "prompt": record["测试提示词"], "scenario": record["测试场景"]}}, ensure_ascii=False, indent=2)); return 0
    except Exception as exc:
        print(json.dumps({"status": "error", "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False), file=sys.stderr); return 1


if __name__ == "__main__": raise SystemExit(main())
