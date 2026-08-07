from pathlib import Path

from openpyxl import load_workbook

from ..models import ExecutionPlan, ValidationCheck


def check_layout(plan: ExecutionPlan, output_path: Path) -> ValidationCheck:
    wb = load_workbook(output_path, read_only=False, data_only=False)
    issues = []
    for step in plan.steps:
        if step.handler == "write_table":
            ws = wb[step.parameters["sheet"]]
            if not ws[step.parameters["anchor"]].value: issues.append("结果表头为空")
            for dimension in ws.column_dimensions.values():
                if dimension.width and dimension.width > 60: issues.append("列宽异常")
    wb.close()
    return ValidationCheck("layout", "important", not issues, "基础布局检查完成", {"issues": issues})
