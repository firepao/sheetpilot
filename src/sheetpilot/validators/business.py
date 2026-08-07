from __future__ import annotations

from pathlib import Path
from datetime import date, datetime

from openpyxl import load_workbook

from ..models import ExecutionPlan, ValidationCheck
from ..workbook.tables import TableData, filter_rows


def check_business(plan: ExecutionPlan, output_path: Path) -> ValidationCheck:
    write_steps = {s.id: s for s in plan.steps if s.handler == "write_table"}
    aggregate_steps = {s.id: s for s in plan.steps if s.handler == "aggregate"}
    if not plan.assertions:
        return ValidationCheck("business_reconciliation", "critical", False, "缺少核心业务断言")
    wb = load_workbook(output_path, read_only=True, data_only=True)
    failures, evidence = [], []
    for assertion in plan.assertions:
        result_id = assertion["result_step"]
        write = write_steps.get(result_id)
        if not write: failures.append("结果步骤不存在"); continue
        aggregate_id = write.parameters["input"]
        aggregate = aggregate_steps.get(aggregate_id) or next(iter(aggregate_steps.values()), None)
        ws = wb[write.parameters["sheet"]]
        anchor = ws[write.parameters["anchor"]]
        headers = [ws.cell(anchor.row, col).value for col in range(anchor.column, ws.max_column + 1)]
        try: metric_col = anchor.column + headers.index(assertion["output_metric"])
        except ValueError: failures.append("输出指标列不存在"); continue
        output_total = sum(ws.cell(row, metric_col).value or 0 for row in range(anchor.row + 1, ws.max_row + 1))
        source = next(s for s in plan.steps if s.id == assertion["source_step"])
        input_wb = load_workbook(plan.input_file, read_only=True, data_only=True); input_ws = input_wb[source.parameters["sheet"]]
        source_col = source.parameters["columns"][assertion["metric"]]
        from openpyxl.utils import column_index_from_string
        metric_index = column_index_from_string(source_col)
        date_field = aggregate.parameters.get("date_field") if aggregate else None
        date_index = column_index_from_string(source.parameters["columns"][date_field]) if date_field else None
        period = aggregate.parameters.get("current_period") or [None, None]
        source_rows = []
        column_indexes = {field: column_index_from_string(column) - 1 for field, column in source.parameters["columns"].items()}
        first_data_row = max(source.parameters.get("header_rows", [1])) + 1
        for values in input_ws.iter_rows(min_row=first_data_row, values_only=True):
            source_rows.append({field: values[index] if index < len(values) else None for field, index in column_indexes.items()})
        input_wb.close()
        source_table = TableData(list(source.parameters["columns"]), source_rows)
        filter_step = next((s for s in plan.steps if s.handler == "filter_rows"), None)
        if filter_step:
            try: source_table = filter_rows(source_table, filter_step.parameters["where"], filter_step.parameters.get("invalid_value_policy", "exclude"))
            except (TypeError, ValueError): failures.append("源数据过滤条件无法独立复算")
        source_total = sum(row.get(assertion["metric"]) or 0 for row in source_table.rows if isinstance(row.get(assertion["metric"]), (int, float)))
        tolerance = assertion.get("tolerance", 1e-6)
        if abs(source_total - output_total) > tolerance: failures.append("汇总值与源数据不一致")
        evidence.append({"source_total": source_total, "output_total": output_total})
    wb.close()
    return ValidationCheck("business_reconciliation", "critical", not failures, "核心指标独立复算完成", {"failures": failures, "reconciliation": evidence})
