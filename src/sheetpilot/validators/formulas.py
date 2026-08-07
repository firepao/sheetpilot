from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ..models import ExecutionPlan, ValidationCheck


ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")


def check_formulas(plan: ExecutionPlan, output_path: Path) -> ValidationCheck:
    if "formula_scan" not in plan.required_validations:
        return ValidationCheck("formula_scan", "important", True, "计划未要求公式扫描")
    wb = load_workbook(output_path, read_only=True, data_only=False)
    errors = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.startswith("=") and any(token in value for token in ERROR_TOKENS):
                        errors.append(f"{ws.title}!{cell.coordinate}")
    finally:
        wb.close()
    return ValidationCheck("formula_scan", "critical", not errors, "新增和现有公式文本扫描完成", {"error_cells": errors[:100]})
