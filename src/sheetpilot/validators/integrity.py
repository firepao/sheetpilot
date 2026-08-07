from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from ..models import ValidationCheck
from ..workspace import sha256_file


def check_integrity(input_path: Path, output_path: Path, expected_hash: str, expected_sheets: list[str]) -> list[ValidationCheck]:
    checks = []
    checks.append(ValidationCheck("input_unchanged", "critical", sha256_file(input_path) == expected_hash, "输入文件哈希未变化"))
    try:
        with ZipFile(output_path) as archive:
            required = {"[Content_Types].xml", "xl/workbook.xml"}
            zip_ok = required.issubset(archive.namelist())
        wb = load_workbook(output_path, read_only=True, data_only=False)
        missing = [name for name in expected_sheets if name not in wb.sheetnames]; wb.close()
        checks.append(ValidationCheck("file_integrity", "critical", zip_ok and not missing, "输出文件可打开且必需工作表存在", {"missing_sheets": missing}))
    except (BadZipFile, Exception) as exc:
        checks.append(ValidationCheck("file_integrity", "critical", False, "输出文件损坏", {"reason": str(exc)}))
    return checks
