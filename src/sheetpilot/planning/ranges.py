from __future__ import annotations

import re
from dataclasses import dataclass

from openpyxl.utils.cell import column_index_from_string, get_column_letter

from ..errors import ErrorCode, SheetPilotError

CELL = re.compile(r"^\$?([A-Z]{1,3})\$?([1-9][0-9]*)$")


@dataclass(frozen=True)
class RectRange:
    sheet: str
    min_row: int
    min_col: int
    max_row: int
    max_col: int

    @property
    def area(self) -> int:
        return (self.max_row - self.min_row + 1) * (self.max_col - self.min_col + 1)

    def intersects(self, other: "RectRange") -> bool:
        return self.sheet == other.sheet and not (self.max_row < other.min_row or other.max_row < self.min_row or self.max_col < other.min_col or other.max_col < self.min_col)

    def contains(self, other: "RectRange") -> bool:
        return self.sheet == other.sheet and self.min_row <= other.min_row <= other.max_row <= self.max_row and self.min_col <= other.min_col <= other.max_col <= self.max_col

    def to_a1(self) -> str:
        sheet = "'" + self.sheet.replace("'", "''") + "'"
        return f"{sheet}!{get_column_letter(self.min_col)}{self.min_row}:{get_column_letter(self.max_col)}{self.max_row}"


def parse_range(value: str, default_sheet: str | None = None) -> RectRange:
    if "!" in value:
        raw_sheet, address = value.rsplit("!", 1)
        sheet = raw_sheet.strip("'").replace("''", "'")
    else:
        sheet, address = default_sheet, value
    if not sheet:
        raise SheetPilotError(ErrorCode.PLAN_INVALID, "范围缺少工作表名", {"range": value})
    parts = address.split(":")
    if len(parts) > 2:
        raise SheetPilotError(ErrorCode.PLAN_INVALID, "非法 A1 范围", {"range": value})
    matches = [CELL.match(part.upper()) for part in parts]
    if any(match is None for match in matches):
        raise SheetPilotError(ErrorCode.PLAN_INVALID, "不允许整行、整列或非法范围", {"range": value})
    points = [(int(m.group(2)), column_index_from_string(m.group(1))) for m in matches if m]
    if len(points) == 1: points.append(points[0])
    (r1, c1), (r2, c2) = points
    return RectRange(sheet, min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2))
