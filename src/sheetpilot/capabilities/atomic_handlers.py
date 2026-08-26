from __future__ import annotations

from copy import copy
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.drawing.image import Image

from ..workbook.formulas import add_formula_column
from ..workbook.operations import append_rows, apply_style, clear_range, copy_range, copy_style, inspect_sheet, inspect_workbook, range_ref, read_range, validate_structure, write_range, write_table
from ..workbook.tables import aggregate, concat_tables, count_values, deduplicate, describe, fill_missing, filter_rows, join_tables, melt_table, pivot_table, read_table, rename_fields, select_columns, sort_rows, unique_values, value_counts


def _engine(ctx): return ctx.engine
def _table(results, params): return results[params["input"]]


def execute_atomic(op: str, ctx, p: dict[str, Any], results: dict[str, Any]):
    engine = _engine(ctx); wb = engine._workbook
    if op == "workbook.inspect": return inspect_workbook(engine)
    if op == "workbook.list_sheets": return engine.sheet_names()
    if op == "workbook.validate_structure": return validate_structure(engine, p.get("required_sheets"), p.get("forbidden_sheets"))
    if op == "sheet.inspect": return inspect_sheet(engine, p["sheet"])
    if op == "sheet.create": engine.create_sheet(p["sheet"]); return {"sheet": p["sheet"]}
    if op == "sheet.delete": wb.remove(wb[p["sheet"]]); return {"deleted": p["sheet"]}
    if op == "sheet.rename": wb[p["sheet"]].title = p["name"]; return {"sheet": p["name"]}
    if op == "sheet.copy": ws=wb.copy_worksheet(wb[p["sheet"]]); ws.title=p["name"]; return {"sheet": p["name"]}
    if op == "sheet.set_visibility": wb[p["sheet"]].sheet_state=p["visibility"]; return {"sheet":p["sheet"],"visibility":p["visibility"]}
    if op == "cell.read": return engine.read_cell(p["sheet"], p["row"], p["column"])
    if op == "cell.write": engine.write_cell(p["sheet"],p["row"],p["column"],p.get("value")); return {"sheet":p["sheet"],"row":p["row"],"column":p["column"]}
    if op == "cell.clear": engine.write_cell(p["sheet"],p["row"],p["column"],None); return {"sheet":p["sheet"],"row":p["row"],"column":p["column"]}
    if op == "range.read": return read_range(engine,p["sheet"],p["range"])
    if op == "range.write": return write_range(engine,p["sheet"],p["anchor"],p["values"])
    if op == "range.write_table": return write_table(engine,p["sheet"],p.get("anchor","A1"),_table(results,p),p.get("include_header",True))
    if op == "range.append_rows": return append_rows(engine,p["sheet"],p["values"])
    if op == "range.copy": return copy_range(engine,p["source_sheet"],p["source_range"],p["target_sheet"],p["anchor"],p.get("include_style",True))
    if op == "range.clear": return clear_range(engine,p["sheet"],p["range"])
    if op == "range.merge": wb[p["sheet"]].merge_cells(p["range"]); return range_ref(p["sheet"],p["range"])
    if op == "range.unmerge": wb[p["sheet"]].unmerge_cells(p["range"]); return range_ref(p["sheet"],p["range"])
    if op in {"rows.insert","rows.delete","columns.insert","columns.delete"}:
        ws=wb[p["sheet"]]; index=p["index"]; amount=p.get("amount",1)
        {"rows.insert":ws.insert_rows,"rows.delete":ws.delete_rows,"columns.insert":ws.insert_cols,"columns.delete":ws.delete_cols}[op](index,amount)
        return {"sheet":p["sheet"],"index":index,"amount":amount}
    if op == "formula.read": return wb[p["sheet"]][p["cell"]].value
    if op == "formula.write": wb[p["sheet"]][p["cell"]]=p["formula"]; return {"sheet":p["sheet"],"cell":p["cell"]}
    if op in {"formula.fill","formula.copy"}:
        ws=wb[p["sheet"]]; source=ws[p["source"]]; target=range_ref(p["sheet"],p["target"])
        for row in ws.iter_rows(min_row=target["min_row"],max_row=target["max_row"],min_col=target["min_col"],max_col=target["max_col"]):
            for cell in row: cell.value=source.value
        return target
    if op == "style.apply": return apply_style(engine,p.get("ref") or results[p["input"]],p["preset"])
    if op == "style.copy": return copy_style(engine,p["source_sheet"],p["source_range"],p["target_sheet"],p["target_range"])
    if op == "style.apply_table_default": return apply_style(engine,p.get("ref") or results[p["input"]],"business_table")
    if op == "number_format.apply":
        ref=range_ref(p["sheet"],p["range"]); ws=wb[p["sheet"]]
        for row in ws.iter_rows(min_row=ref["min_row"],max_row=ref["max_row"],min_col=ref["min_col"],max_col=ref["max_col"]):
            for cell in row: cell.number_format=p["format"]
        return ref
    if op == "column.set_width": wb[p["sheet"]].column_dimensions[p["column"]].width=p["width"]; return p
    if op == "row.set_height": wb[p["sheet"]].row_dimensions[p["row"]].height=p["height"]; return p
    if op == "freeze_panes.set": wb[p["sheet"]].freeze_panes=p.get("cell"); return p
    if op == "auto_filter.set": wb[p["sheet"]].auto_filter.ref=p["range"]; return p
    if op == "protection.set": wb[p["sheet"]].protection.sheet=bool(p.get("enabled",True)); return p
    if op == "table.read": return read_table(engine,p["sheet"],p.get("header_row",1),p["columns"])
    if op == "table.select": return select_columns(_table(results,p),p["fields"])
    if op == "table.filter": return filter_rows(_table(results,p),p["where"],p.get("invalid_value_policy","exclude"))
    if op == "table.group": return {"table":_table(results,p),"fields":p["fields"]}
    if op == "table.aggregate":
        source=_table(results,p); groups=source if not isinstance(source,dict) else source["table"]
        group_by=p.get("group_by", source.get("fields") if isinstance(source,dict) else [])
        return aggregate(groups,group_by,p["metrics"],p.get("date_field"),p.get("current_period"))
    if op == "table.sort": return sort_rows(_table(results,p),p["keys"])
    if op == "table.deduplicate": return deduplicate(_table(results,p),p.get("keys"),p.get("keep","first"))
    if op == "table.fill_missing": return fill_missing(_table(results,p),p.get("fields"),p.get("value",""))
    if op == "table.describe": return describe(_table(results,p),p.get("fields"))
    if op == "table.value_counts": return value_counts(_table(results,p),p["field"],p.get("as","数量"))
    if op == "table.count": return count_values(_table(results,p),p.get("field"),p.get("mode","rows"))
    if op == "table.unique": return unique_values(_table(results,p),p["field"],p.get("sort",True))
    if op == "table.rename_fields": return rename_fields(_table(results,p),p["mapping"])
    if op == "table.concat": return concat_tables([results[item] for item in p["inputs"]])
    if op == "table.join": return join_tables(results[p["left"]],results[p["right"]],p["on"],p.get("how","inner"))
    if op == "table.pivot": return pivot_table(_table(results,p),p["index"],p["columns"],p["values"],p.get("function","sum"))
    if op == "table.melt": return melt_table(_table(results,p),p["id_fields"],p["value_fields"],p.get("variable_name","变量"),p.get("value_name","值"))
    if op == "excel_table.create":
        ws=wb[p["sheet"]]; table=Table(displayName=p["name"],ref=p["range"]); table.tableStyleInfo=TableStyleInfo(name=p.get("style","TableStyleMedium2"),showRowStripes=True); ws.add_table(table); return {"sheet":p["sheet"],"name":p["name"],"range":p["range"]}
    if op == "chart.create":
        ws=wb[p["sheet"]]; ref=range_ref(p["sheet"],p["range"]); chart={"bar":BarChart,"column":BarChart,"line":LineChart,"pie":PieChart}[p["chart_type"]](); chart.title=p.get("title","")
        chart.add_data(Reference(ws,min_col=ref["min_col"]+1,max_col=ref["max_col"],min_row=ref["min_row"],max_row=ref["max_row"]),titles_from_data=True); chart.set_categories(Reference(ws,min_col=ref["min_col"],min_row=ref["min_row"]+1,max_row=ref["max_row"])); ws.add_chart(chart,p["anchor"]); return {"sheet":p["sheet"],"chart_index":len(ws._charts)-1}
    if op == "hyperlink.create": wb[p["sheet"]][p["cell"]].hyperlink=p["target"]; return p
    if op == "formula.inspect_dependencies":
        formula=wb[p["sheet"]][p["cell"]].value or ""; import re
        return {"sheet":p["sheet"],"cell":p["cell"],"references":sorted(set(re.findall(r"(?:'[^']+'|[A-Za-z0-9_]+)!?\$?[A-Z]{1,3}\$?\d+|\$?[A-Z]{1,3}\$?\d+",str(formula))))}
    if op == "formula.find_errors":
        ws=wb[p["sheet"]]; errors=[]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value,str) and cell.value.startswith("#"): errors.append(cell.coordinate)
        return {"sheet":p["sheet"],"errors":errors}
    if op == "chart.add_series":
        ws=wb[p["sheet"]]; chart=ws._charts[p["chart_index"]]; ref=range_ref(p["sheet"],p["range"]); chart.add_data(Reference(ws,min_col=ref["min_col"],max_col=ref["max_col"],min_row=ref["min_row"],max_row=ref["max_row"]),titles_from_data=p.get("titles_from_data",True)); return p
    if op == "chart.set_categories":
        ws=wb[p["sheet"]]; chart=ws._charts[p["chart_index"]]; ref=range_ref(p["sheet"],p["range"]); chart.set_categories(Reference(ws,min_col=ref["min_col"],max_col=ref["max_col"],min_row=ref["min_row"],max_row=ref["max_row"])); return p
    if op == "chart.inspect":
        ws=wb[p["sheet"]]; chart=ws._charts[p["chart_index"]]; return {"sheet":p["sheet"],"chart_index":p["chart_index"],"title":str(chart.title) if chart.title else None,"series_count":len(chart.ser)}
    if op == "validation.create":
        ws=wb[p["sheet"]]; validation=DataValidation(type=p["validation_type"],formula1=p.get("formula1"),formula2=p.get("formula2"),allow_blank=p.get("allow_blank",True)); ws.add_data_validation(validation); validation.add(p["range"]); return {"sheet":p["sheet"],"range":p["range"]}
    if op == "conditional_format.create":
        ws=wb[p["sheet"]]; rule=CellIsRule(operator=p["operator"],formula=p.get("formula",[]),fill=None); ws.conditional_formatting.add(p["range"],rule); return {"sheet":p["sheet"],"range":p["range"]}
    if op == "named_range.create":
        wb.defined_names.add(DefinedName(p["name"],attr_text=f"'{p['sheet']}'!{p['range']}")); return {"name":p["name"]}
    if op == "comment.create": wb[p["sheet"]][p["cell"]].comment=Comment(p["text"],p.get("author","SheetPilot")); return p
    if op == "image.insert":
        image=Image(p["path"]); wb[p["sheet"]].add_image(image,p.get("anchor","A1")); return {"sheet":p["sheet"],"anchor":p.get("anchor","A1")}
    raise ValueError(f"unsupported atomic operation: {op}")


def handler(op: str):
    return lambda ctx, params, results: execute_atomic(op, ctx, params, results)
