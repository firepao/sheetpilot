"""轻量 Excel MVP：带契约的 Atom/Molecule、执行证据和发布门禁。"""
from __future__ import annotations

import json
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from .engines import OpenPyxlEngine
from .errors import ErrorCode, SheetPilotError
from .runtime import TransformManifest, TransformRunner
from .workbook.tables import TableData, aggregate, derive_column, filter_rows, read_table, sort_rows
from .workspace import Workspace, sha256_file


Atom = Callable[[Any, dict[str, Any], dict[str, Any]], Any]


@dataclass(frozen=True)
class Operation:
    name: str
    kind: str
    required: tuple[str, ...]
    input_kind: str
    output_kind: str
    handler: Atom | None = None
    optional: tuple[str, ...] = ()
    steps: tuple[dict[str, Any], ...] = ()
    supported_functions: tuple[str, ...] = ()
    effects: tuple[str, ...] = ()
    validators: tuple[str, ...] = ()

    def public_dict(self) -> dict[str, Any]:
        value = {"name": self.name, "kind": self.kind, "required": list(self.required), "optional": list(self.optional), "input_kind": self.input_kind, "output_kind": self.output_kind, "effects": list(self.effects), "validators": list(self.validators)}
        if self.steps: value["steps"] = list(self.steps)
        if self.supported_functions: value["supported_functions"] = list(self.supported_functions)
        return value


def _read(engine, p, _): return read_table(engine, p["sheet"], p.get("header_row", 1), p["columns"])
def _filter(_, p, r): return filter_rows(r[p["input"]], p["where"], p.get("invalid_value_policy", "exclude"))
def _derive(_, p, r): return derive_column(r[p["input"]], p["name"], p["template"], p["arguments"])
def _aggregate(_, p, r): return aggregate(r[p["input"]], p["group_by"], p["metrics"], p.get("date_field"), p.get("period"))
def _sort(_, p, r): return sort_rows(r[p["input"]], p.get("keys", p.get("sort", [])))
def _create(engine, p, _):
    if p["sheet"] in engine.sheet_names(): raise SheetPilotError(ErrorCode.PLAN_INVALID, "目标工作表已存在", {"sheet": p["sheet"]})
    engine.create_sheet(p["sheet"]); return {"sheet": p["sheet"]}
def _write(engine, p, r):
    table = r[p["input"]]
    if not isinstance(table, TableData): raise SheetPilotError(ErrorCode.PLAN_INVALID, "write_table 输入不是 TableData")
    row, column = coordinate_to_tuple(p.get("anchor", "A1")); ws = engine.worksheet(p["sheet"])
    if p.get("include_header", True):
        for offset, header in enumerate(table.headers): ws.cell(row, column + offset).value = header
        row += 1
    for row_offset, item in enumerate(table.rows):
        for column_offset, header in enumerate(table.headers): ws.cell(row + row_offset, column + column_offset).value = item.get(header)
    return {"sheet": p["sheet"], "anchor": p.get("anchor", "A1"), "rows": len(table.rows), "columns": list(table.headers)}
def _save(_, p, __): return {"requested": True}


ATOMS = {
    "read_table": Operation("read_table", "atom", ("sheet", "columns"), "WorkbookRef", "TableData", _read, ("header_row",), effects=("read_workbook",)),
    "filter_rows": Operation("filter_rows", "atom", ("input", "where"), "TableData", "TableData", _filter, ("invalid_value_policy",)),
    "derive_column": Operation("derive_column", "atom", ("input", "name", "template", "arguments"), "TableData", "TableData", _derive),
    "aggregate": Operation("aggregate", "atom", ("input", "group_by", "metrics"), "TableData", "TableData", _aggregate, ("date_field", "period"), supported_functions=("sum", "count", "average"), validators=("aggregate_reconciliation",)),
    "sort_rows": Operation("sort_rows", "atom", ("input", "keys"), "TableData", "TableData", _sort, validators=("sort_order",)),
    "create_sheet": Operation("create_sheet", "atom", ("sheet",), "WorkbookRef", "SheetRef", _create, effects=("create_sheet",)),
    "write_table": Operation("write_table", "atom", ("input", "sheet"), "TableData", "RangeRef", _write, ("anchor", "include_header"), effects=("write_workbook",)),
    "save_workbook": Operation("save_workbook", "atom", (), "WorkbookRef", "WorkbookRef", _save, effects=("save_temporary",)),
}

MOLECULES = {
    "summarize_by_dimension": Operation("summarize_by_dimension", "molecule", ("input", "group_by", "metrics"), "TableData", "TableData", optional=("where", "sort"), steps=({"op":"filter_rows","when":"where"},{"op":"aggregate"},{"op":"sort_rows","when":"sort"}), validators=("row_filter_count","aggregate_reconciliation","sort_order")),
    "add_metric": Operation("add_metric", "molecule", ("input", "name", "template", "arguments"), "TableData", "TableData", steps=({"op":"derive_column"},)),
}


def manifest() -> dict[str, Any]:
    return {"schema_version": "1.0", "atoms": [item.public_dict() for item in ATOMS.values()], "molecules": [item.public_dict() for item in MOLECULES.values()], "dynamic_transform": {"input_kind":"TableData", "output_kind":"TableData", "forbidden":["file","network","subprocess","environment","workbook_libraries"]}}


def _fail(message: str, details: dict | None = None) -> None: raise SheetPilotError(ErrorCode.PLAN_INVALID, message, details)


def validate_plan(plan: dict[str, Any]) -> None:
    required = {"schema_version", "input_file", "output_file", "steps", "requirements"}
    if not isinstance(plan, dict) or not required <= set(plan): _fail("轻量计划缺少必需字段", {"missing": sorted(required - set(plan or {}))})
    if plan["schema_version"] != "1.0" or not isinstance(plan["steps"], list): _fail("轻量计划版本或 steps 不合法")
    ids: set[str] = set()
    for step in plan["steps"]:
        if not isinstance(step, dict) or not isinstance(step.get("id"), str) or step["id"] in ids: _fail("步骤 id 缺失或重复")
        ids.add(step["id"]); op = step.get("op")
        definition = ATOMS.get(op) or MOLECULES.get(op)
        if op == "dynamic_transform":
            if "input" not in step: _fail("dynamic_transform 缺少 input")
            continue
        if definition is None: raise SheetPilotError(ErrorCode.CAPABILITY_UNSUPPORTED, "未注册操作", {"op": op})
        missing = [name for name in definition.required if name not in step]
        allowed = {"id", "op", *definition.required, *definition.optional}
        unknown = sorted(set(step) - allowed)
        if missing or unknown: _fail("能力参数不符合契约", {"op": op, "missing": missing, "unknown": unknown})
        if "input" in step and step["input"] not in ids: _fail("步骤引用不存在或尚未执行", {"id": step["id"], "input": step["input"]})
        if op in {"aggregate", "summarize_by_dimension"}:
            aliases = set()
            for metric in step["metrics"]:
                function = metric.get("function")
                if function not in ATOMS["aggregate"].supported_functions: _fail("不支持的聚合函数", {"function": function})
                if not metric.get("as") or metric["as"] in aliases: _fail("聚合指标别名缺失或重复")
                aliases.add(metric["as"])
                if function in {"sum", "average"} and not metric.get("field"): _fail("聚合函数缺少 field", {"function": function})
                if function == "count" and metric.get("mode") not in {"rows", "non_empty"}: _fail("count 必须声明 rows 或 non_empty 模式")
                if function == "count" and metric.get("mode") == "non_empty" and not metric.get("field"): _fail("non_empty count 缺少 field")


def execute(engine, steps: list[dict[str, Any]], dynamic=None, record: dict[str, Any] | None = None) -> dict[str, Any]:
    results: dict[str, Any] = {}; record = record if record is not None else {}
    record.setdefault("requested_operations", []); record.setdefault("expanded_operations", []); record.setdefault("executed_atoms", []); record.setdefault("dynamic_transforms", []); record.setdefault("workbook_writes", [])
    for step in steps:
        op = step["op"]; record["requested_operations"].append({"id":step["id"],"op":op})
        fragments = []
        if op in ATOMS: fragments = [(step["id"], op, step)]
        elif op in MOLECULES:
            current = step["input"]
            for index, fragment in enumerate(MOLECULES[op].steps):
                if fragment.get("when") and fragment["when"] not in step: continue
                atom = fragment["op"]; params = {**step, "op": atom, "input": current}
                if atom == "sort_rows": params["keys"] = step["sort"]
                fragment_id = step["id"] if index == len(MOLECULES[op].steps)-1 else f"{step['id']}__{atom}"
                fragments.append((fragment_id, atom, params)); current = fragment_id
            record["expanded_operations"].append({"id":step["id"],"op":op,"atoms":[item[1] for item in fragments]})
        elif op == "dynamic_transform":
            if dynamic is None: raise SheetPilotError(ErrorCode.PLAN_INVALID, "未提供受控动态脚本")
            results[step["id"]] = dynamic(results[step["input"]], step); record["dynamic_transforms"].append({"id":step["id"],"input_rows":len(results[step["input"]].rows),"output_rows":len(results[step["id"]].rows)}); continue
        for fragment_id, atom, params in fragments:
            input_value = results.get(params.get("input"))
            try: value = ATOMS[atom].handler(engine, params, results)
            except SheetPilotError: raise
            except (KeyError, TypeError, ValueError) as exc: raise SheetPilotError(ErrorCode.PLAN_INVALID, "能力执行参数无效", {"op":atom,"reason":str(exc)}) from exc
            atom_record = {"id":fragment_id,"op":atom}
            if isinstance(input_value, TableData): atom_record["input_rows"] = len(input_value.rows)
            if isinstance(value, TableData): atom_record["output_rows"] = len(value.rows)
            results[fragment_id] = value; record["executed_atoms"].append(atom_record)
            if atom in {"create_sheet","write_table","save_workbook"}: record["workbook_writes"].append({"id":fragment_id,"op":atom,"result":value})
        if fragments: results[step["id"]] = results[fragments[-1][0]]
    return results


def run_plan(plan: dict[str, Any], run_dir: Path, dynamic_script: Path | None = None) -> dict[str, Any]:
    validate_plan(plan); run_dir = run_dir.resolve()
    workspace = Workspace(run_dir, allowed_root=Path(plan["output_file"]).resolve().parent)
    source = workspace.validate_input(Path(plan["input_file"])); output = workspace.validate_output(source, Path(plan["output_file"]))
    if any(run_dir.iterdir()): raise SheetPilotError(ErrorCode.INPUT_INVALID, "run-dir 必须为空且独立", {"run_dir":str(run_dir)})
    working = workspace.create_working_copy(source); input_hash = sha256_file(source); record: dict[str, Any] = {"schema_version":"1.0","input_file":str(source),"input_sha256":input_hash,"output_file":str(output),"validation_results":[]}
    runner = None
    if dynamic_script:
        source_code = dynamic_script.read_text(encoding="utf-8"); transform_runner = TransformRunner()
        runner = lambda table, step: transform_runner.run(source_code, table, step.get("params", {}), TransformManifest(dynamic_script.stem), run_dir)
    engine = OpenPyxlEngine.open(working)
    try:
        execute(engine, plan["steps"], runner, record); temporary = run_dir / "temporary_output.xlsx"; engine.save(temporary)
    finally: engine.close()
    plan_copy = dict(plan); plan_copy["input_file"] = str(source); plan_copy["output_file"] = str(output)
    (run_dir / "plan.json").write_text(json.dumps(plan_copy, ensure_ascii=False, indent=2), encoding="utf-8")
    record["temporary_output"] = str(temporary); record["temporary_sha256"] = sha256_file(temporary)
    (run_dir / "execution.json").write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"status":"READY_FOR_VALIDATION","run_dir":str(run_dir),"temporary_output":str(temporary),"executed_atoms":record["executed_atoms"],"expanded_molecules":record["expanded_operations"],"dynamic_transforms":record["dynamic_transforms"]}


def validate_run(run_dir: Path) -> dict[str, Any]:
    run_dir = run_dir.resolve(); plan = json.loads((run_dir / "plan.json").read_text(encoding="utf-8")); record = json.loads((run_dir / "execution.json").read_text(encoding="utf-8"))
    source, temporary, output = Path(plan["input_file"]), run_dir / "temporary_output.xlsx", Path(plan["output_file"])
    checks = []
    def check(name, passed, details=None): checks.append({"check":name,"status":"PASS" if passed else "FAIL","details":details or {}})
    check("source_unchanged", source.is_file() and sha256_file(source) == record["input_sha256"])
    check("temporary_unchanged", temporary.is_file() and sha256_file(temporary) == record["temporary_sha256"])
    previous_evidence = run_dir / "evidence.json"
    if previous_evidence.is_file() and output.is_file():
        previous = json.loads(previous_evidence.read_text(encoding="utf-8"))
        if previous.get("status") == "PASS": check("published_unchanged", sha256_file(output) == previous.get("published_sha256"))
    wb = load_workbook(temporary, data_only=True, read_only=True)
    try:
        requirements = plan.get("requirements", {})
        for sheet in requirements.get("required_sheets", []): check("required_sheet", sheet in wb.sheetnames, {"sheet":sheet})
        for sheet, columns in requirements.get("required_columns", {}).items():
            actual = [cell.value for cell in next(wb[sheet].iter_rows(min_row=1,max_row=1))] if sheet in wb.sheetnames else []
            check("required_columns", all(column in actual for column in columns), {"sheet":sheet,"required":columns,"actual":actual})
        for requirement in requirements.get("checks", []):
            item = {"type":requirement} if isinstance(requirement, str) else requirement; kind = item.get("type")
            if kind in {"source_unchanged"}: continue
            if kind == "sort_order":
                sheet, field, direction = item["sheet"], item["field"], item.get("direction","asc"); ws=wb[sheet]; headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]; index=headers.index(field); values=[row[index].value for row in ws.iter_rows(min_row=2) if row[index].value is not None]; check(kind, values == sorted(values, reverse=direction=="desc"), {"sheet":sheet,"field":field})
            elif kind == "aggregate_reconciliation":
                source_ws, target_ws = wb[item["source_sheet"]], wb[item["target_sheet"]]
                source_headers=[c.value for c in next(source_ws.iter_rows(min_row=1,max_row=1))]; target_headers=[c.value for c in next(target_ws.iter_rows(min_row=1,max_row=1))]
                source_index, target_index = source_headers.index(item["source_field"]), target_headers.index(item["target_field"])
                where=item.get("where",{}); source_total=0
                for row in source_ws.iter_rows(min_row=2,values_only=True):
                    mapped=dict(zip(source_headers,row))
                    if all(mapped.get(c["field"]) == c.get("value") for c in where.get("all",[])) and isinstance(row[source_index],(int,float)): source_total += row[source_index]
                target_total=sum(row[target_index] or 0 for row in target_ws.iter_rows(min_row=2,values_only=True)); check(kind, abs(source_total-target_total)<=item.get("tolerance",0.000001), {"source_total":source_total,"target_total":target_total})
            elif kind == "row_filter_count":
                filters=[atom for atom in record["executed_atoms"] if atom["op"] == "filter_rows"]
                check(kind, bool(filters) and all("input_rows" in atom and "output_rows" in atom for atom in filters), {"filters":filters})
            else: check(kind or "unknown", False, {"reason":"unsupported requirement check"})
    finally: wb.close()
    record["validation_results"] = checks
    (run_dir / "execution.json").write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
    if any(item["status"] == "FAIL" for item in checks):
        evidence={"status":"FAIL","checks":checks}; (run_dir/"evidence.json").write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding="utf-8"); raise SheetPilotError(ErrorCode.VALIDATION_FAILED,"轻量任务验证失败",evidence)
    output.parent.mkdir(parents=True,exist_ok=True); shutil.copy2(temporary,output); output_hash=sha256_file(output)
    evidence={"status":"PASS","checks":checks,"validated_sha256":record["temporary_sha256"],"published_sha256":output_hash}; (run_dir/"evidence.json").write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding="utf-8")
    result={"status":"PASS","run_dir":str(run_dir),"output_file":str(output),"validation":evidence,"executed_atoms":record["executed_atoms"],"expanded_molecules":record["expanded_operations"],"dynamic_transforms":record["dynamic_transforms"]}; (run_dir/"result.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8"); return result
