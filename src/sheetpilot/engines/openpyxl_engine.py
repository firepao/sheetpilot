from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..errors import ErrorCode, SheetPilotError


class OpenPyxlEngine:
    engine_name = "openpyxl"
    def __init__(self, workbook):
        self._workbook = workbook

    @classmethod
    def open(cls, path: Path) -> "OpenPyxlEngine":
        try:
            return cls(load_workbook(path, data_only=False, keep_links=True))
        except Exception as exc:
            raise SheetPilotError(ErrorCode.EXECUTION_FAILED, "无法打开工作副本", {"reason": str(exc)}) from exc

    def sheet_names(self) -> list[str]: return self._workbook.sheetnames
    def worksheet(self, name: str): return self._workbook[name]
    def read_cell(self, sheet: str, row: int, column: int) -> Any: return self._workbook[sheet].cell(row, column).value
    def write_cell(self, sheet: str, row: int, column: int, value: Any) -> None: self._workbook[sheet].cell(row, column).value = value
    def create_sheet(self, name: str) -> None: self._workbook.create_sheet(name)

    def save(self, path: Path) -> None:
        try: self._workbook.save(path)
        except Exception as exc: raise SheetPilotError(ErrorCode.EXECUTION_FAILED, "保存临时输出失败", {"reason": str(exc)}) from exc

    def close(self) -> None: self._workbook.close()
