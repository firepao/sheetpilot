from __future__ import annotations

import copy
import json
import os
import shutil
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from ..engines import OpenPyxlEngine
from ..workbook.tables import TableData, aggregate, filter_rows, read_table, select_columns, sort_rows
from ..workspace import sha256_file
from .compiler import collect_fields, compile_plan
from .contract import acceptance_snapshot, invalid_response, stable_hash, task_type_manifest, validate_request
from .inventory import build_inventory, cap_entries, read_headers


def default_state_root() -> Path:
    configured = os.environ.get("SHEETPILOT_STATE_ROOT")
    if configured:
        return Path(configured).expanduser()
    if os.name == "nt":
        base = Path(os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local"))
    else:
        base = Path(os.environ.get("XDG_STATE_HOME", Path.home() / ".local" / "state"))
    return base / "SheetPilot"


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temporary, path)


class TaskRuntime:
    def __init__(self, state_root: Path | None = None):
        self.state_root = (state_root or default_state_root()).resolve()
        self.tasks_root = self.state_root / "tasks"

    def _query_error(self, code: str, message: str, action: str = "AMEND_REQUEST", allowed: list[dict[str, Any]] | None = None) -> dict[str, Any]:
        retryable = action == "AMEND_REQUEST"
        return {"schema_version": "1.0", "status": "REQUEST_INVALID", "task_id": None, "request_revision": None, "attempt_id": None,
                "error": {"code": code, "phase": "workbook_inspection", "message": message, "retryable": retryable,
                          "diagnostics": [{"code": code, "path": None, "message": message}],
                          "recovery": {"action": action, "retryable": retryable, "base_revision": None, "allowed_amendments": allowed or [], "suggested_patch": []}}}

    def task_types(self, input_file: str | None = None, sheet: str | None = None, header_row: int | None = None) -> dict[str, Any]:
        manifest = task_type_manifest()
        if not input_file:
            return manifest
        source = Path(input_file).expanduser().resolve()
        if not source.is_file() or source.suffix.lower() not in {".xlsx", ".xlsm"}:
            return self._query_error("INPUT_NOT_FOUND", "输入文件不存在或不是受支持的工作簿。", "HUMAN_ACTION_REQUIRED")
        from openpyxl import load_workbook
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            visible = [name for name in workbook.sheetnames if workbook[name].sheet_state == "visible"]
        finally:
            workbook.close()
        if header_row is not None and header_row < 1:
            return self._query_error("INVALID_VALUE", "header_row 必须大于等于 1。", allowed=[{"op": "replace", "path": "/query/header_row", "constraints": {"min": 1}}])
        if sheet is not None and sheet not in visible:
            return self._query_error("INVALID_VALUE", "Sheet 不存在或不可见。", allowed=[{"op": "replace", "path": "/query/sheet", "constraints": {"enum": visible}}])
        input_sha256 = sha256_file(source)
        profile = build_inventory(source, input_sha256, sheet=sheet, header_row=header_row)
        if profile["error_code"] == "INVENTORY_TOO_LARGE":
            return self._query_error("INVENTORY_TOO_LARGE", "字段清单超出上限，请使用 --sheet 收窄查询。", allowed=[{"op": "replace", "path": "/query/sheet", "constraints": {"enum": visible}}, {"op": "replace", "path": "/query/header_row", "constraints": {"min": 1}}])
        result = copy.deepcopy(manifest)
        result["input_profile"] = {"input_file": str(source), "input_sha256": input_sha256, **profile}
        return result

    def run(self, request: Any) -> dict[str, Any]:
        if isinstance(request, dict) and "task_id" in request:
            return self._amend(request)
        errors = validate_request(request)
        if errors:
            return invalid_response(errors)
        return self._create_and_run(copy.deepcopy(request))

    def status(self, task_id: str) -> dict[str, Any]:
        task_dir = self._task_dir(task_id)
        state_path = task_dir / "current-state.json"
        if not state_path.is_file():
            return self._error("REQUEST_INVALID", "INVALID_REFERENCE", "contract_validation", "Task 不存在。", None, None, None, False, "NONE")
        state = json.loads(state_path.read_text(encoding="utf-8"))
        task = json.loads((task_dir / "task.json").read_text(encoding="utf-8"))
        response = copy.deepcopy(state)
        response.update({"schema_version": "1.0", "task_id": task_id, "acceptance_hash": task["acceptance_hash"]})
        output = Path(task["output_file"])
        expected = state.get("validated_published_sha256")
        if state["state"] == "RUNTIME_PASS":
            current = sha256_file(output) if output.is_file() else None
            response["artifact_integrity"] = "MATCHED" if current == expected else ("MISSING" if current is None else "MODIFIED")
            response["delivery_valid"] = current == expected
            response["current_published_sha256"] = current
        else:
            response.setdefault("artifact_integrity", None); response.setdefault("delivery_valid", False)
        return response

    def _create_and_run(self, request: dict[str, Any]) -> dict[str, Any]:
        self.tasks_root.mkdir(parents=True, exist_ok=True)
        source, output = Path(request["input_file"]).expanduser().resolve(), Path(request["output_file"]).expanduser().resolve()
        if not source.is_file() or source.suffix.lower() not in {".xlsx", ".xlsm"}:
            return self._error("REQUEST_INVALID", "INPUT_NOT_FOUND", "workbook_inspection", "输入文件不存在或不是受支持的工作簿。", None, None, None, False, "HUMAN_ACTION_REQUIRED")
        if source == output:
            return self._error("REQUEST_INVALID", "INVALID_VALUE", "contract_validation", "输出文件不能覆盖输入文件。", None, None, None, False, "CREATE_NEW_TASK")
        request["input_file"], request["output_file"] = str(source), str(output)
        input_hash = sha256_file(source)
        submission = copy.deepcopy(request); submission.pop("user_request", None)
        submission_hash = stable_hash(submission)
        existing = self._find_submission(submission_hash, input_hash, str(output))
        if existing:
            return self.status(existing)
        if output.exists():
            return self._error("REQUEST_INVALID", "OUTPUT_CONFLICT", "publication", "输出文件已存在。", None, None, None, False, "HUMAN_ACTION_REQUIRED")
        task_id = f"task-{uuid.uuid4().hex[:26]}"
        task_dir = self.tasks_root / task_id
        task_dir.mkdir()
        acceptance = acceptance_snapshot(request); acceptance_hash = stable_hash(acceptance)
        task = {"task_id": task_id, "created_at": _now(), "input_file": str(source), "input_sha256": input_hash, "output_file": str(output), "submission_hash": submission_hash, "acceptance_hash": acceptance_hash}
        _write_json(task_dir / "task.json", task); _write_json(task_dir / "acceptance-snapshot.json", acceptance)
        try:
            binding = self._bind(request, input_hash)
        except Exception as exc:
            response = self._error("EXECUTION_FAILED", "INTERNAL_ERROR", "workbook_inspection", str(exc), task_id, 1, None, False, "HUMAN_ACTION_REQUIRED")
            _write_json(task_dir / "current-state.json", {"state": "FAILED", "terminal": True, "request_revision": 1, "latest_attempt": None, "delivery_valid": False, "error": response["error"]})
            return response
        if binding.get("truncated") and binding.get("unresolved"):
            response = self._error("EXECUTION_FAILED", "INVENTORY_TOO_LARGE", "workbook_inspection", "字段清单超出上限，无法安全闭合绑定候选；请收窄来源查询。", task_id, 1, None, False, "HUMAN_ACTION_REQUIRED")
            _write_json(task_dir / "current-state.json", {"state": "FAILED", "terminal": True, "request_revision": 1, "latest_attempt": None, "delivery_valid": False, "error": response["error"]})
            return response
        binding_hash, request_revision_hash = self._revision_hashes(request, binding, 1)
        revision = {"request_revision": 1, "request": request, "bindings": binding, "binding_hash": binding_hash, "request_revision_hash": request_revision_hash}
        _write_json(task_dir / "revisions" / "revision-001.json", revision)
        if binding["unresolved"]:
            state = self._binding_response(task_id, acceptance_hash, binding)
            _write_json(task_dir / "current-state.json", {key: value for key, value in state.items() if key not in {"schema_version", "status", "binding_slots", "field_inventory", "diagnostics"}} | {"state": "NEEDS_BINDING", "terminal": False})
            return state
        return self._execute(task_dir, task, revision)

    def _bind(self, request: dict[str, Any], input_hash: str) -> dict[str, Any]:
        sheet = request["source"].get("sheet"); header_row = request["source"].get("header_row", 1)
        workbook = load_workbook(request["input_file"], read_only=True, data_only=True)
        try:
            if not sheet or sheet not in workbook.sheetnames:
                return {"source": None, "slots": [], "inventory": [], "truncated": False, "unresolved": [{"kind": "source", "logical_field": None}]}
            full = read_headers(Path(request["input_file"]), input_hash, sheet, header_row)
            inventory, truncated = cap_entries(full)
            by_header: dict[str, list[dict[str, Any]]] = {}
            for item in full: by_header.setdefault(item["header"], []).append(item)
            slots, unresolved = [], []
            for index, (field, references) in enumerate(collect_fields(request), 1):
                uses = self._field_uses(request, field); numeric = any(use in {"metric:sum", "metric:average"} for use in uses)
                matches = [x for x in by_header.get(field, []) if not numeric or x["inferred_type"] == "number"]
                status = "RESOLVED" if len(matches) == 1 and len(by_header.get(field, [])) == 1 else ("AMBIGUOUS" if matches else "UNRESOLVED")
                slot = {"id": f"binding-{index:03d}", "logical_field": field, "references": references, "requirements": {"accepted_types": ["number"] if numeric else ["text", "number", "boolean", "unknown"], "uses": uses}, "status": status, "selected_candidate_id": matches[0]["id"] if status == "RESOLVED" else None, "candidate_ids": [x["id"] for x in full if not numeric or x["inferred_type"] == "number"], "resolution": None}
                if status == "RESOLVED": slot["resolution"] = {"candidate_id": matches[0]["id"], "sheet": sheet, "header_row": header_row, "column": matches[0]["column"], "header": field}
                else: unresolved.append(slot)
                slots.append(slot)
            return {"source": {"id": "source-001", "sheet": sheet, "header_row": header_row}, "slots": slots, "inventory": inventory, "truncated": truncated, "unresolved": unresolved}
        finally: workbook.close()

    def _execute(self, task_dir: Path, task: dict[str, Any], revision: dict[str, Any]) -> dict[str, Any]:
        attempt_id = "attempt-001"; attempt_dir = task_dir / "attempts" / attempt_id; attempt_dir.mkdir(parents=True)
        request, binding = revision["request"], revision["bindings"]
        binding_snapshot = {"source": binding["source"], "slots": binding["slots"]}
        plan, plan_hash = compile_plan(task["task_id"], revision["request_revision"], request, task["acceptance_hash"], binding_snapshot, task["input_sha256"])
        _write_json(attempt_dir / "binding-snapshot.json", binding_snapshot); _write_json(attempt_dir / "internal-plan.json", plan)
        working = attempt_dir / "working-copy.xlsx"; shutil.copy2(task["input_file"], working)
        engine = OpenPyxlEngine.open(working)
        results: dict[str, Any] = {}
        try:
            for step in plan["steps"]:
                op = step["op"]
                if op == "read_table": results[step["id"]] = read_table(engine, step["sheet"], step["header_row"], step["columns"])
                elif op == "summarize_by_dimension":
                    table = results[step["input"]]
                    if "where" in step: table = filter_rows(table, step["where"])
                    table = aggregate(table, step["group_by"], step["metrics"])
                    if "sort" in step: table = sort_rows(table, step["sort"])
                    results[step["id"]] = table
                elif op == "project_columns":
                    results[step["id"]] = select_columns(results[step["input"]], step["fields"])
                elif op == "create_sheet":
                    if step["sheet"] in engine.sheet_names():
                        return self._attempt_failure(task_dir, task, attempt_id, "OUTPUT_CONFLICT", "execution", "目标工作表已存在。")
                    engine.create_sheet(step["sheet"])
                elif op == "write_table": self._write_table(engine, results[step["input"]], step["sheet"], step["anchor"])
            temporary = attempt_dir / "temporary-output.xlsx"; engine.save(temporary)
        except Exception as exc:
            return self._attempt_failure(task_dir, task, attempt_id, "EXECUTION_FAILED", "execution", str(exc))
        finally:
            engine.close()
        if sha256_file(Path(task["input_file"])) != task["input_sha256"]:
            return self._attempt_failure(task_dir, task, attempt_id, "INPUT_CHANGED", "validation", "输入文件在执行期间发生变化。")
        output = Path(task["output_file"]); output.parent.mkdir(parents=True, exist_ok=True)
        if output.exists(): return self._attempt_failure(task_dir, task, attempt_id, "OUTPUT_CONFLICT", "publication", "输出文件已存在。")
        staging = output.with_name(f".{output.name}.sheetpilot-{task['task_id']}-{attempt_id}.tmp")
        shutil.copy2(temporary, staging)
        try:
            os.link(staging, output)
        except FileExistsError:
            staging.unlink(missing_ok=True); return self._attempt_failure(task_dir, task, attempt_id, "OUTPUT_CONFLICT", "publication", "输出文件已存在。")
        except OSError:
            staging.unlink(missing_ok=True); return self._attempt_failure(task_dir, task, attempt_id, "PUBLICATION_FAILED", "publication", "当前文件系统不支持 no-replace 原子发布。")
        staging.unlink(missing_ok=True); published_hash = sha256_file(output)
        evidence = {"runtime_status": "RUNTIME_PASS", "internal_plan_hash": plan_hash, "validated_artifact_sha256": sha256_file(temporary), "published_sha256": published_hash, "coverage": plan["coverage"]}
        evidence["evidence_hash"] = stable_hash(evidence); _write_json(attempt_dir / "evidence.json", evidence)
        state = {"state": "RUNTIME_PASS", "terminal": True, "request_revision": revision["request_revision"], "latest_attempt": {"attempt_id": attempt_id, "state": "RUNTIME_PASS", "request_revision": revision["request_revision"], "finished_at": _now()}, "validated_published_sha256": published_hash, "delivery_valid": True, "artifact_integrity": "MATCHED", "evidence": evidence}
        _write_json(task_dir / "current-state.json", state)
        return {"schema_version": "1.0", "status": "RUNTIME_PASS", "task_id": task["task_id"], "request_revision": revision["request_revision"], "attempt_id": attempt_id, "output_file": str(output), "acceptance_hash": task["acceptance_hash"], "internal_plan_hash": plan_hash, "artifact_integrity": "MATCHED", "delivery_valid": True, "evidence": evidence}

    def _write_table(self, engine: OpenPyxlEngine, table: TableData, sheet: str, anchor: str) -> None:
        row, column = coordinate_to_tuple(anchor); ws = engine.worksheet(sheet)
        for offset, header in enumerate(table.headers): ws.cell(row, column + offset).value = header
        for row_offset, item in enumerate(table.rows, 1):
            for column_offset, header in enumerate(table.headers): ws.cell(row + row_offset, column + column_offset).value = item.get(header)

    def _field_uses(self, request: dict[str, Any], field: str) -> list[str]:
        uses = [f"filter:{item['operator']}" for item in request["filters"] if item["field"] == field]
        uses += ["dimension"] * sum(item["field"] == field for item in request["dimensions"])
        uses += [f"metric:{item['function']}" for item in request["metrics"] if item.get("field") == field]
        return uses

    def _binding_response(self, task_id: str, acceptance_hash: str, binding: dict[str, Any]) -> dict[str, Any]:
        unresolved = [item for item in binding["unresolved"] if item.get("id")]
        resolvable = [item for item in unresolved if item["candidate_ids"]]
        action = "PROVIDE_BINDING" if unresolved and len(resolvable) == len(unresolved) else "HUMAN_ACTION_REQUIRED"
        allowed = [{"op": "add", "path": f"/bindings/{item['id']}", "constraints": {"candidate_ids": item["candidate_ids"]}} for item in resolvable]
        diagnostics = [{"code": "FIELD_BINDING_AMBIGUOUS" if item["candidate_ids"] else "FIELD_BINDING_NOT_FOUND", "path": item["references"][0], "message": f"业务字段“{item['logical_field']}”无法唯一精确绑定。", "expected": {"kind": "single_binding"}, "actual": {"kind": "candidates", "count": len(item["candidate_ids"])}} for item in unresolved]
        if not binding["source"]:
            diagnostics = [{"code": "SOURCE_BINDING_REQUIRED", "path": "/source/sheet", "message": "来源 Sheet 无法唯一确定。", "expected": {"kind": "existing_sheet"}, "actual": {"kind": "unresolved"}}]
        return {"schema_version": "1.0", "status": "NEEDS_BINDING", "task_id": task_id, "request_revision": 1, "attempt_id": None, "acceptance_hash": acceptance_hash, "binding_slots": unresolved, "field_inventory": binding["inventory"], "truncated": binding["truncated"], "diagnostics": diagnostics, "recovery": {"action": action, "retryable": action == "PROVIDE_BINDING", "base_revision": 1, "allowed_amendments": allowed, "suggested_patch": []}}

    def _amend(self, amendment: dict[str, Any]) -> dict[str, Any]:
        task_id = amendment.get("task_id"); base = amendment.get("base_revision"); items = amendment.get("amendments")
        if not isinstance(task_id, str) or not task_id:
            return self._error("REQUEST_INVALID", "INVALID_VALUE", "field_binding", "修订必须携带 task_id。", None, base, None, False, "NONE")
        task_dir = self._task_dir(task_id); task_path = task_dir / "task.json"; state_path = task_dir / "current-state.json"
        if not task_path.is_file() or not state_path.is_file():
            return self._error("REQUEST_INVALID", "INVALID_REFERENCE", "contract_validation", "Task 不存在。", task_id, base, None, False, "NONE")
        task = json.loads(task_path.read_text(encoding="utf-8")); state = json.loads(state_path.read_text(encoding="utf-8")); current = state.get("request_revision", 1)
        if state.get("state") != "NEEDS_BINDING":
            return self._error("REQUEST_INVALID", "INVALID_COMBINATION", "field_binding", "当前任务不在待绑定状态。", task_id, base, None, False, "HUMAN_ACTION_REQUIRED")
        if base != current:
            return self._error("REQUEST_INVALID", "REVISION_CONFLICT", "field_binding", "修订基于过期版本。", task_id, current, None, True, "PROVIDE_BINDING")
        if not isinstance(items, list) or not items:
            return self._error("REQUEST_INVALID", "INVALID_VALUE", "field_binding", "amendments 必须是非空数组。", task_id, base, None, False, "NONE")
        revision = json.loads((task_dir / "revisions" / f"revision-{current:03d}.json").read_text(encoding="utf-8")); binding = copy.deepcopy(revision["bindings"])
        unresolved = {slot["id"]: slot for slot in binding["unresolved"] if slot.get("id")}; entries = {item["id"]: item for item in binding["inventory"]}; selected = {}
        for item in items:
            if not isinstance(item, dict) or item.get("op") not in {"add", "replace"}:
                return self._error("REQUEST_INVALID", "INVALID_ENUM", "field_binding", "修订只支持 add/replace。", task_id, base, None, False, "NONE")
            path = item.get("path", "")
            if not path.startswith("/bindings/"):
                return self._error("REQUEST_INVALID", "INVALID_REFERENCE", "field_binding", "该路径会改变冻结请求组件，请创建新 Task。", task_id, base, None, False, "CREATE_NEW_TASK")
            slot = unresolved.get(path.removeprefix("/bindings/")); candidate_id = (item.get("value") or {}).get("candidate_id")
            if slot is None:
                return self._error("REQUEST_INVALID", "INVALID_REFERENCE", "field_binding", "修订路径不在允许范围内。", task_id, base, None, False, "PROVIDE_BINDING")
            if candidate_id not in slot["candidate_ids"] or candidate_id not in entries:
                return self._error("REQUEST_INVALID", "INVALID_VALUE", "field_binding", "候选不在该 Slot 的允许范围内。", task_id, base, None, False, "PROVIDE_BINDING")
            selected[slot["id"]] = candidate_id
        if set(selected) != set(unresolved):
            return self._error("REQUEST_INVALID", "INVALID_COMBINATION", "field_binding", "一次修订必须解决全部待绑定 Slot。", task_id, base, None, False, "PROVIDE_BINDING")
        for slot_id, candidate_id in selected.items():
            slot = unresolved[slot_id]; entry = entries[candidate_id]
            slot.update({"status": "RESOLVED", "selected_candidate_id": candidate_id, "resolution": {"candidate_id": candidate_id, "sheet": entry["sheet"], "header_row": entry["header_row_start"], "column": entry["column"], "header": entry["header"]}})
        binding["slots"] = [slot if slot["id"] not in unresolved else unresolved[slot["id"]] for slot in binding["slots"]]; binding["unresolved"] = []
        next_revision = {"request_revision": current + 1, "request": revision["request"], "bindings": binding}
        next_revision["binding_hash"], next_revision["request_revision_hash"] = self._revision_hashes(next_revision["request"], binding, current + 1)
        _write_json(task_dir / "revisions" / f"revision-{current + 1:03d}.json", next_revision)
        return self._execute(task_dir, task, next_revision)

    def _revision_hashes(self, request: dict[str, Any], binding: dict[str, Any], revision_number: int) -> tuple[str, str]:
        snapshot = {"source": binding["source"], "bindings": {slot["id"]: slot.get("selected_candidate_id") for slot in binding["slots"] if slot.get("selected_candidate_id")}}
        return stable_hash(snapshot), stable_hash([revision_number, request, snapshot])

    def _attempt_failure(self, task_dir: Path, task: dict[str, Any], attempt_id: str, code: str, phase: str, message: str) -> dict[str, Any]:
        response = self._error("EXECUTION_FAILED" if phase == "execution" else "PUBLICATION_FAILED", code, phase, message, task["task_id"], 1, attempt_id, phase != "publication", "RETRY_ATTEMPT" if phase != "publication" else "HUMAN_ACTION_REQUIRED")
        _write_json(task_dir / "current-state.json", {"state": "FAILED", "terminal": False, "request_revision": 1, "latest_attempt": {"attempt_id": attempt_id, "state": response["status"], "request_revision": 1, "finished_at": _now()}, "delivery_valid": False, "recovery": response["error"]["recovery"]})
        return response

    def _error(self, status: str, code: str, phase: str, message: str, task_id: str | None, revision: int | None, attempt_id: str | None, retryable: bool, action: str) -> dict[str, Any]:
        return {"schema_version": "1.0", "status": status, "task_id": task_id, "request_revision": revision, "attempt_id": attempt_id, "error": {"code": code, "phase": phase, "message": message, "retryable": retryable, "diagnostics": [], "recovery": {"action": action, "retryable": retryable, "base_revision": revision, "allowed_amendments": [], "suggested_patch": []}}}

    def _task_dir(self, task_id: str) -> Path:
        suffix = task_id.removeprefix("task-")
        if not task_id.startswith("task-") or not suffix or any(char not in "0123456789abcdef" for char in suffix):
            return self.tasks_root / "invalid"
        return self.tasks_root / task_id

    def _find_submission(self, submission_hash: str, input_hash: str, output_file: str) -> str | None:
        if not self.tasks_root.is_dir():
            return None
        for path in self.tasks_root.glob("task-*/task.json"):
            try: task = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError): continue
            if task.get("submission_hash") == submission_hash and task.get("input_sha256") == input_hash and task.get("output_file") == output_file and (path.parent / "current-state.json").is_file():
                return task["task_id"]
        return None
