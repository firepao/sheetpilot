from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .contract import stable_hash

MAX_INVENTORY_ENTRIES = 64
MAX_INVENTORY_BYTES = 65536
MAX_SAMPLES = 3
MAX_SAMPLE_CHARS = 64
HEADER_SCORE_FLOOR = 0.5
HEADER_SCORE_MARGIN = 0.1
_PHONE = re.compile(r"^1\d{10}$")
_ID_CARD = re.compile(r"^\d{17}[\dXx]$")
_EMAIL = re.compile(r"^([^@]{1,8})[^@]*(@.+)$")


def _mask(value: str) -> str:
    if _PHONE.match(value):
        return value[:3] + "****" + value[-4:]
    if _ID_CARD.match(value):
        return value[:6] + "********" + value[-4:]
    match = _EMAIL.match(value)
    return match.group(1) + "***" + match.group(2) if match else value


def _samples(values: list[Any]) -> list[Any]:
    result = []
    for value in values[:MAX_SAMPLES]:
        result.append(_mask(value[:MAX_SAMPLE_CHARS]) if isinstance(value, str) else value)
    return result


def read_headers(source: Path, input_sha256: str, sheet: str, header_row: int) -> list[dict[str, Any]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        ws = workbook[sheet]
        if ws.max_column is None or ws.max_row is None:
            ws.calculate_dimension(force=True)
        entries = []
        for column in range(1, (ws.max_column or 0) + 1):
            raw = ws.cell(header_row, column).value
            if raw in (None, ""):
                continue
            values = [ws.cell(row, column).value for row in range(header_row + 1, min(ws.max_row or header_row, header_row + 20) + 1)]
            non_empty = [v for v in values if v not in (None, "")]
            header = str(raw)
            entries.append({
                "id": "candidate-" + stable_hash([input_sha256, sheet, header_row, column, header])[:8],
                "source_id": "source-001", "sheet": sheet,
                "header_row_start": header_row, "header_row_end": header_row,
                "column": get_column_letter(column), "header": header,
                "inferred_type": "number" if non_empty and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in non_empty) else "text",
                "null_ratio": round(1 - len(non_empty) / max(1, len(values)), 4),
                "sample_values": _samples(non_empty),
                "neighbor_headers": [str(ws.cell(header_row, c).value) for c in (column - 1, column + 1) if 1 <= c <= (ws.max_column or 0) and ws.cell(header_row, c).value not in (None, "")],
            })
        return entries
    finally:
        workbook.close()


def cap_entries(entries: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], bool]:
    result = list(entries[:MAX_INVENTORY_ENTRIES])
    truncated = len(result) != len(entries)
    while result and len(json.dumps(result, ensure_ascii=False).encode()) > MAX_INVENTORY_BYTES:
        result.pop(); truncated = True
    return result, truncated


def _score(ws, row: int, max_col: int) -> float:
    values = [ws.cell(row, col).value for col in range(1, max_col + 1)]
    present = [v for v in values if v not in (None, "")]
    if not present: return 0.0
    return len(present) / max_col * .4 + sum(isinstance(v, str) for v in present) / len(present) * .4 + len({str(v) for v in present}) / len(present) * .2


def detect_header_candidates(ws) -> list[int]:
    max_col = min(ws.max_column or 0, 200)
    scored = [(row, _score(ws, row, max_col)) for row in range(1, min(ws.max_row or 1, 10) + 1)]
    best = max((score for _, score in scored), default=0)
    return [row for row, score in scored if best >= HEADER_SCORE_FLOOR and score >= best - HEADER_SCORE_MARGIN]


def build_inventory(source: Path, input_sha256: str, sheet: str | None = None, header_row: int | None = None) -> dict[str, Any]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        names = [n for n in workbook.sheetnames if workbook[n].sheet_state == "visible"]
        if sheet is not None:
            if sheet not in names: raise KeyError(sheet)
            names = [sheet]
        result = {"header_candidates": [], "field_inventory": [], "truncated": False, "error_code": None}
        for name in names:
            ws = workbook[name]; candidates = detect_header_candidates(ws) if header_row is None else []
            row = header_row if header_row is not None else (candidates[0] if len(candidates) == 1 else None)
            if row is None:
                result["header_candidates"].extend({"sheet": name, "header_row": r, "confidence": 1.0, "evidence": ["header_candidate_detected"]} for r in candidates); continue
            entries, truncated = cap_entries(read_headers(source, input_sha256, name, row))
            result["field_inventory"].extend(entries); result["truncated"] |= truncated
        if sheet is None and len(json.dumps(result["field_inventory"], ensure_ascii=False).encode()) > MAX_INVENTORY_BYTES:
            result.update({"header_candidates": [], "field_inventory": [], "truncated": False, "error_code": "INVENTORY_TOO_LARGE"})
        return result
    finally:
        workbook.close()
