from __future__ import annotations

import math
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils import column_index_from_string

from ..workspace import sha256_file


_MISSING = object()
_MAX_EXTRA_CELLS = 100


def _safe_value(value: Any) -> Any:
    """把验收报告中的单元格值转换成稳定、可 JSON 序列化的值。"""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.hex()
    if isinstance(value, dict):
        return {str(key): _safe_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_safe_value(item) for item in value]
    return value


def _values_equal(expected: Any, actual: Any, tolerance: float) -> bool:
    if isinstance(expected, (int, float)) and not isinstance(expected, bool) and isinstance(actual, (int, float)) and not isinstance(actual, bool):
        return math.isclose(float(expected), float(actual), rel_tol=tolerance, abs_tol=tolerance)
    return expected == actual


def _check(checks: list[dict[str, Any]], name: str, passed: bool, expected: Any = _MISSING, actual: Any = _MISSING, path: str | None = None) -> None:
    item: dict[str, Any] = {"name": name, "passed": bool(passed)}
    if expected is not _MISSING:
        item["expected"] = _safe_value(expected)
    if actual is not _MISSING:
        item["actual"] = _safe_value(actual)
    if path is not None:
        item["path"] = path
    checks.append(item)


def _load(path: Path):
    return load_workbook(path, data_only=False, keep_links=True, keep_vba=path.suffix.lower() == ".xlsm")


def _protection_snapshot(ws) -> dict[str, Any]:
    protection = ws.protection
    keys = (
        "sheet", "objects", "scenarios", "formatCells", "formatColumns", "formatRows",
        "insertColumns", "insertRows", "insertHyperlinks", "deleteColumns", "deleteRows",
        "selectLockedCells", "selectUnlockedCells", "sort", "autoFilter", "pivotTables",
    )
    return {key: getattr(protection, key, None) for key in keys} | {"password": protection.password}


def _sheet_structure(ws) -> dict[str, Any]:
    return {
        "sheet_state": ws.sheet_state,
        "dimension": ws.calculate_dimension(),
        "merged_ranges": sorted(str(item) for item in ws.merged_cells.ranges),
        "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes is not None else None,
        "auto_filter": ws.auto_filter.ref,
        "protection": _protection_snapshot(ws),
        "chart_count": len(ws._charts),
        "image_count": len(ws._images),
    }


def _sheet_values(ws) -> list[list[Any]]:
    return [[ws.cell(row, column).value for column in range(1, ws.max_column + 1)] for row in range(1, ws.max_row + 1)]


def _first_value_difference(expected: list[list[Any]], actual: list[list[Any]]) -> dict[str, Any] | None:
    row_count = max(len(expected), len(actual))
    for row_index in range(row_count):
        expected_row = expected[row_index] if row_index < len(expected) else []
        actual_row = actual[row_index] if row_index < len(actual) else []
        column_count = max(len(expected_row), len(actual_row))
        for column_index in range(column_count):
            expected_value = expected_row[column_index] if column_index < len(expected_row) else None
            actual_value = actual_row[column_index] if column_index < len(actual_row) else None
            if expected_value != actual_value:
                return {"cell": f"R{row_index + 1}C{column_index + 1}", "expected": _safe_value(expected_value), "actual": _safe_value(actual_value)}
    return None


def _binding_columns(binding_snapshot: dict[str, Any]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for slot in binding_snapshot.get("slots", []):
        resolution = slot.get("resolution") or {}
        logical_field = slot.get("logical_field")
        column = resolution.get("column")
        if logical_field and column:
            columns[logical_field] = column_index_from_string(column)
    return columns


def _matches(actual: Any, operator: str, expected: Any) -> bool:
    if operator == "eq":
        return actual == expected
    if operator == "ne":
        return actual != expected
    if operator == "gt":
        return actual is not None and actual > expected
    if operator == "gte":
        return actual is not None and actual >= expected
    if operator == "lt":
        return actual is not None and actual < expected
    if operator == "lte":
        return actual is not None and actual <= expected
    if operator == "in":
        return actual in expected
    raise ValueError(f"unsupported filter op: {operator}")


def _sort_key(row: dict[str, Any], field: str) -> tuple[bool, Any]:
    value = row.get(field)
    if isinstance(value, (int, float, date, datetime)) and not isinstance(value, bool):
        return value is None, value
    return value is None, str(value) if value is not None else ""


def _independent_result(input_path: Path, request: dict[str, Any], binding_snapshot: dict[str, Any]) -> tuple[list[str], list[dict[str, Any]], int, int]:
    source = binding_snapshot["source"]
    columns = _binding_columns(binding_snapshot)
    workbook = _load(input_path)
    try:
        ws = workbook[source["sheet"]]
        source_rows: list[dict[str, Any]] = []
        for row_number in range(source["header_row"] + 1, ws.max_row + 1):
            row = {field: ws.cell(row_number, column).value for field, column in columns.items()}
            if any(value not in (None, "") for value in row.values()):
                source_rows.append(row)
    finally:
        workbook.close()

    filtered: list[dict[str, Any]] = []
    for row in source_rows:
        try:
            accepted = all(_matches(row.get(item["field"]), item["operator"], item.get("value")) for item in request["filters"])
        except (TypeError, ValueError):
            accepted = False
        if accepted:
            filtered.append(row)

    dimensions = request["dimensions"]
    metrics = request["metrics"]
    buckets: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in filtered:
        key = tuple(row.get(item["field"]) if row.get(item["field"]) not in (None, "") else "(空)" for item in dimensions)
        bucket = buckets.setdefault(key, {"__rows__": 0, "__averages__": {}})
        bucket["__rows__"] += 1
        for metric in metrics:
            function = metric["function"]
            metric_id = metric["id"]
            value = row.get(metric.get("field")) if metric.get("field") else None
            if function == "count":
                if metric["mode"] == "rows":
                    bucket[metric_id] = bucket.get(metric_id, 0) + 1
                elif value not in (None, ""):
                    bucket[metric_id] = bucket.get(metric_id, 0) + 1
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                if function == "sum":
                    bucket[metric_id] = bucket.get(metric_id, 0) + value
                elif function == "average":
                    bucket["__averages__"].setdefault(metric_id, []).append(value)

    rows: list[dict[str, Any]] = []
    for key in sorted(buckets, key=lambda item: tuple(str(value) for value in item)):
        bucket = buckets[key]
        result = {item["id"]: value for item, value in zip(dimensions, key)}
        for metric in metrics:
            metric_id = metric["id"]
            if metric["function"] == "average":
                values = bucket["__averages__"].get(metric_id, [])
                result[metric_id] = sum(values) / len(values) if values else None
            else:
                result[metric_id] = bucket.get(metric_id, 0)
        rows.append(result)

    # 复刻稳定的多键排序顺序，但排序字段仍使用请求组件 id，避免依赖执行器符号。
    for item in reversed(request["output"].get("sort", [])):
        rows.sort(key=lambda row, field=item["by"]: _sort_key(row, field), reverse=item["direction"] == "desc")
    headers = [item["output_name"] for item in dimensions] + [item["output_name"] for item in metrics]
    projected = []
    for row in rows:
        projected.append({
            **{item["output_name"]: row[item["id"]] for item in dimensions},
            **{item["output_name"]: row[item["id"]] for item in metrics},
        })
    return headers, projected, len(filtered), len(buckets)


def validate_summarize_table_artifact(
    input_path: Path,
    output_path: Path,
    request: dict[str, Any],
    binding_snapshot: dict[str, Any],
    expected_input_sha256: str,
    tolerance: float = 1e-6,
) -> dict[str, Any]:
    checks: list[dict[str, Any]] = []
    result: dict[str, Any] = {"schema_version": "1.0", "passed": False, "checks": checks}
    try:
        actual_input_sha256 = sha256_file(input_path)
        _check(checks, "input_sha256_unchanged", actual_input_sha256 == expected_input_sha256, expected_input_sha256, actual_input_sha256)
        result["input_unchanged"] = actual_input_sha256 == expected_input_sha256

        expected_headers, expected_rows, filtered_count, group_count = _independent_result(input_path, request, binding_snapshot)
        result.update({"filtered_row_count": filtered_count, "group_count": group_count})

        output_workbook = _load(output_path)
        input_workbook = _load(input_path)
        try:
            expected_sheet_names = input_workbook.sheetnames + [request["output"]["sheet"]]
            _check(checks, "sheet_names_and_order", output_workbook.sheetnames == expected_sheet_names, expected_sheet_names, output_workbook.sheetnames)

            for sheet_name in input_workbook.sheetnames:
                if sheet_name not in output_workbook.sheetnames:
                    _check(checks, f"source_sheet_{sheet_name}_preserved", False, sheet_name, None)
                    continue
                expected_ws, actual_ws = input_workbook[sheet_name], output_workbook[sheet_name]
                expected_structure, actual_structure = _sheet_structure(expected_ws), _sheet_structure(actual_ws)
                _check(checks, f"source_sheet_{sheet_name}_structure_preserved", expected_structure == actual_structure, expected_structure, actual_structure)
                expected_values, actual_values = _sheet_values(expected_ws), _sheet_values(actual_ws)
                value_difference = _first_value_difference(expected_values, actual_values)
                _check(checks, f"source_sheet_{sheet_name}_content_preserved", value_difference is None, "unchanged", value_difference or "unchanged")

            target_name = request["output"]["sheet"]
            if target_name not in output_workbook.sheetnames:
                _check(checks, "target_sheet_exists", False, target_name, None)
            else:
                _check(checks, "target_sheet_exists", True, target_name, target_name)
                target = output_workbook[target_name]
                target_structure = {
                    "sheet_state": target.sheet_state,
                    "merged_ranges": sorted(str(item) for item in target.merged_cells.ranges),
                    "freeze_panes": str(target.freeze_panes) if target.freeze_panes is not None else None,
                    "auto_filter": target.auto_filter.ref,
                    "chart_count": len(target._charts),
                    "image_count": len(target._images),
                }
                expected_target_structure = {"sheet_state": "visible", "merged_ranges": [], "freeze_panes": None, "auto_filter": None, "chart_count": 0, "image_count": 0}
                _check(checks, "target_sheet_structure_valid", target_structure == expected_target_structure, expected_target_structure, target_structure)
                start_row, start_column = coordinate_to_tuple(request["output"].get("anchor", "A1"))
                actual_headers = [target.cell(start_row, start_column + offset).value for offset in range(len(expected_headers))]
                _check(checks, "output_headers_match", actual_headers == expected_headers, expected_headers, actual_headers)

                occupied_rows = [
                    row_number
                    for row_number in range(start_row + 1, target.max_row + 1)
                    if any(target.cell(row_number, start_column + offset).value not in (None, "") for offset in range(len(expected_headers)))
                ]
                actual_row_count = max(occupied_rows) - start_row if occupied_rows else 0
                actual_rows = []
                for row_offset in range(actual_row_count):
                    actual_rows.append({header: target.cell(start_row + 1 + row_offset, start_column + offset).value for offset, header in enumerate(expected_headers)})
                row_values_match = len(actual_rows) == len(expected_rows) and all(
                    all(_values_equal(expected.get(header), actual.get(header), tolerance) for header in expected_headers)
                    for expected, actual in zip(expected_rows, actual_rows)
                )
                first_row_difference = None
                if not row_values_match:
                    for index in range(max(len(expected_rows), len(actual_rows))):
                        expected_row = expected_rows[index] if index < len(expected_rows) else None
                        actual_row = actual_rows[index] if index < len(actual_rows) else None
                        if expected_row is None or actual_row is None or any(not _values_equal(expected_row.get(header), actual_row.get(header), tolerance) for header in expected_headers):
                            first_row_difference = {"row": index + 1, "expected": _safe_value(expected_row), "actual": _safe_value(actual_row)}
                            break
                _check(checks, "output_values_and_order_match", row_values_match, "all rows match", first_row_difference or "all rows match")
                result["output_row_count"] = actual_row_count
                _check(checks, "output_row_count_match", actual_row_count == len(expected_rows), len(expected_rows), actual_row_count)

                expected_cells = {(start_row + row_offset, start_column + column_offset) for row_offset in range(len(expected_rows) + 1) for column_offset in range(len(expected_headers))}
                extras = []
                for row in target.iter_rows():
                    for cell in row:
                        if cell.value not in (None, "") and (cell.row, cell.column) not in expected_cells:
                            extras.append({"coordinate": cell.coordinate, "value": _safe_value(cell.value)})
                _check(checks, "output_has_no_extra_non_empty_cells", not extras, [], extras[:_MAX_EXTRA_CELLS])
        finally:
            input_workbook.close()
            output_workbook.close()
    except Exception as exc:
        _check(checks, "validator_completed", False, "无异常", str(exc))

    result["passed"] = all(item["passed"] for item in checks)
    return result
