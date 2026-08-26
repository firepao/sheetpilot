from __future__ import annotations

from copy import copy
from typing import Any

from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

from .styles import apply_preset
from .tables import TableData


def range_ref(sheet: str, address: str) -> dict[str, Any]:
    min_col, min_row, max_col, max_row = range_boundaries(address)
    return {"sheet": sheet, "min_row": min_row, "min_col": min_col, "max_row": max_row, "max_col": max_col}


def inspect_workbook(engine) -> dict:
    return {"sheets": [inspect_sheet(engine, name) for name in engine.sheet_names()]}


def inspect_sheet(engine, name: str) -> dict:
    ws = engine.worksheet(name)
    return {"name": name, "visibility": ws.sheet_state, "max_row": ws.max_row, "max_column": ws.max_column, "merged_ranges": [str(x) for x in ws.merged_cells.ranges]}


def validate_structure(engine, required_sheets=None, forbidden_sheets=None) -> dict:
    names = set(engine.sheet_names()); required = set(required_sheets or []); forbidden = set(forbidden_sheets or [])
    return {"valid": required <= names and not (forbidden & names), "missing": sorted(required - names), "forbidden_present": sorted(forbidden & names)}


def read_range(engine, sheet: str, address: str) -> list[list[Any]]:
    ws = engine.worksheet(sheet); ref = range_ref(sheet, address)
    return [[ws.cell(row, col).value for col in range(ref["min_col"], ref["max_col"] + 1)] for row in range(ref["min_row"], ref["max_row"] + 1)]


def write_range(engine, sheet: str, anchor: str, values: list[list[Any]]) -> dict:
    row, col = coordinate_to_tuple(anchor); ws = engine.worksheet(sheet)
    width = max((len(item) for item in values), default=0)
    for r_offset, values_row in enumerate(values):
        for c_offset, value in enumerate(values_row): ws.cell(row + r_offset, col + c_offset, value)
    return {"sheet": sheet, "min_row": row, "min_col": col, "max_row": row + max(len(values) - 1, 0), "max_col": col + max(width - 1, 0)}


def write_table(engine, sheet: str, anchor: str, data: TableData, include_header=True) -> dict:
    values = ([data.headers] if include_header else []) + [[row.get(field) for field in data.headers] for row in data.rows]
    ref = write_range(engine, sheet, anchor, values); ref["data"] = data
    return ref


def append_rows(engine, sheet: str, values: list[list[Any]]) -> dict:
    ws = engine.worksheet(sheet); start = ws.max_row + 1
    for row in values: ws.append(row)
    return {"sheet": sheet, "start_row": start, "end_row": ws.max_row, "rows": len(values)}


def clear_range(engine, sheet: str, address: str) -> dict:
    ref = range_ref(sheet, address); ws = engine.worksheet(sheet)
    for row in ws.iter_rows(min_row=ref["min_row"], max_row=ref["max_row"], min_col=ref["min_col"], max_col=ref["max_col"]):
        for cell in row: cell.value = None
    return ref


def copy_range(engine, source_sheet: str, source: str, target_sheet: str, anchor: str, include_style=True) -> dict:
    source_ref = range_ref(source_sheet, source); target_row, target_col = coordinate_to_tuple(anchor)
    source_ws, target_ws = engine.worksheet(source_sheet), engine.worksheet(target_sheet)
    for r_offset, row in enumerate(source_ws.iter_rows(min_row=source_ref["min_row"], max_row=source_ref["max_row"], min_col=source_ref["min_col"], max_col=source_ref["max_col"])):
        for c_offset, source_cell in enumerate(row):
            target = target_ws.cell(target_row + r_offset, target_col + c_offset, source_cell.value)
            if include_style:
                target._style = copy(source_cell._style); target.number_format = source_cell.number_format
    return {"sheet": target_sheet, "min_row": target_row, "min_col": target_col, "max_row": target_row + source_ref["max_row"] - source_ref["min_row"], "max_col": target_col + source_ref["max_col"] - source_ref["min_col"]}


def apply_style(engine, ref: dict, preset: str) -> dict:
    apply_preset(engine.worksheet(ref["sheet"]), ref["min_row"], ref["min_col"], ref["max_row"], ref["max_col"], preset); return ref


def copy_style(engine, source_sheet: str, source: str, target_sheet: str, target: str) -> dict:
    source_ref, target_ref = range_ref(source_sheet, source), range_ref(target_sheet, target)
    if (source_ref["max_row"]-source_ref["min_row"], source_ref["max_col"]-source_ref["min_col"]) != (target_ref["max_row"]-target_ref["min_row"], target_ref["max_col"]-target_ref["min_col"]): raise ValueError("style ranges must have equal shape")
    sws, tws = engine.worksheet(source_sheet), engine.worksheet(target_sheet)
    for r in range(source_ref["max_row"]-source_ref["min_row"]+1):
        for c in range(source_ref["max_col"]-source_ref["min_col"]+1): tws.cell(target_ref["min_row"]+r,target_ref["min_col"]+c)._style=copy(sws.cell(source_ref["min_row"]+r,source_ref["min_col"]+c)._style)
    return target_ref
