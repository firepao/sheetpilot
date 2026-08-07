from __future__ import annotations

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

from ..errors import ErrorCode, SheetPilotError
from .charts import create_chart as add_chart
from .formulas import add_formula_column
from .styles import apply_preset
from .tables import TableData, aggregate as aggregate_data, derive_column as derive_table, filter_rows as filter_table, read_table as load_table, select_columns as select_table, sort_rows as sort_table


class WorkbookContext:
    def __init__(self, engine, recorder, run_dir=None):
        self.engine, self.recorder = engine, recorder
        self.run_dir = run_dir
        self.created_sheets: set[str] = set()

    def read_table(self, sheet, header_rows, columns):
        data = load_table(self.engine, sheet, max(header_rows), columns)
        self.recorder.record("read", sheet=sheet, rows=len(data.rows), fields=list(columns))
        return data

    def aggregate(self, data, **params): return aggregate_data(data, params["group_by"], params["metrics"], params.get("date_field"), params.get("current_period"))

    def filter_rows(self, data, where, invalid_value_policy="exclude"):
        result = filter_table(data, where, invalid_value_policy)
        self.recorder.record("transform", operation="filter_rows", input_rows=len(data.rows), output_rows=len(result.rows), excluded_rows=len(data.rows) - len(result.rows), count=0)
        return result

    def select_columns(self, data, fields): return select_table(data, fields)
    def sort_rows(self, data, keys): return sort_table(data, keys)
    def derive_column(self, data, name, template, arguments): return derive_table(data, name, template, arguments)

    def create_sheet(self, name):
        if name in self.engine.sheet_names(): raise SheetPilotError(ErrorCode.PLAN_INVALID, "工作表已存在")
        self.engine.create_sheet(name); self.created_sheets.add(name); self.recorder.record("create_sheet", sheet=name)

    def write_table(self, sheet: str, anchor: str, data: TableData):
        if sheet not in self.created_sheets: raise SheetPilotError(ErrorCode.POLICY_BLOCKED, "禁止写入非新建工作表")
        row, col = coordinate_to_tuple(anchor); ws = self.engine.worksheet(sheet)
        for offset, header in enumerate(data.headers): ws.cell(row, col + offset, header)
        for r_offset, item in enumerate(data.rows, 1):
            for c_offset, header in enumerate(data.headers): ws.cell(row + r_offset, col + c_offset, item.get(header))
        count = len(data.headers) * (len(data.rows) + 1)
        self.recorder.record("write", sheet=sheet, range=f"{anchor}:{get_column_letter(col + len(data.headers)-1)}{row + len(data.rows)}", count=count)
        return {"sheet": sheet, "min_row": row, "min_col": col, "max_row": row + len(data.rows), "max_col": col + len(data.headers) - 1, "data": data}

    def apply_style(self, ref, preset):
        apply_preset(self.engine.worksheet(ref["sheet"]), ref["min_row"], ref["min_col"], ref["max_row"], ref["max_col"], preset)
        self.recorder.record("style", sheet=ref["sheet"], count=0)

    def freeze(self, sheet, cell): self.engine.worksheet(sheet).freeze_panes = cell; self.recorder.record("freeze", sheet=sheet, cell=cell, count=0)

    def create_chart(self, ref, chart_type, anchor, title, category_fields=None, series_fields=None):
        add_chart(self.engine.worksheet(ref["sheet"]), ref["min_row"], ref["min_col"], ref["max_row"], ref["max_col"], chart_type, anchor, title, category_fields, series_fields)
        self.recorder.record("chart", sheet=ref["sheet"], anchor=anchor)

    def add_formula_column(self, ref, header, template, arguments, number_format=None):
        result = add_formula_column(self.engine.worksheet(ref["sheet"]), ref, header, template, arguments, number_format)
        self.recorder.record("formula", sheet=ref["sheet"], template=template, count=0)
        return result
