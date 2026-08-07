from __future__ import annotations

from openpyxl.chart import BarChart, LineChart, PieChart, Reference

from ..errors import ErrorCode, SheetPilotError


def create_chart(ws, min_row: int, min_col: int, max_row: int, max_col: int, chart_type: str, anchor: str, title: str, category_fields: list[str] | None = None, series_fields: list[str] | None = None):
    classes = {"bar": BarChart, "column": BarChart, "line": LineChart, "pie": PieChart}
    if chart_type not in classes or max_row <= min_row or max_col <= min_col:
        raise SheetPilotError(ErrorCode.CAPABILITY_UNSUPPORTED, "图表类型不支持或数据为空")
    headers = {ws.cell(min_row, col).value: col for col in range(min_col, max_col + 1)}
    category_fields = category_fields or [ws.cell(min_row, min_col).value]
    series_fields = series_fields or [ws.cell(min_row, max_col).value]
    if any(field not in headers for field in [*category_fields, *series_fields]):
        raise SheetPilotError(ErrorCode.PLAN_INVALID, "图表字段不在结果区域中")
    chart = classes[chart_type](); chart.title = title
    for field in series_fields:
        chart.add_data(Reference(ws, min_col=headers[field], max_col=headers[field], min_row=min_row, max_row=max_row), titles_from_data=True)
    category_columns = [headers[field] for field in category_fields]
    if category_columns != list(range(min(category_columns), max(category_columns) + 1)):
        raise SheetPilotError(ErrorCode.PLAN_INVALID, "图表分类字段必须连续")
    chart.set_categories(Reference(ws, min_col=min(category_columns), max_col=max(category_columns), min_row=min_row + 1, max_row=max_row))
    chart.height, chart.width = 8, 14
    ws.add_chart(chart, anchor)
    return chart
