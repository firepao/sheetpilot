from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def apply_business_table(ws, min_row: int, min_col: int, max_row: int, max_col: int) -> None:
    blue, white, border = "1F4E78", "FFFFFF", Side(style="thin", color="D9E2F3")
    for cell in ws[min_row]:
        if min_col <= cell.column <= max_col:
            cell.font = Font(bold=True, color=white); cell.fill = PatternFill("solid", fgColor=blue); cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row: cell.border = Border(bottom=border)
    for col in range(min_col, max_col + 1):
        letter = ws.cell(1, col).column_letter
        width = max(len(str(ws.cell(row, col).value or "")) for row in range(min_row, max_row + 1)) + 2
        ws.column_dimensions[letter].width = min(max(width, 10), 30)


def apply_preset(ws, min_row: int, min_col: int, max_row: int, max_col: int, preset: str) -> None:
    if preset == "business_table": return apply_business_table(ws, min_row, min_col, max_row, max_col)
    if preset == "title":
        for cell in ws[min_row]:
            if min_col <= cell.column <= max_col: cell.font = Font(bold=True, size=14, color="1F4E78")
        return
    if preset == "total_row":
        for cell in ws[max_row]:
            if min_col <= cell.column <= max_col: cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="E2F0D9")
        return
    if preset in {"input_area", "formula_area"}:
        fill = "FFF2CC" if preset == "input_area" else "EDEDED"
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row: cell.fill = PatternFill("solid", fgColor=fill)
        return
    raise ValueError(f"unknown style preset: {preset}")
