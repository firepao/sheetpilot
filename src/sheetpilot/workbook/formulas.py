from __future__ import annotations

from openpyxl.utils import get_column_letter

from ..errors import ErrorCode, SheetPilotError


def add_formula_column(ws, ref: dict, header: str, template: str, arguments: dict[str, str], number_format: str | None = None) -> dict:
    if template not in {"safe_ratio", "subtract", "add", "multiply"}:
        raise SheetPilotError(ErrorCode.CAPABILITY_UNSUPPORTED, "公式模板未注册", {"template": template})
    headers = {ws.cell(ref["min_row"], col).value: col for col in range(ref["min_col"], ref["max_col"] + 1)}
    if header in headers or any(field not in headers for field in arguments.values()):
        raise SheetPilotError(ErrorCode.PLAN_INVALID, "公式列参数或来源列无效")
    target = ref["max_col"] + 1
    ws.cell(ref["min_row"], target, header)
    for row in range(ref["min_row"] + 1, ref["max_row"] + 1):
        def cell(name): return f"{get_column_letter(headers[name])}{row}"
        if template == "safe_ratio": formula = f'=IFERROR({cell(arguments["numerator"])}/{cell(arguments["denominator"])},0)'
        elif template == "subtract": formula = f"={cell(arguments['left'])}-{cell(arguments['right'])}"
        elif template == "add": formula = f"={cell(arguments['left'])}+{cell(arguments['right'])}"
        else: formula = f"={cell(arguments['left'])}*{cell(arguments['right'])}"
        ws.cell(row, target, formula)
        if number_format: ws.cell(row, target).number_format = number_format
    return {**ref, "max_col": target, "formula_template": template, "representative_formula": ws.cell(ref["min_row"] + 1, target).value if ref["max_row"] > ref["min_row"] else None}
