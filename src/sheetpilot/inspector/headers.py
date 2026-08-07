from __future__ import annotations

from openpyxl.utils import get_column_letter

from ..models import FieldProfile, HeaderCandidate
from .samples import summarize


def detect_headers(ws, scan_rows: int = 10) -> list[HeaderCandidate]:
    best_row, best_score = 1, -1.0
    max_col = min(ws.max_column, 200)
    for row in range(1, min(ws.max_row, scan_rows) + 1):
        values = [ws.cell(row, col).value for col in range(1, max_col + 1)]
        present = [v for v in values if v not in (None, "")]
        if not present:
            continue
        text_ratio = sum(isinstance(v, str) for v in present) / len(present)
        uniqueness = len({str(v) for v in present}) / len(present)
        score = len(present) / max(max_col, 1) * 0.4 + text_ratio * 0.4 + uniqueness * 0.2
        if score > best_score:
            best_row, best_score = row, score
    fields = []
    for col in range(1, max_col + 1):
        raw = ws.cell(best_row, col).value
        if raw in (None, ""):
            continue
        values = [ws.cell(row, col).value for row in range(best_row + 1, min(ws.max_row, best_row + 1000) + 1)]
        kind, samples, null_ratio, minimum, maximum = summarize(str(raw), values)
        fields.append(FieldProfile(get_column_letter(col), str(raw), kind, samples, null_ratio, ws.cell(best_row + 1, col).number_format if best_row < ws.max_row else "General", minimum, maximum))
    return [HeaderCandidate(best_row, best_row, round(max(0.0, min(best_score, 1.0)), 3), fields)] if fields else []
