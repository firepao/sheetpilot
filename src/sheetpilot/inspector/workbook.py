from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..errors import ErrorCode, SheetPilotError
from ..models import SheetProfile, WorkbookProfile
from ..workspace import sha256_file
from .headers import detect_headers
from .ooxml_risk import inspect_ooxml


def inspect_workbook(path: Path) -> WorkbookProfile:
    path = path.resolve()
    risks = inspect_ooxml(path)
    try:
        wb = load_workbook(path, read_only=False, data_only=False, keep_links=True)
        sheets = []
        for ws in wb.worksheets:
            formula_count = sum(1 for row in ws.iter_rows() for cell in row if cell.data_type == "f")
            used_range = f"A1:{get_column_letter(max(ws.max_column, 1))}{max(ws.max_row, 1)}"
            sheets.append(SheetProfile(ws.title, ws.sheet_state, used_range, detect_headers(ws), formula_count, len(ws._charts)))
        wb.close()
    except Exception as exc:
        raise SheetPilotError(ErrorCode.INPUT_INVALID, "无法打开工作簿", {"reason": str(exc)}) from exc
    return WorkbookProfile("1.0", {"path": str(path), "sha256": sha256_file(path), "size_bytes": path.stat().st_size, "format": path.suffix.lower().lstrip(".")}, sheets, risks, [])
