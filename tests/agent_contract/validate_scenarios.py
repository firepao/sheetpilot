#!/usr/bin/env python3
"""Validate the scenario registry and report non-runnable fixtures."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_FIELDS = ("scenario_key", "id", "user_prompt", "prompt_file", "data_file", "expected_outcome")


def validate(root: Path) -> list[dict[str, object]]:
    scenarios_dir = root / "scenarios"
    data_dir = root / "data"
    prompts_dir = root / "prompts"
    results: list[dict[str, object]] = []
    keys: set[str] = set()
    for path in sorted(scenarios_dir.glob("*.json")):
        item = json.loads(path.read_text(encoding="utf-8"))
        errors: list[str] = []
        key = item.get("scenario_key")
        if key != path.stem:
            errors.append("scenario_key_must_equal_filename_stem")
        if item.get("id") != key:
            errors.append("id_must_equal_scenario_key")
        if key in keys:
            errors.append("duplicate_scenario_key")
        keys.add(str(key))
        for field in REQUIRED_FIELDS:
            if not item.get(field):
                errors.append(f"missing_{field}")
        configured_data_file = str(item.get("data_file", ""))
        if configured_data_file != Path(configured_data_file).name:
            errors.append("data_file_must_be_a_filename")
        data_file = data_dir / Path(configured_data_file).name
        if not data_file.is_file():
            errors.append("missing_data_file")
        prompt_file = prompts_dir / str(item.get("prompt_file", ""))
        if not prompt_file.is_file():
            errors.append("missing_prompt_file")
        if configured_data_file and configured_data_file not in str(item.get("user_prompt", "")):
            errors.append("prompt_data_file_mismatch")
        if item.get("blocking_issues"):
            errors.append("blocking_issues")
        if item.get("evaluation_mode") not in {"deterministic_request", "semantic", "correct_stop"}:
            errors.append("invalid_evaluation_mode")
        if item.get("evaluation_mode") == "deterministic_request" and not item.get("task_request_template"):
            errors.append("missing_task_request_template")
        source = item.get("source")
        if data_file.is_file():
            if not isinstance(source, dict) or not source.get("sheet") or not source.get("header_row"):
                errors.append("missing_source_contract")
            else:
                workbook = load_workbook(data_file, read_only=True, data_only=True)
                try:
                    if source["sheet"] not in workbook.sheetnames:
                        errors.append("source_sheet_not_found")
                    else:
                        header = [cell.value for cell in workbook[source["sheet"]][int(source["header_row"])]]
                        if not header or any(value in (None, "") for value in header):
                            errors.append("invalid_header_row")
                        missing_fields = sorted(set(item.get("expected_source_fields", [])) - set(header))
                        if missing_fields:
                            errors.append(f"missing_source_fields:{','.join(missing_fields)}")
                finally:
                    workbook.close()
        declared_status = item.get("status")
        runnable = not errors
        status_matches = declared_status == ("active" if runnable else "draft")
        results.append({
            "scenario_key": key,
            "file": path.name,
            "status": declared_status,
            "runnable": runnable,
            "valid": status_matches,
            "errors": errors,
        })
    return results


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate SheetPilot agent-contract scenarios")
    parser.add_argument("--root", type=Path, default=Path(__file__).parent)
    args = parser.parse_args()
    results = validate(args.root)
    active = sum(item["status"] == "active" for item in results)
    valid = all(bool(item["valid"]) for item in results)
    print(json.dumps({"valid": valid, "total": len(results), "active": active, "draft": len(results) - active, "scenarios": results}, ensure_ascii=False, indent=2))
    return 0 if valid else 1


if __name__ == "__main__":
    raise SystemExit(main())
