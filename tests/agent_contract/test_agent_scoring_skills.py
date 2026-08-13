from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
PARSER = ROOT / "skills/sheetpilot-run-review/scripts/parse_agent_replay.py"
APPENDER = ROOT / "skills/sheetpilot-run-scorer/scripts/append_score_excel.py"


class AgentScoringSkillsTest(unittest.TestCase):
    def test_parser_emits_essential_and_compact_views(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); replay = root / "replay.json"
            start = 1_786_437_070_000
            replay.write_text(json.dumps({"messages": [
                {"role": "user", "content": "汇总销售额", "timestamp_ms": start},
                {"role": "agent", "content_type": "step", "step_id": "s1", "timestamp_ms": start + 1_000, "finishedTimestamp": start + 4_000},
                {"role": "agent", "content": "完成", "timestamp_ms": start + 5_000, "finishedTimestamp": start + 8_000},
            ], "steps": [{"id": "s1", "name": "shell", "command": "python wrapper task-types", "result": {"status": "ok"}, "timestamp": start + 1_000, "lastEventTimestamp": start + 4_000, "status": "success"}]}), encoding="utf-8")
            essential, compact = root / "essential.json", root / "compact.json"
            for mode, output in (("essential", essential), ("compact", compact)):
                result = subprocess.run([sys.executable, str(PARSER), str(replay), "--format", mode, "-o", str(output)], capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)
            compact_result = json.loads(compact.read_text(encoding="utf-8")); essential_result = json.loads(essential.read_text(encoding="utf-8"))
            self.assertEqual(compact_result["steps"][0]["duration_ms"], 3_000)
            self.assertEqual(compact_result["timing"]["total_duration_ms"], 8_000)
            self.assertEqual(compact_result["timing"]["end_ms"], start + 8_000)
            self.assertIn("+08:00", compact_result["timing"]["start_local"])
            self.assertEqual(essential_result["stats"]["step_count"], 1)
            self.assertEqual(essential_result["stats"]["step_time"]["p95_ms"], 3_000)
            self.assertEqual(essential_result["stats"]["agent_response_duration_ms"], 4_000)

    def test_excel_archive_recomputes_score_and_groups_versions(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); score_path, excel = root / "score.json", root / "scores.xlsx"
            score = {"scenario_id": "S1", "task_result": "PASS", "quality_score": 1, "diagnostic_score": 1, "prompt": "按部门汇总销售额", "platform": "Codex", "model": "test", "skill_version": "v1", "semantic_review": {"status": "accepted"}, "critical_violations": {}, "metrics": {"elapsed_ms": 8_000, "wall_clock_duration_ms": 8_000, "timing_source": "replay_timestamps", "timing_complete": True, "tool_duration_sum_ms": 3_000, "agent_response_duration_sum_ms": 4_000, "step_count": 1, "step_avg_ms": 3_000, "step_p95_ms": 3_000, "agent_operations": 3, "sheetpilot_cli_calls": 3, "filesystem_exploration": 0}, "score_breakdown": {"interface_compliance": 30, "execution_efficiency": 20, "recovery_behavior": 15, "evidence_and_report": 20, "output_usability": 15}}
            score_path.write_text(json.dumps(score, ensure_ascii=False), encoding="utf-8")
            first = subprocess.run([sys.executable, str(APPENDER), str(excel), "--score-json", str(score_path)], capture_output=True, text=True)
            self.assertEqual(first.returncode, 0, first.stderr)
            second = subprocess.run([sys.executable, str(APPENDER), str(excel), "--score-json", str(score_path), "--connect", "2"], capture_output=True, text=True)
            self.assertEqual(second.returncode, 0, second.stderr)
            workbook = load_workbook(excel); sheet = workbook["SheetPilotReviews"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, headers["质量分"]).value, 100)
            self.assertEqual(sheet.cell(2, headers["总用时"]).value, "0:08")
            self.assertEqual(sheet.cell(2, headers["总用时(ms)"]).value, 8_000)
            self.assertEqual(sheet.cell(2, headers["时间证据完整"]).value, "是")
            self.assertEqual(sheet.cell(2, headers["步骤P95耗时(ms)"]).value, 3_000)
            self.assertEqual(sheet.max_row, 3); workbook.close()

    def test_excel_archive_keeps_missing_timing_blank(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); score_path, excel = root / "score.json", root / "scores.xlsx"
            score = {"scenario_id": "S1", "task_result": "FAIL", "prompt": "缺少时间", "critical_violations": {}, "metrics": {}, "score_breakdown": {}}
            score_path.write_text(json.dumps(score, ensure_ascii=False), encoding="utf-8")
            result = subprocess.run([sys.executable, str(APPENDER), str(excel), "--score-json", str(score_path)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            workbook = load_workbook(excel); sheet = workbook["SheetPilotReviews"]; headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertIsNone(sheet.cell(2, headers["总用时(ms)"]).value)
            self.assertEqual(sheet.cell(2, headers["时间证据完整"]).value, "否"); workbook.close()


if __name__ == "__main__": unittest.main()
