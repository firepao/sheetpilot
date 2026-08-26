from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from tests.agent_contract.evaluation.evaluate_batch import evaluate_manifest


class BatchEvaluationTest(unittest.TestCase):
    def test_mixed_runs_are_isolated_and_scored(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            good = {"scenario": {"id": "S2_full", "expected_outcome": "success"}, "commands": ["task-run"], "task_status": {"state": "RUNTIME_PASS", "artifact_integrity": "MATCHED", "delivery_valid": True}, "oracle_result": {"passed": True, "input_unchanged": True}, "final_report": {"runtime_status": "RUNTIME_PASS"}, "task_request": {"filters": [], "dimensions": [], "metrics": [], "output": {"sort": []}, "acceptance": {"required_filters": [], "required_dimensions": [], "required_metrics": [], "required_sort": []}}}
            bad = {"scenario": {"id": "S3_full", "expected_outcome": "success"}, "commands": [], "task_request": {}}
            (root / "good.json").write_text(json.dumps(good), encoding="utf-8")
            (root / "bad.json").write_text(json.dumps(bad), encoding="utf-8")
            manifest = {"schema_version": "1.0", "runs": [{"scenario_key": "S2_full", "run_id": "a", "bundle_file": "good.json"}, {"scenario_key": "S3_full", "run_id": "b", "bundle_file": "bad.json"}]}
            (root / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            report = evaluate_manifest(root / "manifest.json", root / "out")
            self.assertEqual(report["counts"]["PASS"], 1)
            self.assertEqual(report["counts"]["ERROR"], 1)
            self.assertEqual(report["runs"][0]["score_policy"], "benchmark-v1")
            self.assertTrue((root / "out" / "semantic-queue.json").is_file())
            self.assertTrue((root / "out" / "scores.json").is_file())
            self.assertTrue((root / "out" / "S2_full" / "a" / "review-packet.json").is_file())
            self.assertTrue((root / "out" / "scores.xlsx").is_file())
            self.assertIsNone(report["runs"][1]["benchmark_score"])
            self.assertTrue((root / "out" / "S3_full" / "b" / "score.json").is_file())
            workbook = load_workbook(root / "out" / "scores.xlsx", read_only=True, data_only=True)
            try:
                self.assertEqual(workbook["BatchScores"].max_row, 3)
                self.assertEqual(workbook["BatchScores"]["E2"].value, "benchmark-v1")
            finally:
                workbook.close()

    def test_duplicate_key_and_run_id_is_configuration_error(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            bundle = {"scenario": {"id": "S2", "expected_outcome": "success"}}
            (root / "bundle.json").write_text(json.dumps(bundle), encoding="utf-8")
            manifest = {"schema_version": "1.0", "runs": [{"scenario_key": "S2", "run_id": "same", "bundle_file": "bundle.json"}, {"scenario_key": "S2", "run_id": "same", "bundle_file": "bundle.json"}]}
            (root / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            report = evaluate_manifest(root / "manifest.json", root / "out")
            self.assertEqual(report["status"], "SCENARIO_CONFIG_ERROR")

    def test_resume_merges_semantic_review(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            bundle = {"scenario": {"id": "S2", "expected_outcome": "success", "semantic_review_required": True}, "commands": ["task-run"], "task_status": {"state": "RUNTIME_PASS", "artifact_integrity": "MATCHED", "delivery_valid": True}, "oracle_result": {"passed": True, "input_unchanged": True}, "final_report": {"runtime_status": "RUNTIME_PASS"}, "task_request": {"filters": [], "dimensions": [], "metrics": [], "output": {"sort": []}, "acceptance": {"required_filters": [], "required_dimensions": [], "required_metrics": [], "required_sort": []}}, "semantic_review": {"status": "not_assessed"}}
            (root / "bundle.json").write_text(json.dumps(bundle), encoding="utf-8")
            (root / "manifest.json").write_text(json.dumps({"schema_version": "1.0", "runs": [{"scenario_key": "S2", "run_id": "r1", "bundle_file": "bundle.json"}]}), encoding="utf-8")
            (root / "semantic.json").write_text(json.dumps({"runs": [{"scenario_key": "S2", "run_id": "r1", "semantic_review": {"status": "accepted", "rationale": "evidence"}}]}), encoding="utf-8")
            report = evaluate_manifest(root / "manifest.json", root / "out", root / "semantic.json")
            self.assertEqual(report["runs"][0]["task_result"], "PASS")
            self.assertEqual(report["runs"][0]["score_status"], "FINAL")

    def test_replay_only_manifest_prepares_explicit_evidence(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            replay = {"run_bundle": {"scenario": {"id": "S2", "expected_outcome": "success"}, "commands": ["task-run"], "task_request": {"filters": [], "dimensions": [], "metrics": [], "output": {"sort": []}, "acceptance": {"required_filters": [], "required_dimensions": [], "required_metrics": [], "required_sort": []}}, "task_status": {"state": "RUNTIME_PASS", "artifact_integrity": "MATCHED", "delivery_valid": True}, "oracle_result": {"passed": True, "input_unchanged": True}, "final_report": {"runtime_status": "RUNTIME_PASS"}}}
            (root / "replay.json").write_text(json.dumps(replay), encoding="utf-8")
            (root / "manifest.json").write_text(json.dumps({"schema_version": "1.0", "runs": [{"scenario_key": "S2", "run_id": "r1", "replay_file": "replay.json"}]}), encoding="utf-8")
            report = evaluate_manifest(root / "manifest.json", root / "out")
            self.assertEqual(report["runs"][0]["task_result"], "PASS")
            self.assertTrue((root / "out" / "S2" / "r1" / "replay.raw.json").is_file())

    def test_scenario_registry_file_is_authoritative(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            bundle = {"scenario": {"id": "S2", "expected_outcome": "success"}, "commands": ["task-run"], "task_request": {"filters": [], "dimensions": [], "metrics": [], "output": {"sort": []}, "acceptance": {"required_filters": [], "required_dimensions": [], "required_metrics": [], "required_sort": []}}, "task_status": {"state": "RUNTIME_PASS", "artifact_integrity": "MATCHED", "delivery_valid": True}, "oracle_result": {"passed": True, "input_unchanged": True}, "final_report": {"runtime_status": "RUNTIME_PASS"}}
            registry = {"scenario_key": "S2_full", "id": "S2", "status": "active", "expected_outcome": "success"}
            (root / "bundle.json").write_text(json.dumps(bundle), encoding="utf-8")
            (root / "scenario.json").write_text(json.dumps(registry), encoding="utf-8")
            (root / "manifest.json").write_text(json.dumps({"schema_version": "1.0", "runs": [{"scenario_key": "S2_full", "run_id": "r1", "bundle_file": "bundle.json", "scenario_file": "scenario.json"}]}), encoding="utf-8")
            report = evaluate_manifest(root / "manifest.json", root / "out")
            self.assertEqual(report["runs"][0]["task_result"], "PASS")

    def test_incomplete_raw_replay_keeps_all_evidence_views_and_errors(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "replay.json").write_text(json.dumps({"messages": [], "steps": []}), encoding="utf-8")
            (root / "manifest.json").write_text(json.dumps({"schema_version": "1.0", "runs": [{"scenario_key": "S2", "run_id": "r1", "replay_file": "replay.json"}]}), encoding="utf-8")
            report = evaluate_manifest(root / "manifest.json", root / "out")
            self.assertEqual(report["runs"][0]["task_result"], "ERROR")
            run_dir = root / "out" / "S2" / "r1"
            self.assertEqual({p.name for p in run_dir.iterdir()}, {"replay.raw.json", "replay.essential.json", "replay.compact.json", "score.json"})

    def test_deterministic_cache_hits_and_invalidates_on_replay_change(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            replay = {"run_bundle": {"scenario": {"id": "S2", "expected_outcome": "success"}, "commands": ["task-run"], "task_request": {"filters": [], "dimensions": [], "metrics": [], "output": {"sort": []}, "acceptance": {"required_filters": [], "required_dimensions": [], "required_metrics": [], "required_sort": []}}, "task_status": {"state": "RUNTIME_PASS", "artifact_integrity": "MATCHED", "delivery_valid": True}, "oracle_result": {"passed": True, "input_unchanged": True}, "final_report": {"runtime_status": "RUNTIME_PASS"}}}
            (root / "replay.json").write_text(json.dumps(replay), encoding="utf-8")
            (root / "manifest.json").write_text(json.dumps({"schema_version": "1.0", "runs": [{"scenario_key": "S2", "run_id": "r1", "replay_file": "replay.json"}]}), encoding="utf-8")
            first = evaluate_manifest(root / "manifest.json", root / "out")
            second = evaluate_manifest(root / "manifest.json", root / "out")
            self.assertEqual(first["runs"][0].get("cache"), None)
            self.assertEqual(second["runs"][0].get("cache"), "hit")
            (root / "replay.json").write_text(json.dumps({**replay, "changed": True}), encoding="utf-8")
            third = evaluate_manifest(root / "manifest.json", root / "out")
            self.assertEqual(third["runs"][0].get("cache"), None)

    def test_cache_invalidates_when_semantic_results_change(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            replay = {"run_bundle": {"scenario": {"id": "S2", "expected_outcome": "success", "semantic_review_required": True}, "commands": ["task-run"], "task_request": {"filters": [], "dimensions": [], "metrics": [], "output": {"sort": []}, "acceptance": {"required_filters": [], "required_dimensions": [], "required_metrics": [], "required_sort": []}}, "task_status": {"state": "RUNTIME_PASS", "artifact_integrity": "MATCHED", "delivery_valid": True}, "oracle_result": {"passed": True, "input_unchanged": True}, "final_report": {"runtime_status": "RUNTIME_PASS"}, "semantic_review": {"status": "not_assessed"}}}
            (root / "replay.json").write_text(json.dumps(replay), encoding="utf-8")
            (root / "semantic.json").write_text(json.dumps({"runs": []}), encoding="utf-8")
            (root / "manifest.json").write_text(json.dumps({"schema_version": "1.0", "runs": [{"scenario_key": "S2", "run_id": "r1", "replay_file": "replay.json"}]}), encoding="utf-8")
            evaluate_manifest(root / "manifest.json", root / "out", root / "semantic.json")
            second = evaluate_manifest(root / "manifest.json", root / "out", root / "semantic.json")
            self.assertEqual(second["runs"][0].get("cache"), "hit")
            (root / "semantic.json").write_text(json.dumps({"runs": [{"scenario_key": "S2", "run_id": "r1", "semantic_review": {"status": "accepted"}}]}), encoding="utf-8")
            third = evaluate_manifest(root / "manifest.json", root / "out", root / "semantic.json")
            self.assertEqual(third["runs"][0].get("cache"), None)
            self.assertEqual(third["runs"][0]["task_result"], "PASS")

    def test_scenario_root_resolves_full_dataset_keys(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); scenarios = root / "scenarios"; scenarios.mkdir()
            scenario = {"scenario_key": "S2_full", "id": "S2", "status": "active", "expected_outcome": "success"}
            (scenarios / "S2_full.json").write_text(json.dumps(scenario), encoding="utf-8")
            bundle = {"scenario": {"id": "S2", "expected_outcome": "success"}, "commands": ["task-run"], "task_request": {"filters": [], "dimensions": [], "metrics": [], "output": {"sort": []}, "acceptance": {"required_filters": [], "required_dimensions": [], "required_metrics": [], "required_sort": []}}, "task_status": {"state": "RUNTIME_PASS", "artifact_integrity": "MATCHED", "delivery_valid": True}, "oracle_result": {"passed": True, "input_unchanged": True}, "final_report": {"runtime_status": "RUNTIME_PASS"}}
            (root / "bundle.json").write_text(json.dumps(bundle), encoding="utf-8")
            manifest = {"schema_version": "1.0", "scenario_root": "scenarios", "runs": [{"scenario_key": "S2_full", "run_id": "r1", "bundle_file": "bundle.json"}]}
            (root / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            report = evaluate_manifest(root / "manifest.json", root / "out")
            self.assertEqual(report["runs"][0]["task_result"], "PASS")

    def test_archive_failure_is_reported_without_losing_json_scores(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            bundle = {"scenario": {"id": "S2", "expected_outcome": "success"}, "commands": ["task-run"], "task_request": {"filters": [], "dimensions": [], "metrics": [], "output": {"sort": []}, "acceptance": {"required_filters": [], "required_dimensions": [], "required_metrics": [], "required_sort": []}}, "task_status": {"state": "RUNTIME_PASS", "artifact_integrity": "MATCHED", "delivery_valid": True}, "oracle_result": {"passed": True, "input_unchanged": True}, "final_report": {"runtime_status": "RUNTIME_PASS"}}
            (root / "bundle.json").write_text(json.dumps(bundle), encoding="utf-8")
            (root / "manifest.json").write_text(json.dumps({"schema_version": "1.0", "runs": [{"scenario_key": "S2", "run_id": "r1", "bundle_file": "bundle.json"}]}), encoding="utf-8")
            with patch("tests.agent_contract.evaluation.evaluate_batch._write_scores_excel", side_effect=OSError("archive unavailable")):
                report = evaluate_manifest(root / "manifest.json", root / "out")
            self.assertIn("archive_error", report)
            self.assertTrue((root / "out" / "scores.json").is_file())

    def test_scores_excel_appends_to_existing_history(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            bundle = {"scenario": {"id": "S2", "expected_outcome": "success"}, "commands": ["task-run"], "task_request": {"filters": [], "dimensions": [], "metrics": [], "output": {"sort": []}, "acceptance": {"required_filters": [], "required_dimensions": [], "required_metrics": [], "required_sort": []}}, "task_status": {"state": "RUNTIME_PASS", "artifact_integrity": "MATCHED", "delivery_valid": True}, "oracle_result": {"passed": True, "input_unchanged": True}, "final_report": {"runtime_status": "RUNTIME_PASS"}}
            (root / "bundle.json").write_text(json.dumps(bundle), encoding="utf-8")
            manifest = {"schema_version": "1.0", "runs": [{"scenario_key": "S2", "run_id": "r1", "bundle_file": "bundle.json"}]}
            (root / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            evaluate_manifest(root / "manifest.json", root / "out")
            evaluate_manifest(root / "manifest.json", root / "out")
            workbook = load_workbook(root / "out" / "scores.xlsx", read_only=True, data_only=True)
            try:
                self.assertEqual(workbook["BatchScores"].max_row, 3)
            finally:
                workbook.close()

    def test_workers_process_runs_in_manifest_order(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); runs = []
            for index in range(3):
                bundle = {"scenario": {"id": f"S{index}", "expected_outcome": "success"}, "commands": ["task-run"], "task_request": {"filters": [], "dimensions": [], "metrics": [], "output": {"sort": []}, "acceptance": {"required_filters": [], "required_dimensions": [], "required_metrics": [], "required_sort": []}}, "task_status": {"state": "RUNTIME_PASS", "artifact_integrity": "MATCHED", "delivery_valid": True}, "oracle_result": {"passed": True, "input_unchanged": True}, "final_report": {"runtime_status": "RUNTIME_PASS"}}
                path = root / f"b{index}.json"; path.write_text(json.dumps(bundle), encoding="utf-8")
                runs.append({"scenario_key": f"S{index}", "run_id": "r1", "bundle_file": str(path)})
            (root / "manifest.json").write_text(json.dumps({"schema_version": "1.0", "workers": 3, "runs": runs}), encoding="utf-8")
            report = evaluate_manifest(root / "manifest.json", root / "out")
            self.assertEqual([item["scenario_key"] for item in report["runs"]], ["S0", "S1", "S2"])
            self.assertEqual(report["counts"]["PASS"], 3)

    def test_invalid_workers_is_configuration_error(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            manifest = {"schema_version": "1.0", "workers": 0, "runs": [{"scenario_key": "S2", "run_id": "r1", "bundle_file": "missing.json"}]}
            (root / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            report = evaluate_manifest(root / "manifest.json", root / "out")
            self.assertEqual(report["status"], "SCENARIO_CONFIG_ERROR")
            self.assertIn("workers_must_be_integer_1_to_64", report["errors"])

    def test_manifest_build_errors_block_partial_dataset_evaluation(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            manifest = {"schema_version": "1.0", "build": {"errors": ["missing_evidence_directory:M3"]}, "runs": [{"scenario_key": "S2", "run_id": "r1", "bundle_file": "missing.json"}]}
            (root / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            report = evaluate_manifest(root / "manifest.json", root / "out")
            self.assertEqual(report["status"], "SCENARIO_CONFIG_ERROR")
            self.assertIn("missing_evidence_directory:M3", report["errors"])
