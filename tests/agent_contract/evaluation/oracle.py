from __future__ import annotations

import hashlib
import math
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _compare(value: Any, operator: str, expected: Any) -> bool:
    if operator == "eq": return value == expected
    if operator == "ne": return value != expected
    if operator == "gt": return value is not None and value > expected
    if operator == "gte": return value is not None and value >= expected
    if operator == "lt": return value is not None and value < expected
    if operator == "lte": return value is not None and value <= expected
    if operator == "in": return value in expected
    raise ValueError(f"Unsupported filter operator: {operator}")


def _read_source(request: dict[str, Any]) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(request["input_file"], read_only=True, data_only=True)
    try:
        sheet = workbook[request["source"]["sheet"]]
        header_row = request["source"].get("header_row", 1)
        headers = [cell.value for cell in sheet[header_row]]
        if any(value in (None, "") for value in headers):
            raise ValueError("Oracle requires a contiguous, non-empty header row")
        if len(headers) != len(set(headers)):
            raise ValueError("Oracle cannot evaluate unresolved duplicate headers")
        rows = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=header_row + 1, values_only=True)]
        return headers, rows
    finally:
        workbook.close()


def _sheet_values(workbook: Any, sheet_name: str) -> list[tuple[Any, ...]]:
    return list(workbook[sheet_name].iter_rows(values_only=True))


def _same(actual: Any, expected: Any, tolerance: float) -> bool:
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        return math.isclose(float(actual), float(expected), rel_tol=tolerance, abs_tol=tolerance)
    return actual == expected


def evaluate_summarize_table(
    request: dict[str, Any],
    *,
    input_sha256_before: str | None = None,
    tolerance: float = 1e-6,
) -> dict[str, Any]:
    """Independently recompute a summarize_table request without importing SheetPilot."""
    input_path, output_path = Path(request["input_file"]), Path(request["output_file"])
    checks: list[dict[str, Any]] = []
    current_input_hash = sha256_file(input_path)
    input_unchanged = input_sha256_before is None or current_input_hash == input_sha256_before
    checks.append({"name": "input_unchanged", "passed": input_unchanged})
    if not output_path.is_file():
        checks.append({"name": "output_exists", "passed": False})
        return {"schema_version": "1.0", "passed": False, "input_unchanged": input_unchanged, "checks": checks}

    _, source_rows = _read_source(request)
    filtered = [row for row in source_rows if all(_compare(row[item["field"]], item["operator"], item["value"]) for item in request["filters"])]
    dimensions = request["dimensions"]
    grouped: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
    for row in filtered:
        grouped[tuple(row[item["field"]] for item in dimensions)].append(row)

    expected_rows: list[list[Any]] = []
    for key, rows in grouped.items():
        result = list(key)
        for metric in request["metrics"]:
            function = metric["function"]
            values = [row[metric["field"]] for row in rows if metric.get("field") and row[metric["field"]] not in (None, "")]
            if function == "sum": result.append(sum(values))
            elif function == "average": result.append(sum(values) / len(values) if values else None)
            elif function == "count" and metric["mode"] == "rows": result.append(len(rows))
            elif function == "count" and metric["mode"] == "non_empty": result.append(len(values))
            else: raise ValueError(f"Unsupported metric: {metric}")
        expected_rows.append(result)

    ids = [item["id"] for item in dimensions + request["metrics"]]
    for sort in reversed(request["output"].get("sort", [])):
        index = ids.index(sort["by"])
        expected_rows.sort(key=lambda row: (row[index] is None, row[index]), reverse=sort["direction"] == "desc")
    expected_headers = [item["output_name"] for item in dimensions + request["metrics"]]

    source_workbook = load_workbook(input_path, read_only=True, data_only=False)
    workbook = load_workbook(output_path, read_only=True, data_only=False)
    try:
        target = request["output"]["sheet"]
        expected_sheets = source_workbook.sheetnames + [target]
        sheets_pass = workbook.sheetnames == expected_sheets
        source_sheets_pass = all(_sheet_values(source_workbook, name) == _sheet_values(workbook, name) for name in source_workbook.sheetnames)
        checks.extend([
            {"name": "sheet_set_and_order_match", "passed": sheets_pass, "expected": expected_sheets, "actual": workbook.sheetnames},
            {"name": "source_sheets_unchanged", "passed": source_sheets_pass},
        ])
        sheet_exists = target in workbook.sheetnames
        checks.append({"name": "target_sheet_exists", "passed": sheet_exists})
        if not sheet_exists:
            return {"schema_version": "1.0", "passed": False, "input_unchanged": input_unchanged, "checks": checks}
        sheet = workbook[target]
        actual = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
        source_workbook.close()

    actual_headers = list(actual[0][:len(expected_headers)]) if actual else []
    actual_rows = [list(row[:len(expected_headers)]) for row in actual[1:] if any(value is not None for value in row)]
    headers_pass = actual_headers == expected_headers
    shape_pass = len(actual_rows) == len(expected_rows)
    values_pass = shape_pass and all(all(_same(a, e, tolerance) for a, e in zip(actual_row, expected_row)) for actual_row, expected_row in zip(actual_rows, expected_rows))
    checks.extend([
        {"name": "headers_match", "passed": headers_pass, "expected": expected_headers, "actual": actual_headers},
        {"name": "row_count_match", "passed": shape_pass, "expected": len(expected_rows), "actual": len(actual_rows)},
        {"name": "group_values_and_order_match", "passed": values_pass},
    ])
    passed = all(item["passed"] for item in checks)
    return {"schema_version": "1.0", "passed": passed, "input_unchanged": input_unchanged, "filtered_row_count": len(filtered), "group_count": len(expected_rows), "checks": checks}
