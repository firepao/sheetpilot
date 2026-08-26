#!/usr/bin/env python3
"""Run the deterministic evaluator over one or more prepared run bundles.

The manifest is deliberately explicit: scenario_key is the join key and a
bundle path is the evidence boundary. This keeps batch orchestration separate
from replay acquisition and semantic review.
"""
from __future__ import annotations

import argparse
import ast
import hashlib
import json
import os
import subprocess
import sys
import tempfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .evaluator import evaluate_run
from .oracle import evaluate_summarize_table


def _read(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"JSON object required: {path}")
    return value


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _cache_key(item: dict[str, Any], base: Path, semantic_results_path: Path | None) -> str:
    values: dict[str, Any] = {"policy": "benchmark-v1", "scenario_key": item.get("scenario_key"), "run_id": item.get("run_id")}
    file_values = [("bundle", item.get("bundle_file") or item.get("review_dir")), ("replay", item.get("replay_file")), ("scenario", item.get("scenario_file")), ("semantic", str(semantic_results_path) if semantic_results_path else None)]
    for label, value in file_values:
        if not value:
            values[label] = None
            continue
        path = _resolve(base, value)
        if path.is_dir(): path = path / "run-bundle.json"
        values[label] = {"path": str(path), "sha256": _sha256(path) if path.is_file() else None}
    request_path = None
    bundle_value = item.get("bundle_file") or item.get("review_dir")
    if bundle_value:
        candidate = _resolve(base, bundle_value); candidate = candidate / "run-bundle.json" if candidate.is_dir() else candidate
        if candidate.is_file():
            try:
                request_path = _read(candidate).get("task_request", {}).get("output_file")
            except (OSError, ValueError, json.JSONDecodeError):
                request_path = None
    if request_path:
        output = Path(request_path)
        values["output"] = {"path": str(output), "sha256": _sha256(output) if output.is_file() else None}
    payload = json.dumps(values, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _resolve(base: Path, value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else (base / path).resolve()


def _find_first(value: Any, key: str) -> Any:
    """Find an exact evidence field without interpreting arbitrary text."""
    if isinstance(value, dict):
        if key in value:
            return value[key]
        for child in value.values():
            found = _find_first(child, key)
            if found is not None:
                return found
    elif isinstance(value, list):
        for child in value:
            found = _find_first(child, key)
            if found is not None:
                return found
    elif isinstance(value, str):
        text = value.strip()
        if text[:1] in {"{", "["}:
            try:
                return _find_first(json.loads(text), key)
            except json.JSONDecodeError:
                pass
    return None


def _find_task_request(value: Any) -> dict[str, Any] | None:
    if isinstance(value, dict):
        if value.get("task_type") and value.get("input_file") and value.get("source"):
            return value
        for child in value.values():
            found = _find_task_request(child)
            if found is not None:
                return found
    elif isinstance(value, list):
        for child in value:
            found = _find_task_request(child)
            if found is not None:
                return found
    elif isinstance(value, str) and value.lstrip()[:1] in {"{", "["}:
        try:
            return _find_task_request(json.loads(value))
        except json.JSONDecodeError:
            return None
    return None


def _find_task_status(value: Any) -> dict[str, Any] | None:
    if isinstance(value, dict):
        if value.get("state") and ("artifact_integrity" in value or "delivery_valid" in value):
            return value
        for child in value.values():
            found = _find_task_status(child)
            if found is not None:
                return found
    elif isinstance(value, list):
        for child in value:
            found = _find_task_status(child)
            if found is not None:
                return found
    elif isinstance(value, str) and value.lstrip()[:1] in {"{", "["}:
        try:
            return _find_task_status(json.loads(value))
        except json.JSONDecodeError:
            return None
    return None


def _json_object(text: Any) -> dict[str, Any] | None:
    if not isinstance(text, str):
        return None
    try:
        value = json.loads(text.strip())
    except (json.JSONDecodeError, TypeError):
        return None
    return value if isinstance(value, dict) else None


def _snapshot_messages(raw: dict[str, Any]) -> list[dict[str, Any]]:
    snapshot: Any = raw
    data = raw.get("data")
    if isinstance(data, dict) and isinstance(data.get("snapshot_data"), str):
        snapshot = _json_object(data["snapshot_data"])
    if not isinstance(snapshot, dict):
        return []
    flattened: list[dict[str, Any]] = []
    for message in snapshot.get("messages", []):
        if not isinstance(message, dict):
            continue
        exploration = message.get("exploration")
        children = exploration.get("messages") if isinstance(exploration, dict) else None
        flattened.extend(child for child in children if isinstance(child, dict)) if isinstance(children, list) else flattened.append(message)
    return flattened


def _scenario_segment(messages: list[dict[str, Any]], scenario: dict[str, Any]) -> list[dict[str, Any]]:
    segments: list[list[dict[str, Any]]] = []
    current: list[dict[str, Any]] = []
    for message in messages:
        if message.get("role") == "user" and isinstance(message.get("content"), str):
            if current:
                segments.append(current)
            current = [message]
        elif current:
            current.append(message)
    if current:
        segments.append(current)
    data_file = str(scenario.get("data_file") or "").lower()
    prompt = "".join(str(scenario.get("user_prompt") or "").lower().split())
    matches = []
    for segment in segments:
        user_text = str(segment[0].get("content") or "")
        normalized = "".join(user_text.lower().split())
        if (data_file and data_file in user_text.lower()) or (prompt and (prompt in normalized or normalized in prompt)):
            matches.append(segment)
    if not matches:
        discovered = sorted({Path(token.replace("\\\\", "\\")).name for segment in segments for token in str(segment[0].get("content") or "").split() if token.lower().endswith((".xlsx", ".xlsm"))})
        raise ValueError(f"scenario_not_found_in_replay:{scenario.get('scenario_key')}; discovered={','.join(discovered) or 'none'}")
    if len(matches) > 1:
        raise ValueError(f"scenario_ambiguous_in_replay:{scenario.get('scenario_key')}")
    return matches[0]


def _balanced_literal(text: str, start: int) -> str | None:
    opening = text[start:start + 1]
    closing = {"{": "}", "[": "]"}.get(opening)
    if not closing:
        return None
    depth, quote, escaped = 0, None, False
    for index in range(start, len(text)):
        char = text[index]
        if quote:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == quote:
                quote = None
            continue
        if char in {"'", '"'}:
            quote = char
        elif char == opening:
            depth += 1
        elif char == closing:
            depth -= 1
            if depth == 0:
                return text[start:index + 1]
    return None


def _request_from_command(command: str) -> dict[str, Any] | None:
    marker_index = command.find("json.dump(")
    if marker_index >= 0:
        start = command.find("{", marker_index)
        literal = _balanced_literal(command, start) if start >= 0 else None
        if literal:
            try:
                value = ast.literal_eval(literal)
                if isinstance(value, dict) and value.get("task_type"):
                    return value
            except (SyntaxError, ValueError):
                pass
    task_index = command.find('"task_type"')
    if task_index >= 0:
        start = command.rfind("{", 0, task_index)
        value = _json_object(_balanced_literal(command, start)) if start >= 0 else None
        if value and value.get("task_type"):
            return value
    return None


def _evidence_from_segment(segment: list[dict[str, Any]]) -> dict[str, Any]:
    request = status = run_result = None
    commands: list[str] = []
    final_text = ""
    for message in segment:
        step = message.get("step")
        if isinstance(step, dict):
            inputs = step.get("inputs", {}).get("values", {})
            outputs = step.get("outputs", {}).get("values", {})
            command = inputs.get("command") if isinstance(inputs, dict) else None
            stdout = outputs.get("stdout") if isinstance(outputs, dict) else None
            if isinstance(command, str):
                commands.append(command)
                request = _request_from_command(command) or request
            payload = _json_object(stdout)
            if payload and payload.get("state") and "task-status" in str(command):
                status = payload
            elif payload and payload.get("status") and "task-run" in str(command):
                run_result = payload
        elif message.get("role") in {"robot", "agent", "assistant"} and isinstance(message.get("content"), str) and message.get("content").strip():
            final_text = message["content"]
    if request and run_result and run_result.get("output_file"):
        request["output_file"] = run_result["output_file"]
    report: dict[str, Any] | None = None
    if final_text:
        evidence = status or run_result or {}
        report = {"text": final_text, "task_id": evidence.get("task_id"), "runtime_status": evidence.get("state") or evidence.get("status"), "artifact_integrity": evidence.get("artifact_integrity"), "delivery_valid": evidence.get("delivery_valid"), "output_file": (run_result or {}).get("output_file")}
        marker_index = final_text.rfind('"semantic_assessment"')
        if marker_index >= 0:
            start = final_text.rfind("{", 0, marker_index)
            semantic = _json_object(_balanced_literal(final_text, start)) if start >= 0 else None
            if semantic:
                report.update(semantic)
    return {"task_request": request, "task_status": status, "final_report": report, "commands": commands}


def _parse_replay(replay_path: Path, run_dir: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    """Use the repository's review parser and retain both evidence views."""
    parser = Path(__file__).parents[2] / ".." / "skills" / "sheetpilot-run-review" / "scripts" / "parse_agent_replay.py"
    parser = parser.resolve()
    if not parser.is_file():
        raise ValueError(f"replay_parser_not_found:{parser}")
    run_dir.mkdir(parents=True, exist_ok=True)
    views = {}
    for view in ("essential", "compact"):
        target = run_dir / f"replay.{view}.json"
        completed = subprocess.run(
            [sys.executable, str(parser), str(replay_path), "--format", view, "-o", str(target)],
            capture_output=True, text=True, check=False,
        )
        if completed.returncode != 0:
            raise ValueError(f"replay_parse_failed:{view}:{completed.stderr.strip() or completed.stdout.strip()}")
        views[view] = _read(target)
    return views["essential"], views["compact"]


def _bundle_from_replay(replay_path: Path, item: dict[str, Any]) -> dict[str, Any]:
    preparation_dir = item.get("_preparation_dir")
    if not isinstance(preparation_dir, str):
        raise ValueError("internal_missing_preparation_dir")
    original = _read(replay_path)
    essential, compact = _parse_replay(replay_path, Path(preparation_dir))
    raw = original
    scenario = item.get("scenario") if isinstance(item.get("scenario"), dict) else {"scenario_key": item.get("scenario_key"), "id": item.get("scenario_key"), "expected_outcome": "success"}
    direct = raw.get("run_bundle") if isinstance(raw, dict) else None
    if isinstance(direct, dict):
        bundle = direct
    else:
        extracted: dict[str, Any] = {}
        messages = _snapshot_messages(raw)
        if messages and scenario.get("data_file"):
            extracted = _evidence_from_segment(_scenario_segment(messages, scenario))
        bundle = {
            "scenario": scenario,
            "task_request": extracted.get("task_request") or _find_first(raw, "task_request") or _find_task_request(raw),
            "task_status": extracted.get("task_status") or _find_first(raw, "task_status") or _find_task_status(raw),
            "oracle_result": _find_first(raw, "oracle_result"),
            "semantic_review": _find_first(raw, "semantic_review") or {"status": "not_assessed"},
            "final_report": extracted.get("final_report") or _find_first(raw, "final_report"),
            "commands": extracted.get("commands") or _find_first(raw, "commands") or [],
            "replay_stats": compact.get("timing") or _find_first(raw, "replay_stats"),
        }
        if isinstance(bundle.get("task_request"), dict) and not bundle.get("oracle_result"):
            bundle["oracle_result"] = evaluate_summarize_table(bundle["task_request"])
    bundle["replay_essential"] = essential
    bundle["replay_compact"] = compact
    if not isinstance(bundle.get("scenario"), dict):
        bundle["scenario"] = {"scenario_key": item.get("scenario_key"), "id": item.get("scenario_key")}
    bundle["scenario"].setdefault("scenario_key", item.get("scenario_key"))
    bundle["scenario"].setdefault("id", item.get("scenario_key"))
    if not bundle.get("task_request") or not bundle.get("task_status") or not bundle.get("final_report"):
        raise ValueError("replay_evidence_incomplete: task_request/task_status/final_report required")
    return bundle


def _write_scores_excel(path: Path, entries: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(path) if path.is_file() else Workbook()
    sheet = workbook["BatchScores"] if "BatchScores" in workbook.sheetnames else workbook.create_sheet("BatchScores")
    columns = ["scenario_key", "run_id", "task_result", "benchmark_score", "score_policy", "score_status", "failed_gates", "elapsed_ms", "error", "score_file"]
    existing_header = [sheet.cell(1, index).value for index in range(1, max(sheet.max_column, len(columns)) + 1)] if sheet.max_row else []
    if not any(existing_header):
        if sheet.max_row:
            sheet.delete_rows(1, sheet.max_row)
        sheet.append(columns)
    elif existing_header[:len(columns)] != columns:
        raise ValueError("existing scores.xlsx has incompatible BatchScores header")
    for entry in entries:
        sheet.append([
            entry.get("scenario_key"), entry.get("run_id"), entry.get("task_result"), entry.get("benchmark_score"), entry.get("score_policy"),
            entry.get("score_status"), ", ".join(entry.get("failed_gates", [])),
            (entry.get("metrics") or {}).get("elapsed_ms"), entry.get("error"), entry.get("score_file"),
        ])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill; cell.font = Font(color="FFFFFF", bold=True)
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    widths = [32, 20, 22, 18, 18, 18, 42, 16, 52, 70]
    for index, width in enumerate(widths, 1): sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(prefix=f".{path.stem}-", suffix=path.suffix, dir=path.parent, delete=False) as stream:
            temporary = Path(stream.name)
        workbook.save(temporary)
        workbook.close()
        os.replace(temporary, path)
    finally:
        workbook.close()
        if temporary and temporary.exists(): temporary.unlink()


def _write_review_packet(path: Path, bundle: dict[str, Any]) -> None:
    required = ("task_request", "task_status", "final_report", "oracle_result", "commands")
    packet = {
        "schema_version": "1.0",
        "scenario": bundle.get("scenario", {}),
        "prompt": bundle.get("prompt") or bundle.get("scenario", {}).get("user_prompt"),
        "task_request": bundle.get("task_request"),
        "task_status": bundle.get("task_status"),
        "final_report": bundle.get("final_report"),
        "oracle_result": bundle.get("oracle_result"),
        "commands": bundle.get("commands", []),
        "replay_stats": bundle.get("replay_stats"),
        "missing_evidence": [field for field in required if not bundle.get(field)],
    }
    path.write_text(json.dumps(packet, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _validate_manifest(manifest: dict[str, Any], base: Path) -> list[str]:
    errors: list[str] = []
    if manifest.get("schema_version") != "1.0":
        errors.append("schema_version_must_be_1.0")
    runs = manifest.get("runs")
    if not isinstance(runs, list) or not runs:
        return errors + ["runs_must_be_non_empty_list"]
    workers = manifest.get("workers", 1)
    if isinstance(workers, bool) or not isinstance(workers, int) or not 1 <= workers <= 64:
        errors.append("workers_must_be_integer_1_to_64")
    seen: set[tuple[str, str]] = set()
    scenario_root_value = manifest.get("scenario_root")
    scenario_root = _resolve(base, scenario_root_value) if isinstance(scenario_root_value, str) else None
    if scenario_root_value and not scenario_root.is_dir():
        errors.append(f"missing_scenario_root:{scenario_root_value}")
    for index, item in enumerate(runs):
        if not isinstance(item, dict):
            errors.append(f"runs[{index}]_must_be_object"); continue
        key, run_id = item.get("scenario_key"), item.get("run_id")
        if not isinstance(key, str) or not key:
            errors.append(f"runs[{index}]_missing_scenario_key")
        if not isinstance(run_id, str) or not run_id:
            errors.append(f"runs[{index}]_missing_run_id")
        pair = (str(key), str(run_id))
        if pair in seen:
            errors.append(f"duplicate_scenario_key_and_run_id:{key}:{run_id}")
        seen.add(pair)
        bundle = item.get("bundle_file") or item.get("review_dir")
        replay = item.get("replay_file")
        if not isinstance(bundle, str) and not isinstance(replay, str):
            errors.append(f"runs[{index}]_missing_bundle_file")
        elif isinstance(bundle, str) and not _resolve(base, bundle).exists():
            errors.append(f"runs[{index}]_missing_bundle:{bundle}")
        if isinstance(replay, str) and not _resolve(base, replay).is_file():
            errors.append(f"runs[{index}]_missing_replay:{replay}")
        scenario_file = item.get("scenario_file")
        if not scenario_file and scenario_root:
            scenario_file = str(scenario_root / f"{key}.json")
        if scenario_file:
            scenario_path = _resolve(base, scenario_file)
            if not scenario_path.is_file():
                errors.append(f"runs[{index}]_missing_scenario:{scenario_file}")
            else:
                try:
                    scenario = _read(scenario_path)
                    if scenario.get("scenario_key") != key:
                        errors.append(f"scenario_key_mismatch:{scenario.get('scenario_key')}:{key}")
                    if scenario.get("status") != "active":
                        errors.append(f"scenario_not_active:{key}")
                except (OSError, ValueError, json.JSONDecodeError) as exc:
                    errors.append(f"invalid_scenario:{key}:{exc}")
    return errors


def _score(bundle: dict[str, Any]) -> dict[str, Any]:
    result = evaluate_run(bundle)
    gates = result.get("gates", {})
    failed = [name for name, passed in gates.items() if not passed]
    benchmark = result.get("benchmark", {})
    if not bundle.get("commands"):
        result.update({"task_result": "ERROR", "benchmark_score": None, "score_status": "UNAVAILABLE"})
    elif bundle.get("scenario", {}).get("semantic_review_required") and (bundle.get("semantic_review") or {}).get("status") == "not_assessed":
        result.update({"task_result": "NEEDS_SEMANTIC_REVIEW", "benchmark_score": benchmark.get("score"), "score_policy": benchmark.get("policy"), "score_status": "PROVISIONAL", "failed_gates": failed})
    else:
        result.update({
            "benchmark_score": benchmark.get("score"),
            "score_policy": benchmark.get("policy"),
            "score_status": "FINAL",
            "failed_gates": failed,
        })
    return result


def _process_item(item: dict[str, Any], base: Path, output_dir: Path, cache_dir: Path, semantic_results: dict[tuple[str, str], dict[str, Any]], semantic_results_path: Path | None, use_cache: bool) -> tuple[dict[str, Any], dict[str, Any] | None]:
    key, run_id = item.get("scenario_key"), item.get("run_id")
    run_dir = output_dir / str(key) / str(run_id)
    run_dir.mkdir(parents=True, exist_ok=True)
    item = {**item, "_preparation_dir": str(run_dir)}
    cache_file = cache_dir / f"{_cache_key(item, base, semantic_results_path)}.json" if use_cache else None
    if cache_file and cache_file.is_file():
        cached = _read(cache_file); result = cached["result"]
        (run_dir / "score.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        packet_path = run_dir / "review-packet.json"
        bundle_path_for_packet = run_dir / "run-bundle.json"
        _write_review_packet(packet_path, _read(bundle_path_for_packet) if bundle_path_for_packet.is_file() else {})
        entry = {"scenario_key": key, "run_id": run_id, **result, "cache": "hit", "score_file": str(run_dir / "score.json")}
        queued = {"scenario_key": key, "run_id": run_id, "packet_file": str(packet_path)} if result.get("task_result") == "NEEDS_SEMANTIC_REVIEW" or result.get("score_status") == "PROVISIONAL" else None
        return entry, queued
    try:
        bundle_value, replay_value = item.get("bundle_file") or item.get("review_dir"), item.get("replay_file")
        source = _resolve(base, bundle_value) if bundle_value else None
        bundle_path = source / "run-bundle.json" if source and source.is_dir() else source
        entry_hash = None
        if replay_value:
            replay_path = _resolve(base, replay_value)
            if not replay_path.is_file(): raise ValueError(f"missing replay_file: {replay_value}")
            (run_dir / "replay.raw.json").write_bytes(replay_path.read_bytes()); entry_hash = _sha256(replay_path)
        if bundle_path is not None:
            bundle = _read(bundle_path)
            if replay_value:
                essential, compact = _parse_replay(_resolve(base, replay_value), run_dir)
                bundle["replay_essential"], bundle["replay_compact"] = essential, compact
        else:
            scenario = _read(_resolve(base, item["scenario_file"])) if item.get("scenario_file") else None
            bundle = _bundle_from_replay(_resolve(base, replay_value), {**item, "scenario": scenario})
        if item.get("scenario_file"): bundle["scenario"] = _read(_resolve(base, item["scenario_file"]))
        if bundle.get("scenario", {}).get("scenario_key", bundle.get("scenario", {}).get("id")) != key:
            if item.get("scenario_file"): raise ValueError("scenario_key does not match scenario registry")
            bundle.setdefault("scenario", {})["scenario_key"] = key
        semantic = semantic_results.get((key, run_id))
        if semantic: bundle["semantic_review"] = semantic.get("semantic_review", semantic)
        result = _score(bundle)
        (run_dir / "run-bundle.json").write_text(json.dumps(bundle, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        packet_path = run_dir / "review-packet.json"; _write_review_packet(packet_path, bundle)
        (run_dir / "score.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        if cache_file: cache_file.write_text(json.dumps({"schema_version": "1.0", "result": result}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        entry = {"scenario_key": key, "run_id": run_id, **result, "replay_sha256": entry_hash, "score_file": str(run_dir / "score.json")}
        queued = {"scenario_key": key, "run_id": run_id, "packet_file": str(packet_path)} if result["task_result"] == "NEEDS_SEMANTIC_REVIEW" or result.get("score_status") == "PROVISIONAL" else None
        return entry, queued
    except Exception as exc:
        error_result = {"schema_version": "1.0", "scenario_key": key, "run_id": run_id, "task_result": "ERROR", "benchmark_score": None, "score_status": "UNAVAILABLE", "error": str(exc)}
        error_path = run_dir / "score.json"; error_path.write_text(json.dumps(error_result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return {**error_result, "score_file": str(error_path)}, None


def evaluate_manifest(manifest_path: Path, output_dir: Path, semantic_results_path: Path | None = None, use_cache: bool = True) -> dict[str, Any]:
    manifest = _read(manifest_path)
    base = manifest_path.parent.resolve()
    semantic_results: dict[tuple[str, str], dict[str, Any]] = {}
    if semantic_results_path:
        payload = _read(semantic_results_path)
        for item in payload.get("runs", []):
            if isinstance(item, dict) and item.get("scenario_key") and item.get("run_id"):
                semantic_results[(item["scenario_key"], item["run_id"])] = item
    errors = _validate_manifest(manifest, base)
    build_info = manifest.get("build")
    if isinstance(build_info, dict):
        errors.extend(str(error) for error in build_info.get("errors", []) if error)
    output_dir.mkdir(parents=True, exist_ok=True)
    if errors:
        report = {"schema_version": "1.0", "status": "SCENARIO_CONFIG_ERROR", "errors": errors, "runs": []}
        (output_dir / "batch-summary.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return report
    entries: list[dict[str, Any]] = []
    semantic_queue: list[dict[str, Any]] = []
    scenario_root_value = manifest.get("scenario_root")
    scenario_root = _resolve(base, scenario_root_value) if isinstance(scenario_root_value, str) else None
    cache_dir = output_dir / ".cache"
    if use_cache: cache_dir.mkdir(parents=True, exist_ok=True)
    prepared_items = [({**item, "scenario_file": str(scenario_root / f"{item['scenario_key']}.json")} if not item.get("scenario_file") and scenario_root else item) for item in manifest["runs"]]
    workers = int(manifest.get("workers", 1) or 1)
    if workers > 1 and len(prepared_items) > 1:
        with ThreadPoolExecutor(max_workers=workers) as executor:
            processed = list(executor.map(lambda item: _process_item(item, base, output_dir, cache_dir, semantic_results, semantic_results_path, use_cache), prepared_items))
    else:
        processed = [_process_item(item, base, output_dir, cache_dir, semantic_results, semantic_results_path, use_cache) for item in prepared_items]
    for entry, queued in processed:
        entries.append(entry)
        if queued:
            semantic_queue.append(queued)
    summary = {
        "schema_version": "1.0", "status": "completed", "created_at": datetime.now(timezone.utc).isoformat(),
        "counts": {status: sum(entry.get("task_result") == status for entry in entries) for status in ("PASS", "FAIL", "NEEDS_SEMANTIC_REVIEW", "ERROR")},
        "pass_rate": sum(e.get("task_result") == "PASS" for e in entries) / len(entries) if entries else 0,
        "error_rate": sum(e.get("task_result") == "ERROR" for e in entries) / len(entries) if entries else 0,
        "average_benchmark_score": (sum(e["benchmark_score"] for e in entries if isinstance(e.get("benchmark_score"), (int, float))) / sum(isinstance(e.get("benchmark_score"), (int, float)) for e in entries)) if any(isinstance(e.get("benchmark_score"), (int, float)) for e in entries) else None,
        "runs": entries,
    }
    (output_dir / "batch-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (output_dir / "scores.json").write_text(json.dumps({"schema_version": "1.0", "score_policy": "benchmark-v1", "runs": entries}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (output_dir / "semantic-queue.json").write_text(json.dumps({"schema_version": "1.0", "runs": semantic_queue}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    scores_value = manifest.get("scores_file")
    scores_path = _resolve(base, scores_value) if isinstance(scores_value, str) else output_dir / "scores.xlsx"
    try:
        _write_scores_excel(scores_path, entries)
        summary["scores_file"] = str(scores_path)
        (output_dir / "batch-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        (output_dir / "scores.json").write_text(json.dumps({"schema_version": "1.0", "score_policy": "benchmark-v1", "scores_file": str(scores_path), "runs": entries}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    except Exception as exc:
        # Excel commonly locks the destination workbook. Preserve the current
        # batch in a new file instead of leaving the user with stale scores.xlsx.
        fallback = scores_path.with_name(f"{scores_path.stem}-latest{scores_path.suffix}")
        try:
            _write_scores_excel(fallback, entries)
            summary["scores_file"] = str(fallback)
            summary["archive_error"] = str(exc)
            summary["archive_note"] = f"scores.xlsx 可能正在被 Excel 占用。请关闭该文件；本批次已自动写入: {fallback}"
        except Exception as fallback_exc:
            summary["archive_error"] = str(exc)
            summary["fallback_archive_error"] = str(fallback_exc)
        (output_dir / "batch-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        (output_dir / "scores.json").write_text(json.dumps({"schema_version": "1.0", "score_policy": "benchmark-v1", "scores_file": summary.get("scores_file"), "runs": entries}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Batch-evaluate prepared SheetPilot run bundles")
    parser.add_argument("manifest", type=Path)
    parser.add_argument("-o", "--output-dir", type=Path, required=True)
    parser.add_argument("--resume", type=Path, help="Merge semantic-results.json before rescoring")
    parser.add_argument("--no-cache", action="store_true", help="Disable deterministic result cache")
    args = parser.parse_args(argv)
    report = evaluate_manifest(args.manifest, args.output_dir, args.resume, use_cache=not args.no_cache)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 2 if report.get("status") == "SCENARIO_CONFIG_ERROR" or report.get("counts", {}).get("ERROR") or report.get("archive_error") else 0


if __name__ == "__main__":
    raise SystemExit(main())
