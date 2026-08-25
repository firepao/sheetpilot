#!/usr/bin/env python3
"""
批量评估Agent执行结果

基于本地replay文件，批量运行Oracle验证和质量评分。
不需要LLM，纯Python计算，支持并行处理。

用法:
    # 评估单个场景
    python batch_evaluate.py replays/M3/

    # 评估多个场景
    python batch_evaluate.py replays/M3/ replays/M4/ replays/M5/

    # 评估整个目录（自动发现子目录）
    python batch_evaluate.py replays/ --auto-discover

    # 并行处理
    python batch_evaluate.py replays/ --auto-discover --parallel 4

    # 生成Excel报告
    python batch_evaluate.py replays/ --auto-discover -o results/scores.xlsx
"""
from __future__ import annotations

import argparse
import json
import sys
import traceback
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

# 导入现有的评估逻辑
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

try:
    from tests.agent_contract.evaluation.evaluator import evaluate_run, inspect_commands
    from tests.agent_contract.evaluation.oracle import evaluate_summarize_table
    from skills.sheetpilot_run_review.scripts.parse_agent_replay import _normalize, _essential
except ImportError as e:
    print(f"Error: Cannot import required modules: {e}")
    print("Please run from SheetPilot repository root")
    sys.exit(1)


def extract_commands_from_replay(replay: dict) -> list[str]:
    """从replay中提取所有命令"""
    commands = []
    messages = replay.get("messages", [])

    for msg in messages:
        if msg.get("content_type") == "tool_use":
            tool_name = msg.get("tool_name", "")
            if "bash" in tool_name.lower() or "command" in tool_name.lower():
                # 提取命令
                input_data = msg.get("input", {})
                command = input_data.get("command", "")
                if command:
                    commands.append(command)

    return commands


def extract_task_request_from_replay(replay: dict) -> dict | None:
    """从replay中提取task request"""
    messages = replay.get("messages", [])

    for msg in messages:
        if msg.get("content_type") == "tool_use":
            tool_name = msg.get("tool_name", "")
            input_data = msg.get("input", {})

            # 查找task-run调用
            if "task-run" in tool_name or "task_run" in str(input_data):
                # 尝试多种方式提取request
                request = input_data.get("request")
                if request:
                    if isinstance(request, str):
                        try:
                            return json.loads(request)
                        except json.JSONDecodeError:
                            pass
                    elif isinstance(request, dict):
                        return request

                # 尝试从命令中提取JSON
                command = input_data.get("command", "")
                if "task-run" in command:
                    # 提取 --request 后面的JSON
                    import re
                    match = re.search(r'--request["\s]+({.*?})\s*(?:--|$)', command, re.DOTALL)
                    if match:
                        try:
                            return json.loads(match.group(1))
                        except json.JSONDecodeError:
                            pass

    return None


def extract_task_status_from_replay(replay: dict) -> dict | None:
    """从replay中提取task status"""
    messages = replay.get("messages", [])

    for msg in messages:
        if msg.get("content_type") == "tool_result":
            result_data = msg.get("result", {})

            # 查找task-status结果
            if isinstance(result_data, dict):
                if "state" in result_data or "status" in result_data:
                    return result_data

            # 尝试解析字符串结果
            if isinstance(result_data, str):
                try:
                    parsed = json.loads(result_data)
                    if "state" in parsed or "status" in parsed:
                        return parsed
                except json.JSONDecodeError:
                    pass

    return None


def build_bundle_from_replay(scenario_dir: Path, scenario_id: str) -> dict:
    """从replay构建evaluation bundle"""
    replay_file = scenario_dir / "replay.json"

    if not replay_file.exists():
        raise FileNotFoundError(f"replay.json not found in {scenario_dir}")

    # 读取原始replay
    raw_replay = json.loads(replay_file.read_text(encoding="utf-8"))

    # 解析replay（使用现有的parse逻辑）
    compact = _normalize(raw_replay, max_string_length=0)
    essential = _essential(compact)

    # 提取关键信息
    commands = extract_commands_from_replay(raw_replay)
    task_request = extract_task_request_from_replay(raw_replay)
    task_status = extract_task_status_from_replay(raw_replay)

    # 安全地加载场景定义（防止路径遍历）
    scenario = {}
    # 消毒scenario_id：只允许字母数字和连字符
    safe_scenario_id = re.sub(r'[^\w-]', '_', scenario_id)
    if re.match(r'^[A-Za-z0-9_-]+$', safe_scenario_id):
        scenarios_dir = Path(__file__).parent / "scenarios"
        scenario_json = scenarios_dir / f"{safe_scenario_id}.json"

        # 验证解析后的路径仍在scenarios目录内
        try:
            scenario_json_resolved = scenario_json.resolve()
            scenarios_dir_resolved = scenarios_dir.resolve()

            # Python 3.9+ 使用 is_relative_to
            if hasattr(scenario_json_resolved, 'is_relative_to'):
                is_safe = scenario_json_resolved.is_relative_to(scenarios_dir_resolved)
            else:
                # Python 3.8 兼容
                try:
                    scenario_json_resolved.relative_to(scenarios_dir_resolved)
                    is_safe = True
                except ValueError:
                    is_safe = False

            if is_safe and scenario_json_resolved.exists():
                scenario = json.loads(scenario_json_resolved.read_text(encoding="utf-8"))
        except (ValueError, OSError):
            # 路径解析失败，使用空scenario
            pass

    # 构建bundle
    bundle = {
        "scenario_id": scenario_id,
        "scenario": scenario,
        "replay": essential,
        "commands": commands,
        "task_request": task_request,
        "task_status": task_status,
    }

    # 运行Oracle（如果有输出文件）
    output_file = scenario_dir / "output.xlsx"
    if output_file.exists() and task_request:
        try:
            oracle_result = evaluate_summarize_table(
                task_request,
                str(output_file),
                scenario
            )
            bundle["oracle"] = oracle_result
        except Exception as e:
            bundle["oracle_error"] = str(e)

    return bundle


def evaluate_scenario(scenario_dir: Path) -> dict:
    """评估单个场景（纯Python，无LLM调用）"""
    scenario_id = scenario_dir.name

    try:
        print(f"Processing {scenario_id}...", end=" ", flush=True)

        # 构建bundle
        bundle = build_bundle_from_replay(scenario_dir, scenario_id)

        # 运行评估（evaluate_run是纯Python逻辑）
        score = evaluate_run(bundle)

        # 提取关键指标
        result = {
            "scenario_id": scenario_id,
            "task_result": score.get("task_result"),
            "quality_score": score.get("quality_score"),
            "diagnostic_score": score.get("diagnostic_score"),
            "oracle_pass": score.get("oracle", {}).get("passed"),
            "critical_violations": score.get("critical_violations", {}),
            "gates": score.get("gates", {}),
            "pass": score.get("task_result") == "PASS",
            "evaluated_at": datetime.now(timezone.utc).isoformat(),
        }

        status = "✓ PASS" if result["pass"] else "✗ FAIL"
        print(f"{status}  score={result.get('quality_score', 'N/A')}")

        return result

    except Exception as e:
        print(f"✗ ERROR: {str(e)[:50]}")
        return {
            "scenario_id": scenario_id,
            "error": str(e),
            "traceback": traceback.format_exc(),
            "pass": False,
        }


def discover_scenario_dirs(root_dir: Path) -> list[Path]:
    """自动发现包含replay.json的场景目录"""
    scenario_dirs = []

    for path in root_dir.iterdir():
        if path.is_dir():
            replay_file = path / "replay.json"
            if replay_file.exists():
                scenario_dirs.append(path)

    return sorted(scenario_dirs, key=lambda p: p.name)


def save_summary_json(results: list[dict], output_file: Path):
    """保存JSON格式的汇总报告"""
    passed = sum(1 for r in results if r.get("pass"))
    failed = len(results) - passed
    errors = sum(1 for r in results if "error" in r)

    summary = {
        "evaluated_at": datetime.now(timezone.utc).isoformat(),
        "total": len(results),
        "passed": passed,
        "failed": failed,
        "errors": errors,
        "pass_rate": f"{passed/len(results)*100:.1f}%" if results else "0%",
        "results": sorted(results, key=lambda x: x["scenario_id"]),
    }

    output_file.parent.mkdir(parents=True, exist_ok=True)
    output_file.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )


def save_summary_excel(results: list[dict], output_file: Path):
    """保存Excel格式的汇总报告"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = Workbook()
        ws = wb.active
        ws.title = "Evaluation Summary"

        # 表头
        headers = [
            "场景ID", "任务结果", "质量分数", "诊断分数",
            "Oracle通过", "关键违规", "通过"
        ]
        ws.append(headers)

        # 样式
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # 数据行
        for r in sorted(results, key=lambda x: x["scenario_id"]):
            critical_count = sum(r.get("critical_violations", {}).values())

            row = [
                r["scenario_id"],
                r.get("task_result", "N/A"),
                r.get("quality_score"),
                r.get("diagnostic_score"),
                "是" if r.get("oracle_pass") else "否",
                critical_count,
                "✓" if r.get("pass") else "✗",
            ]
            ws.append(row)

            # 为失败的行着色
            if not r.get("pass"):
                for cell in ws[ws.max_row]:
                    cell.fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")

        # 调整列宽
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        output_file.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_file)

    except ImportError:
        print("Warning: openpyxl not installed, skipping Excel output")
        print("Install with: pip install openpyxl")


def main() -> int:
    parser = argparse.ArgumentParser(
        description="批量评估Agent执行结果",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__.split("用法:")[1]
    )

    parser.add_argument(
        "scenario_dirs",
        nargs="+",
        type=Path,
        help="场景目录路径（包含replay.json）"
    )
    parser.add_argument(
        "--auto-discover",
        action="store_true",
        help="自动发现子目录中的场景"
    )
    parser.add_argument(
        "--parallel",
        type=int,
        default=1,
        help="并行处理的worker数量，默认1"
    )
    parser.add_argument(
        "-o", "--output",
        type=Path,
        help="输出文件路径（.json或.xlsx）"
    )

    args = parser.parse_args()

    # 收集所有场景目录
    all_scenario_dirs = []

    for input_dir in args.scenario_dirs:
        if not input_dir.exists():
            print(f"Warning: {input_dir} does not exist, skipping")
            continue

        if args.auto_discover:
            # 自动发现子目录
            discovered = discover_scenario_dirs(input_dir)
            all_scenario_dirs.extend(discovered)
        else:
            # 直接使用指定目录
            if (input_dir / "replay.json").exists():
                all_scenario_dirs.append(input_dir)
            else:
                print(f"Warning: {input_dir} does not contain replay.json, skipping")

    if not all_scenario_dirs:
        print("Error: No valid scenario directories found")
        return 1

    print(f"Found {len(all_scenario_dirs)} scenarios to evaluate")
    print(f"Processing with {args.parallel} worker(s)...\n")

    # 评估所有场景
    results = []

    if args.parallel == 1:
        # 串行处理
        for scenario_dir in all_scenario_dirs:
            result = evaluate_scenario(scenario_dir)
            results.append(result)
    else:
        # 并行处理
        with ProcessPoolExecutor(max_workers=args.parallel) as executor:
            future_to_dir = {
                executor.submit(evaluate_scenario, d): d
                for d in all_scenario_dirs
            }

            for future in as_completed(future_to_dir):
                result = future.result()
                results.append(result)

    # 统计结果
    passed = sum(1 for r in results if r.get("pass"))
    failed = len(results) - passed
    errors = sum(1 for r in results if "error" in r)

    # 打印汇总
    print("\n" + "=" * 60)
    print("Evaluation Summary:")
    print(f"  Total:   {len(results)}")
    print(f"  PASS:    {passed}")
    print(f"  FAIL:    {failed}")
    print(f"  ERROR:   {errors}")
    print(f"  Rate:    {passed/len(results)*100:.1f}%" if results else "  Rate:    0%")
    print("=" * 60)

    # 详细列表
    print("\nDetailed Results:")
    for r in sorted(results, key=lambda x: x["scenario_id"]):
        status = "✓" if r.get("pass") else "✗"
        score = r.get("quality_score", "N/A")
        print(f"{status} {r['scenario_id']:8s}  score={score}")

    # 保存报告
    if args.output:
        if args.output.suffix == ".xlsx":
            save_summary_excel(results, args.output)
            print(f"\n✓ Excel report saved to: {args.output}")
        else:
            # 默认JSON格式
            output_json = args.output.with_suffix(".json")
            save_summary_json(results, output_json)
            print(f"\n✓ JSON report saved to: {output_json}")
    else:
        # 默认保存JSON到当前目录
        default_output = Path("evaluation_summary.json")
        save_summary_json(results, default_output)
        print(f"\n✓ JSON report saved to: {default_output}")

    return 0 if failed == 0 and errors == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
