from __future__ import annotations
import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from sheetpilot.task_api.inventory import MAX_INVENTORY_BYTES, MAX_INVENTORY_ENTRIES, build_inventory, cap_entries, read_headers

class InventoryTest(unittest.TestCase):
    def make(self, d, rows):
        p=Path(d)/"source.xlsx"; w=Workbook(); s=w.active; s.title="明细"
        for row in rows: s.append(row)
        w.save(p); w.close(); return p
    def test_read_headers_masks_and_orders(self):
        with tempfile.TemporaryDirectory() as d:
            e=read_headers(self.make(d, [["净销售额","手机号","证件号","邮箱"],[100,"13800138000","110101199001011234","zhangsan@example.com"]]), "h", "明细", 1)
            self.assertEqual([x["header"] for x in e],["净销售额","手机号","证件号","邮箱"]); self.assertEqual(e[1]["sample_values"],["138****8000"]); self.assertEqual(e[2]["sample_values"],["110101********1234"]); self.assertEqual(e[3]["sample_values"],["zhangsan***@example.com"])
    def test_cap_entries(self):
        entries=[{"id":str(i),"header":"x"*1000} for i in range(MAX_INVENTORY_ENTRIES+1)]; capped,truncated=cap_entries(entries); self.assertTrue(truncated); self.assertLessEqual(len(capped),MAX_INVENTORY_ENTRIES); self.assertLessEqual(len(json.dumps(capped,ensure_ascii=False).encode()),MAX_INVENTORY_BYTES)
    def test_ambiguous_and_data_rows(self):
        with tempfile.TemporaryDirectory() as d:
            p=self.make(d, [["城市","金额"],["地区","金额"],["北京",10]])
            self.assertEqual([x["header_row"] for x in build_inventory(p,"h")["header_candidates"]],[1,2])
            p=self.make(d, [["城市","金额"],["北京",10],["上海",30]])
            self.assertEqual(build_inventory(p,"h")["header_candidates"],[])
    def test_multi_sheet_budget(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/"wide.xlsx"; w=Workbook(); w.remove(w.active)
            for n in ("甲","乙","丙"):
                s=w.create_sheet(n); s.append(["列"+"甲"*600 for _ in range(40)]); s.append([1]*40)
            w.save(p); w.close(); self.assertEqual(build_inventory(p,"h",header_row=1)["error_code"],"INVENTORY_TOO_LARGE")

if __name__ == "__main__": unittest.main()
