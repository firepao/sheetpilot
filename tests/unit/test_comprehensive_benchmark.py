import json
import unittest
from pathlib import Path

from openpyxl import load_workbook

from sheetpilot.capabilities import DEFAULT_REGISTRY


ROOT = Path(__file__).resolve().parents[1] / "agent_contract"


class ComprehensiveBenchmarkTest(unittest.TestCase):
    def test_cases_have_real_workbooks_and_registered_coverage(self):
        benchmark = json.loads((ROOT / "comprehensive_benchmark_v2.json").read_text(encoding="utf-8"))
        registered = {definition.name for definition in DEFAULT_REGISTRY.definitions()}
        self.assertEqual(len(benchmark["cases"]), 15)
        for case in benchmark["cases"]:
            path = ROOT.parents[1] / case["input_file"]
            self.assertTrue(path.is_file(), case["id"])
            prompt_path = ROOT / "prompts" / case["prompt_file"]
            self.assertTrue(prompt_path.is_file(), case["id"])
            self.assertIn(Path(case["input_file"]).name, prompt_path.read_text(encoding="utf-8"))
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                self.assertTrue(set(case["sheets"]) <= set(workbook.sheetnames), case["id"])
            finally:
                workbook.close()
            self.assertTrue(set(case["coverage"]) <= registered, case["id"])
            self.assertGreaterEqual(len(case["coverage"]), 3)


if __name__ == "__main__":
    unittest.main()
