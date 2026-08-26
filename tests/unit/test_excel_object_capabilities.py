import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from sheetpilot.atomic.context import ExecutionContext
from sheetpilot.capabilities import DEFAULT_REGISTRY
from sheetpilot.engines import OpenPyxlEngine


class ExcelObjectCapabilityTest(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook(); self.ws = self.wb.active; self.ws.title = "结果"
        self.ws.append(["类别", "金额"]); self.ws.append(["A", 10]); self.ws.append(["B", 20])
        self.ctx = ExecutionContext(OpenPyxlEngine(self.wb), {})

    def call(self, op, params):
        definition = DEFAULT_REGISTRY.require(op)
        return definition.handler(self.ctx, params, {})

    def test_formula_dependency_and_error_scan(self):
        self.call("formula.write", {"sheet":"结果", "cell":"C2", "formula":"=B2*2"})
        self.assertIn("B2", self.call("formula.inspect_dependencies", {"sheet":"结果", "cell":"C2"})["references"])
        self.ws["D2"] = "#VALUE!"
        self.assertIn("D2", self.call("formula.find_errors", {"sheet":"结果"})["errors"])

    def test_objects_and_chart(self):
        self.call("excel_table.create", {"sheet":"结果", "name":"ResultTable", "range":"A1:B3"})
        chart = self.call("chart.create", {"sheet":"结果", "range":"A1:B3", "chart_type":"bar", "anchor":"E2"})
        self.call("chart.add_series", {"sheet":"结果", "chart_index":chart["chart_index"], "range":"B1:B3"})
        self.assertEqual(self.call("chart.inspect", {"sheet":"结果", "chart_index":chart["chart_index"]})["series_count"], 2)
        self.call("comment.create", {"sheet":"结果", "cell":"A1", "text":"说明"})
        self.assertEqual(self.ws["A1"].comment.text, "说明")
        self.call("named_range.create", {"name":"ResultAmount", "sheet":"结果", "range":"$B$2:$B$3"})
        self.assertIsNotNone(self.wb.defined_names["ResultAmount"])


if __name__ == "__main__": unittest.main()
