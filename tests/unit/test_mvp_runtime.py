from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from sheetpilot.engines import OpenPyxlEngine
from sheetpilot.errors import ErrorCode, SheetPilotError
from sheetpilot.mvp import execute, manifest, run_plan, validate_plan, validate_run
from sheetpilot.runtime import analyze_transform


class MvpRuntimeTest(unittest.TestCase):
    def test_manifest_exposes_parameter_contracts(self):
        value = manifest(); atoms = {item["name"]: item for item in value["atoms"]}
        self.assertEqual(set(atoms), {"read_table","filter_rows","derive_column","aggregate","sort_rows","create_sheet","write_table","save_workbook"})
        self.assertEqual(atoms["aggregate"]["supported_functions"], ["sum","count","average"])
        molecule = next(item for item in value["molecules"] if item["name"] == "summarize_by_dimension")
        self.assertEqual([step["op"] for step in molecule["steps"]], ["filter_rows","aggregate","sort_rows"])

    def test_molecule_aggregates_multiple_metrics_and_sorts_numbers(self):
        wb=Workbook(); ws=wb.active; ws.title="数据"; ws.append(["区域","订单","营收","状态"])
        ws.append(["北","A",10,"有效"]); ws.append(["北","B",20,"有效"]); ws.append(["南","C",9,"有效"]); ws.append(["南",None,99,"无效"])
        result=execute(OpenPyxlEngine(wb), [{"id":"rows","op":"read_table","sheet":"数据","columns":{"区域":"A","订单":"B","营收":"C","状态":"D"}}, {"id":"summary","op":"summarize_by_dimension","input":"rows","where":{"all":[{"field":"状态","op":"eq","value":"有效"}]},"group_by":["区域"],"metrics":[{"field":"营收","function":"sum","as":"销售收入"},{"function":"count","mode":"rows","as":"订单数"},{"field":"订单","function":"count","mode":"non_empty","as":"非空订单"},{"field":"营收","function":"average","as":"均值"}],"sort":[{"field":"销售收入","direction":"desc"}]}])
        self.assertEqual(result["summary"].rows[0], {"区域":"北","销售收入":30,"订单数":2,"非空订单":2,"均值":15})

    def test_unknown_aggregate_is_plan_invalid(self):
        plan={"schema_version":"1.0","input_file":"in.xlsx","output_file":"out.xlsx","requirements":{},"steps":[{"id":"x","op":"aggregate","input":"missing","group_by":[],"metrics":[{"field":"x","function":"median","as":"m"}]}]}
        with self.assertRaises(SheetPilotError) as caught: validate_plan(plan)
        self.assertEqual(caught.exception.code, ErrorCode.PLAN_INVALID)

    def test_dynamic_file_access_is_rejected(self):
        with self.assertRaises(SheetPilotError) as caught: analyze_transform("def transform(table, params):\n    return open('x').read()\n")
        self.assertEqual(caught.exception.code, ErrorCode.DYNAMIC_TRANSFORM_REJECTED)

    def test_run_validate_publish_and_stale_evidence(self):
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary); source=root/"source.xlsx"; output=root/"output.xlsx"; run_dir=root/"run"
            wb=Workbook(); ws=wb.active; ws.title="清洗明细"; ws.append(["城市","订单号","销售额","是否退货","清洗状态"]); ws.append(["北京","1",10,"否","有效"]); ws.append(["上海","2",30,"否","有效"]); ws.append(["北京","3",99,"是","有效"]); wb.save(source); wb.close()
            plan={"schema_version":"1.0","input_file":str(source),"output_file":str(output),"steps":[{"id":"source","op":"read_table","sheet":"清洗明细","columns":{"城市":"A","订单号":"B","销售额":"C","是否退货":"D","清洗状态":"E"}},{"id":"summary","op":"summarize_by_dimension","input":"source","where":{"all":[{"field":"清洗状态","op":"eq","value":"有效"},{"field":"是否退货","op":"eq","value":"否"}]},"group_by":["城市"],"metrics":[{"field":"销售额","function":"sum","as":"销售收入"},{"function":"count","mode":"rows","as":"订单数量"},{"field":"销售额","function":"average","as":"平均订单金额"}],"sort":[{"field":"销售收入","direction":"desc"}]},{"id":"sheet","op":"create_sheet","sheet":"城市经营汇总"},{"id":"write","op":"write_table","input":"summary","sheet":"城市经营汇总","anchor":"A1"}],"requirements":{"required_sheets":["城市经营汇总"],"required_columns":{"城市经营汇总":["城市","销售收入","订单数量","平均订单金额"]},"checks":[{"type":"aggregate_reconciliation","source_sheet":"清洗明细","source_field":"销售额","target_sheet":"城市经营汇总","target_field":"销售收入","where":{"all":[{"field":"清洗状态","value":"有效"},{"field":"是否退货","value":"否"}]}},{"type":"sort_order","sheet":"城市经营汇总","field":"销售收入","direction":"desc"}]}}
            before=source.read_bytes(); ready=run_plan(plan,run_dir); self.assertEqual(ready["status"],"READY_FOR_VALIDATION"); self.assertFalse(output.exists())
            result=validate_run(run_dir); self.assertEqual(result["status"],"PASS"); self.assertEqual(source.read_bytes(),before)
            workbook=load_workbook(output,data_only=True); self.assertEqual(workbook["城市经营汇总"]["B2"].value,30); workbook.close()
            workbook=load_workbook(run_dir/"temporary_output.xlsx"); workbook["城市经营汇总"]["B2"]=999; workbook.save(run_dir/"temporary_output.xlsx"); workbook.close()
            with self.assertRaises(SheetPilotError) as caught: validate_run(run_dir)
            self.assertEqual(caught.exception.code, ErrorCode.VALIDATION_FAILED)

    def test_modified_published_output_invalidates_evidence(self):
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary); source=root/"source.xlsx"; output=root/"output.xlsx"; run_dir=root/"run"
            wb=Workbook(); ws=wb.active; ws.title="数据"; ws.append(["城市","金额"]); ws.append(["北京",10]); wb.save(source); wb.close()
            plan={"schema_version":"1.0","input_file":str(source),"output_file":str(output),"steps":[{"id":"source","op":"read_table","sheet":"数据","columns":{"城市":"A","金额":"B"}},{"id":"sheet","op":"create_sheet","sheet":"结果"},{"id":"write","op":"write_table","input":"source","sheet":"结果"}],"requirements":{"required_sheets":["结果"],"required_columns":{"结果":["城市","金额"]},"checks":[]}}
            run_plan(plan,run_dir); validate_run(run_dir)
            workbook=load_workbook(output); workbook["结果"]["B2"]=999; workbook.save(output); workbook.close()
            with self.assertRaises(SheetPilotError) as caught: validate_run(run_dir)
            self.assertEqual(caught.exception.code,ErrorCode.VALIDATION_FAILED)

    def test_real_dirty_orders_file_end_to_end(self):
        source=Path(__file__).resolve().parents[1]/"fixtures"/"legacy-mvp"/"U3_builtin_dirty_orders_report.xlsx"
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary); output=root/"result.xlsx"; run_dir=root/"run"
            plan={"schema_version":"1.0","input_file":str(source),"output_file":str(output),"steps":[{"id":"source","op":"read_table","sheet":"清洗明细","columns":{"城市":"D","销售额":"G","是否退货":"J","清洗状态":"L"}},{"id":"summary","op":"summarize_by_dimension","input":"source","where":{"all":[{"field":"清洗状态","op":"eq","value":"有效"},{"field":"是否退货","op":"eq","value":"否"}]},"group_by":["城市"],"metrics":[{"field":"销售额","function":"sum","as":"销售收入"},{"function":"count","mode":"rows","as":"订单数量"},{"field":"销售额","function":"average","as":"平均订单金额"}],"sort":[{"field":"销售收入","direction":"desc"}]},{"id":"sheet","op":"create_sheet","sheet":"城市经营汇总"},{"id":"write","op":"write_table","input":"summary","sheet":"城市经营汇总"}],"requirements":{"required_sheets":["城市经营汇总"],"required_columns":{"城市经营汇总":["城市","销售收入","订单数量","平均订单金额"]},"checks":[{"type":"aggregate_reconciliation","source_sheet":"清洗明细","source_field":"销售额","target_sheet":"城市经营汇总","target_field":"销售收入","where":{"all":[{"field":"清洗状态","value":"有效"},{"field":"是否退货","value":"否"}]}},{"type":"sort_order","sheet":"城市经营汇总","field":"销售收入","direction":"desc"}]}}
            before=source.read_bytes(); run_plan(plan,run_dir); result=validate_run(run_dir)
            self.assertEqual(result["status"],"PASS"); self.assertEqual(source.read_bytes(),before)
            workbook=load_workbook(output,data_only=True); rows=list(workbook["城市经营汇总"].iter_rows(min_row=2,values_only=True)); workbook.close()
            self.assertTrue(rows); self.assertTrue(all(row[2] > 0 and row[3] is not None for row in rows)); self.assertEqual([row[1] for row in rows],sorted((row[1] for row in rows),reverse=True))


if __name__ == "__main__": unittest.main()
