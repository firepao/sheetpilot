from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from sheetpilot.errors import ErrorCode, SheetPilotError
from sheetpilot.inspector import inspect_workbook
from sheetpilot.models import HighLevelPlan, SemanticTask
from sheetpilot.planning.compiler import compile_plan
from sheetpilot.planning.policy import evaluate_policy
from sheetpilot.planning.ranges import parse_range
from sheetpilot.recipes import OperatingSummaryRecipe
from sheetpilot.executor import execute_plan
from sheetpilot.validators import validate_and_publish
from sheetpilot.workspace import sha256_file
from sheetpilot.capabilities import DEFAULT_REGISTRY
from sheetpilot.workbook.tables import TableData, filter_rows


class MvpTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name)
        self.source = self.root / "orders.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "订单明细"
        ws.append(["下单日期", "销售大区", "含税金额", "手机号"])
        ws.append([__import__("datetime").date(2026, 1, 1), "华北", 100.0, "13812345678"])
        ws.append([__import__("datetime").date(2026, 2, 1), "华东", -20.0, "13912345678"])
        ws.append([__import__("datetime").date(2026, 3, 1), None, 50.0, None]); wb.save(self.source); wb.close()

    def tearDown(self): self.temp.cleanup()

    def task(self):
        return SemanticTask.from_dict({"schema_version":"1.0","task_id":"t1","user_intent":"按大区汇总收入","input_file":str(self.source),"output_file":str(self.root / "result.xlsx"),"source":{"sheet":"订单明细","header_rows":[1]},"concepts":{"date":{"field":"下单日期","column":"A","confidence":1,"evidence":[],"alternatives":[]},"region":{"field":"销售大区","column":"B","confidence":1,"evidence":[],"alternatives":[]},"revenue":{"field":"含税金额","column":"C","confidence":1,"evidence":[],"alternatives":[]}},"business_definition":{"current_period":["2026-01-01","2026-12-31"]},"requested_output":{"target_sheet":"经营汇总","dimensions":["region"],"metrics":["revenue"],"chart":"bar"},"ambiguities":[],"confirmations":[]})

    def test_inspector_masks_sensitive_samples(self):
        profile = inspect_workbook(self.source)
        fields = profile.sheets[0].header_candidates[0].fields
        phone = next(field for field in fields if field.header == "手机号")
        self.assertNotIn("13812345678", phone.samples); self.assertEqual(profile.input["sha256"], sha256_file(self.source))

    def test_ranges(self):
        value = parse_range("'经营 汇总'!$A$1:C10")
        self.assertEqual(value.area, 30); self.assertTrue(value.intersects(parse_range("'经营 汇总'!C10:D12")))
        with self.assertRaises(SheetPilotError): parse_range("Sheet1!A:A")

    def test_policy_requires_confirmation(self):
        task = self.task(); data = task.to_dict(); data["ambiguities"] = [{"id":"tax","severity":"key","question":"含税？"}]; task = SemanticTask.from_dict(data)
        profile = inspect_workbook(self.source); plan = compile_plan(task, OperatingSummaryRecipe().compile(task), profile)
        self.assertEqual(evaluate_policy(plan, profile, task).decision, "CONFIRM")

    def test_end_to_end_and_input_unchanged(self):
        before = sha256_file(self.source); task = self.task(); profile = inspect_workbook(self.source)
        plan = compile_plan(task, OperatingSummaryRecipe().compile(task), profile); run_dir = self.root / "run"
        temporary, _ = execute_plan(plan, run_dir); result = validate_and_publish(plan, temporary, run_dir)
        self.assertEqual(result.status, "PASS"); self.assertEqual(before, sha256_file(self.source)); self.assertTrue(Path(result.output_file).exists())
        wb = load_workbook(result.output_file, data_only=True); ws = wb["经营汇总"]
        self.assertEqual(sum(ws.cell(r, 2).value for r in range(4, 7)), 130); self.assertEqual(len(ws._charts), 1); wb.close()

    def test_compiler_rejects_wrong_mapping(self):
        task = self.task(); data = task.to_dict(); data["concepts"]["revenue"]["column"] = "D"; task = SemanticTask.from_dict(data)
        with self.assertRaises(SheetPilotError) as caught: compile_plan(task, OperatingSummaryRecipe().compile(task), inspect_workbook(self.source))
        self.assertEqual(caught.exception.code, ErrorCode.PLAN_INVALID)

    def test_registry_is_executable_source_of_truth(self):
        definitions = {item.name: item for item in DEFAULT_REGISTRY.definitions()}
        capabilities = {"read_table", "filter_rows", "select_columns", "derive_column", "aggregate", "sort_rows", "create_sheet", "write_table", "add_formula_column", "apply_style_preset", "create_chart", "freeze_header"}
        composites = {"build_traceable_detail", "classify_invalid_rows", "calculate_profitability", "summarize_by_period", "summarize_by_dimension", "build_kpi_block", "build_reconciliation_sheet"}
        self.assertTrue(capabilities | composites <= set(definitions))
        self.assertTrue(all(definitions[name].kind == "CAPABILITY" for name in capabilities))
        self.assertTrue(all(definitions[name].kind == "COMPOSITE" for name in composites))

    def test_filter_rows_supports_iso_dates(self):
        data = TableData(["日期", "状态", "金额"], [{"日期": "2026-01-01", "状态": "有效", "金额": 10}, {"日期": "2025-01-01", "状态": "有效", "金额": 99}])
        result = filter_rows(data, {"all": [{"field": "日期", "op": "between", "value": ["2026-01-01", "2026-12-31"], "value_type": "date"}]})
        self.assertEqual(len(result.rows), 1); self.assertEqual(result.rows[0]["金额"], 10)


if __name__ == "__main__": unittest.main()
