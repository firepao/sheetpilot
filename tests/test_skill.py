from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "skills" / "sheetpilot-excel-agent"
CLI = SKILL / "scripts" / "sheetpilot_cli.py"


class SkillWorkflowTest(unittest.TestCase):
    def run_cli(self, *args: str) -> dict:
        environment=dict(os.environ); environment["PYTHONIOENCODING"]="utf-8"; environment["PYTHONUTF8"]="1"
        result=subprocess.run([sys.executable,str(CLI),*args],cwd=ROOT,text=True,capture_output=True,encoding="utf-8",errors="replace",env=environment)
        self.assertEqual(result.returncode,0,(result.stdout or "")+(result.stderr or "")); return json.loads(result.stdout)

    def test_skill_package_contains_only_lightweight_entrypoint(self):
        relative={path.relative_to(SKILL).as_posix() for path in SKILL.rglob("*") if path.is_file()}
        self.assertEqual(relative,{"SKILL.md","agents/openai.yaml","scripts/sheetpilot_cli.py"})

    def test_cli_lightweight_end_to_end(self):
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary); source=root/"source.xlsx"; output=root/"output.xlsx"; run_dir=root/"run"; plan_path=root/"plan.json"
            wb=Workbook(); ws=wb.active; ws.title="数据"; ws.append(["城市","金额"]); ws.append(["北京",10]); ws.append(["上海",30]); wb.save(source); wb.close()
            plan={"schema_version":"1.0","input_file":str(source),"output_file":str(output),"steps":[{"id":"source","op":"read_table","sheet":"数据","columns":{"城市":"A","金额":"B"}},{"id":"summary","op":"summarize_by_dimension","input":"source","group_by":["城市"],"metrics":[{"field":"金额","function":"sum","as":"总额"}],"sort":[{"field":"总额","direction":"desc"}]},{"id":"sheet","op":"create_sheet","sheet":"汇总"},{"id":"write","op":"write_table","input":"summary","sheet":"汇总"}],"requirements":{"required_sheets":["汇总"],"required_columns":{"汇总":["城市","总额"]},"checks":[{"type":"sort_order","sheet":"汇总","field":"总额","direction":"desc"}]}}
            plan_path.write_text(json.dumps(plan,ensure_ascii=False),encoding="utf-8")
            capabilities=self.run_cli("mvp-capabilities"); self.assertIn("aggregate",{item["name"] for item in capabilities["atoms"]})
            ready=self.run_cli("mvp-run","--input",str(source),"--output",str(output),"--plan",str(plan_path),"--run-dir",str(run_dir)); self.assertEqual(ready["status"],"READY_FOR_VALIDATION")
            result=self.run_cli("mvp-validate","--run-dir",str(run_dir)); self.assertEqual(result["status"],"PASS")
            workbook=load_workbook(output,data_only=True); self.assertEqual(workbook["汇总"]["B2"].value,30); workbook.close()


if __name__ == "__main__": unittest.main()
