from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[2]
SKILL = ROOT / "skills" / "sheetpilot-excel-agent"
CLI = SKILL / "scripts" / "sheetpilot_cli.py"


class SkillWorkflowTest(unittest.TestCase):
    def run_cli(self, *args: str) -> dict:
        environment=dict(os.environ); environment["PYTHONIOENCODING"]="utf-8"; environment["PYTHONUTF8"]="1"
        if hasattr(self,"state_root"): environment["SHEETPILOT_STATE_ROOT"]=str(self.state_root)
        result=subprocess.run([sys.executable,str(CLI),*args],cwd=ROOT,text=True,capture_output=True,encoding="utf-8",errors="replace",env=environment)
        self.assertEqual(result.returncode,0,(result.stdout or "")+(result.stderr or "")); return json.loads(result.stdout)

    def test_skill_package_contains_only_lightweight_entrypoint(self):
        relative={path.relative_to(SKILL).as_posix() for path in SKILL.rglob("*") if path.is_file()}
        self.assertEqual(relative,{"SKILL.md","agents/openai.yaml","scripts/sheetpilot_cli.py"})

    def test_wrapper_exposes_no_legacy_commands(self):
        result=subprocess.run([sys.executable,str(CLI),"mvp-capabilities"],cwd=ROOT,text=True,capture_output=True,encoding="utf-8",errors="replace")
        self.assertNotEqual(result.returncode,0)
        self.assertIn("invalid choice",result.stderr)

    def test_skill_exposes_only_agent_facing_task_workflow(self):
        text=(SKILL/"SKILL.md").read_text(encoding="utf-8")
        self.assertIn("task-types",text); self.assertIn("task-run",text); self.assertIn("task-status",text)
        self.assertIn("不要搜索可执行文件",text); self.assertIn("不要读取 SheetPilot 的 Python 文件",text)
        self.assertIn("<AUTO-RESULT-DIR>",text); self.assertIn("用户交付目录",text)
        self.assertIn("不要清空或复用",text); self.assertIn("不要提交 Runtime `run-dir`",text)
        self.assertIn("当前会话禁止自行设置环境变量后继续",text)
        self.assertIn("yyyyMMdd'T'HHmmss",text); self.assertIn("目录名不得包含 `:`",text)
        self.assertNotIn("mvp-run --input",text); self.assertNotIn("创建全新的空运行目录",text)

    def test_configuration_error_is_non_retryable_human_stop(self):
        text=CLI.read_text(encoding="utf-8")
        self.assertIn('"action":"HUMAN_ACTION_REQUIRED"',text)
        self.assertIn('"retryable":false',text)
        self.assertIn('"allowed_amendments":[]',text)

    def test_wrapper_discovers_repo_without_path_search_or_environment(self):
        environment=dict(os.environ); environment.pop("SHEETPILOT_ROOT",None); environment["PYTHONIOENCODING"]="utf-8"; environment["PYTHONUTF8"]="1"
        with tempfile.TemporaryDirectory() as temporary:
            result=subprocess.run([sys.executable,str(CLI),"task-types"],cwd=temporary,text=True,capture_output=True,encoding="utf-8",errors="replace",env=environment)
        self.assertEqual(result.returncode,0,(result.stdout or "")+(result.stderr or ""))
        self.assertEqual(json.loads(result.stdout)["task_types"][0]["name"],"summarize_table")

    def test_wrapper_discovers_repo_when_cwd_is_repo_root(self):
        environment=dict(os.environ); environment.pop("SHEETPILOT_ROOT",None); environment["PYTHONIOENCODING"]="utf-8"; environment["PYTHONUTF8"]="1"
        result=subprocess.run([sys.executable,str(CLI),"task-types"],cwd=ROOT,text=True,capture_output=True,encoding="utf-8",errors="replace",env=environment)
        self.assertEqual(result.returncode,0,(result.stdout or "")+(result.stderr or ""))
        self.assertEqual(json.loads(result.stdout)["task_types"][0]["name"],"summarize_table")

    def test_agent_task_api_end_to_end(self):
        with tempfile.TemporaryDirectory() as temporary:
            root=Path(temporary); source=root/"source.xlsx"; output=root/"output.xlsx"; request_path=root/"request.json"
            self.state_root=root/"state"
            wb=Workbook(); ws=wb.active; ws.title="数据"; ws.append(["城市","金额"]); ws.append(["北京",10]); ws.append(["上海",30]); wb.save(source); wb.close()
            request={"schema_version":"1.0","task_type":"summarize_table","input_file":str(source),"output_file":str(output),"user_request":"按城市汇总金额","source":{"sheet":"数据","header_row":1},"filters":[],"dimensions":[{"id":"city","field":"城市","output_name":"城市"}],"metrics":[{"id":"amount","function":"sum","field":"金额","output_name":"总额"}],"output":{"sheet":"汇总","anchor":"A1","sort":[{"by":"amount","direction":"desc"}]},"acceptance":{"required_filters":[],"required_dimensions":["city"],"required_metrics":["amount"],"required_sort":[{"by":"amount","direction":"desc"}]}}
            request_path.write_text(json.dumps(request,ensure_ascii=False),encoding="utf-8")
            self.assertEqual(self.run_cli("task-types")["task_types"][0]["name"],"summarize_table")
            ready=self.run_cli("task-run","--request",str(request_path)); self.assertEqual(ready["status"],"RUNTIME_PASS")
            result=self.run_cli("task-status","--task-id",ready["task_id"]); self.assertEqual(result["state"],"RUNTIME_PASS")
            workbook=load_workbook(output,data_only=True); self.assertEqual(workbook["汇总"]["B2"].value,30); workbook.close()
            del self.state_root


if __name__ == "__main__": unittest.main()
