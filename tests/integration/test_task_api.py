from __future__ import annotations

import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from sheetpilot.task_api import TaskRuntime
from sheetpilot.task_api.compiler import compile_plan
from sheetpilot.task_api.contract import MINIMAL_EXAMPLE, acceptance_snapshot, stable_hash


class TaskApiTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.source = self.root / "orders.xlsx"
        workbook = Workbook(); sheet = workbook.active; sheet.title = "清洗明细"
        sheet.append(["城市", "订单号", "销售额", "是否退货", "清洗状态", "客户编号"])
        sheet.append(["北京", "A", 10, "否", "有效", "C001"])
        sheet.append(["上海", "B", 30, "否", "有效", "C002"])
        sheet.append(["北京", "C", 90, "是", "有效", None])
        workbook.save(self.source); workbook.close()
        self.runtime = TaskRuntime(self.root / "state")

    def tearDown(self):
        self.temporary.cleanup()

    def request(self) -> dict:
        value = copy.deepcopy(MINIMAL_EXAMPLE)
        value["input_file"] = str(self.source)
        value["output_file"] = str(self.root / "result.xlsx")
        return value

    def test_task_types_is_the_complete_public_contract(self):
        manifest = self.runtime.task_types(); definition = manifest["task_types"][0]
        self.assertEqual(definition["name"], "summarize_table")
        self.assertIn("minimal_example", definition)
        self.assertEqual(definition["supported_features"]["metric_functions"], ["sum", "average", "count.rows", "count.non_empty"])

    def test_contract_failure_does_not_create_task(self):
        request = self.request(); request["metrics"][1]["mode"] = "row"
        result = self.runtime.run(request)
        self.assertEqual(result["status"], "REQUEST_INVALID")
        self.assertIsNone(result["task_id"])
        self.assertEqual(result["error"]["diagnostics"][0]["path"], "/metrics/1/mode")
        tasks = self.root / "state" / "tasks"
        self.assertTrue(not tasks.exists() or not list(tasks.iterdir()))

    def test_runtime_manages_attempt_and_publishes(self):
        request = self.request(); result = self.runtime.run(request)
        self.assertEqual(result["status"], "RUNTIME_PASS")
        self.assertTrue(result["delivery_valid"])
        self.assertNotIn("run_dir", result)
        task_dir = self.root / "state" / "tasks" / result["task_id"]
        self.assertTrue((task_dir / "attempts" / "attempt-001" / "internal-plan.json").is_file())
        workbook = load_workbook(request["output_file"], data_only=True)
        rows = list(workbook["城市经营汇总"].iter_rows(min_row=2, values_only=True)); workbook.close()
        # 过滤：清洗状态=有效 AND 是否退货=否 → 北京A(10,C001) 和 上海B(30,C002)
        # 指标：sales_revenue=sum(销售额), order_count=count.rows, customer_count=count.non_empty(客户编号), avg=average(销售额)
        # 上海: (30, 1, 1, 30.0), 北京: (10, 1, 1, 10.0) 按销售收入降序
        self.assertEqual(rows, [("上海", 30, 1, 1, 30.0), ("北京", 10, 1, 1, 10.0)])
        repeated = self.runtime.run(request)
        self.assertEqual(repeated["task_id"], result["task_id"])
        self.assertEqual(len(list((task_dir / "attempts").iterdir())), 1)

    def test_ambiguous_exact_headers_return_needs_binding_without_attempt(self):
        workbook = load_workbook(self.source); sheet = workbook["清洗明细"]
        sheet.insert_cols(4); sheet["D1"] = "销售额"; sheet["D2"] = 10; sheet["D3"] = 30; sheet["D4"] = 90
        workbook.save(self.source); workbook.close()
        result = self.runtime.run(self.request())
        self.assertEqual(result["status"], "NEEDS_BINDING")
        sales = next(item for item in result["binding_slots"] if item["logical_field"] == "销售额")
        self.assertEqual(sales["status"], "AMBIGUOUS")
        task_dir = self.root / "state" / "tasks" / result["task_id"]
        self.assertFalse((task_dir / "attempts").exists())

    def test_fuzzy_field_needs_binding_with_inventory(self):
        source = self.root / "net_sales.xlsx"
        workbook = Workbook(); sheet = workbook.active; sheet.title = "交易流水"
        sheet.append(["净销售额"]); sheet.append([100]); sheet.append([200]); sheet.append([300])
        workbook.save(source); workbook.close()
        request = copy.deepcopy(MINIMAL_EXAMPLE)
        request["input_file"] = str(source); request["output_file"] = str(self.root / "net-result.xlsx")
        request["source"] = {"sheet": "交易流水", "header_row": 1}
        request["filters"] = []
        request["dimensions"] = [{"id": "net", "field": "净销售收入", "output_name": "净销售收入"}]
        request["metrics"] = [{"id": "total", "function": "sum", "field": "净销售收入", "output_name": "合计"}]
        request["output"] = {"sheet": "汇总", "anchor": "A1", "sort": []}
        request["acceptance"]["required_filters"] = []
        request["acceptance"]["required_dimensions"] = ["net"]
        request["acceptance"]["required_metrics"] = ["total"]
        request["acceptance"]["required_sort"] = []
        result = self.runtime.run(request)
        self.assertEqual(result["status"], "NEEDS_BINDING")
        self.assertEqual(result["binding_slots"][0]["status"], "UNRESOLVED")
        net_candidate = next(c for c in result["field_inventory"] if c["header"] == "净销售额")
        self.assertNotIn("confidence", net_candidate)
        self.assertEqual(net_candidate["sample_values"], [100, 200, 300])
        self.assertEqual(result["recovery"]["action"], "PROVIDE_BINDING")

    def test_unrelated_field_yields_no_candidate_and_human_action(self):
        source = self.root / "unrelated.xlsx"
        workbook = Workbook(); sheet = workbook.active; sheet.title = "交易流水"
        sheet.append(["区域"]); sheet.append(["华东"])
        workbook.save(source); workbook.close()
        request = copy.deepcopy(MINIMAL_EXAMPLE)
        request["input_file"] = str(source); request["output_file"] = str(self.root / "unrelated-result.xlsx")
        request["source"] = {"sheet": "交易流水", "header_row": 1}
        request["filters"] = []
        request["dimensions"] = [{"id": "region", "field": "净利润", "output_name": "净利润"}]
        request["metrics"] = [{"id": "rows", "function": "count", "mode": "rows", "output_name": "行数"}]
        request["output"] = {"sheet": "汇总", "anchor": "A1", "sort": []}
        request["acceptance"]["required_filters"] = []
        request["acceptance"]["required_dimensions"] = ["region"]
        request["acceptance"]["required_metrics"] = ["rows"]
        request["acceptance"]["required_sort"] = []
        result = self.runtime.run(request)
        self.assertEqual(result["status"], "NEEDS_BINDING")
        self.assertEqual(result["binding_slots"][0]["status"], "UNRESOLVED")
        self.assertEqual(result["field_inventory"][0]["header"], "区域")
        self.assertEqual(result["recovery"]["action"], "PROVIDE_BINDING")

    def test_compiler_is_byte_deterministic_and_hides_display_names(self):
        request = self.request(); binding = self.runtime._bind(request, "input-hash")
        snapshot = {"source": binding["source"], "slots": binding["slots"]}
        acceptance_hash = stable_hash(acceptance_snapshot(request))
        first, first_hash = compile_plan("task-abc", 1, request, acceptance_hash, snapshot, "input-hash")
        second, second_hash = compile_plan("task-abc", 1, request, acceptance_hash, snapshot, "input-hash")
        self.assertEqual(json.dumps(first, ensure_ascii=False, sort_keys=True), json.dumps(second, ensure_ascii=False, sort_keys=True))
        self.assertEqual(first_hash, second_hash)
        summarize = next(step for step in first["steps"] if step["id"] == "summarize")
        self.assertTrue(all(field.startswith("__sp_") for field in summarize["group_by"]))
        self.assertNotIn("销售收入", json.dumps(summarize, ensure_ascii=False))

    def test_status_detects_modified_publication(self):
        result = self.runtime.run(self.request())
        workbook = load_workbook(self.root / "result.xlsx"); workbook["城市经营汇总"]["B2"] = 999; workbook.save(self.root / "result.xlsx"); workbook.close()
        status = self.runtime.status(result["task_id"])
        self.assertEqual(status["artifact_integrity"], "MODIFIED")
        self.assertFalse(status["delivery_valid"])

    def test_validation_blocks_wrong_temporary_result_before_publication(self):
        original = self.runtime._write_table

        def corrupt(engine, table, sheet, anchor):
            original(engine, table, sheet, anchor)
            engine.write_cell(sheet, 2, 2, 999999)

        with patch.object(self.runtime, "_write_table", side_effect=corrupt):
            result = self.runtime.run(self.request())
        self.assertEqual(result["status"], "VALIDATION_FAILED")
        self.assertFalse(result["delivery_valid"])
        self.assertFalse(Path(self.request()["output_file"]).exists())
        task_dir = self.root / "state" / "tasks" / result["task_id"]
        validation = json.loads((task_dir / "attempts" / "attempt-001" / "validation.json").read_text(encoding="utf-8"))
        self.assertFalse(validation["passed"])
        self.assertTrue(any(item["name"] == "output_values_and_order_match" and not item["passed"] for item in validation["checks"]))

    def test_validation_blocks_wrong_output_headers(self):
        original = self.runtime._write_table

        def corrupt(engine, table, sheet, anchor):
            original(engine, table, sheet, anchor)
            engine.write_cell(sheet, 1, 1, "错误表头")

        with patch.object(self.runtime, "_write_table", side_effect=corrupt):
            result = self.runtime.run(self.request())
        self.assertEqual(result["status"], "VALIDATION_FAILED")
        self.assertFalse(Path(self.request()["output_file"]).exists())
        self.assertTrue(any(item["name"] == "output_headers_match" and not item["passed"] for item in result["error"]["diagnostics"]))

    def test_validation_blocks_extra_output_cells(self):
        original = self.runtime._write_table

        def corrupt(engine, table, sheet, anchor):
            original(engine, table, sheet, anchor)
            engine.write_cell(sheet, 100, 100, "意外数据")

        with patch.object(self.runtime, "_write_table", side_effect=corrupt):
            result = self.runtime.run(self.request())
        self.assertEqual(result["status"], "VALIDATION_FAILED")
        self.assertTrue(any(item["name"] == "output_has_no_extra_non_empty_cells" and not item["passed"] for item in result["error"]["diagnostics"]))

    def test_validation_accepts_non_a1_anchor_and_records_evidence(self):
        request = self.request()
        request["output"]["anchor"] = "C3"
        result = self.runtime.run(request)
        self.assertEqual(result["status"], "RUNTIME_PASS")
        self.assertTrue(result["evidence"]["validation"]["passed"])
        self.assertEqual(result["evidence"]["validation"]["filtered_row_count"], 2)
        self.assertTrue(result["evidence"]["validation"]["validation_hash"])
        workbook = load_workbook(request["output_file"], data_only=True)
        self.assertEqual(workbook["城市经营汇总"]["C3"].value, "城市")
        self.assertEqual(workbook["城市经营汇总"]["D4"].value, 30)
        workbook.close()

    def test_validation_blocks_source_sheet_mutation(self):
        original = self.runtime._write_table

        def corrupt(engine, table, sheet, anchor):
            original(engine, table, sheet, anchor)
            source_workbook = load_workbook(self.source)
            source_workbook["清洗明细"]["A2"] = "被篡改"
            source_workbook.save(self.source)
            source_workbook.close()

        with patch.object(self.runtime, "_write_table", side_effect=corrupt):
            result = self.runtime.run(self.request())
        self.assertEqual(result["status"], "VALIDATION_FAILED")
        self.assertFalse(Path(self.request()["output_file"]).exists())
        self.assertFalse(result["error"]["diagnostics"][0]["passed"])

    def test_existing_target_sheet_fails_without_publishing(self):
        request=self.request(); request["output"]["sheet"]="清洗明细"
        request["output_file"]=str(self.root/"conflict.xlsx")
        result=self.runtime.run(request)
        self.assertEqual(result["status"],"EXECUTION_FAILED")
        self.assertEqual(result["error"]["code"],"OUTPUT_CONFLICT")
        self.assertFalse(Path(request["output_file"]).exists())

    def test_unsized_worksheet_is_bound_and_executed(self):
        request = copy.deepcopy(MINIMAL_EXAMPLE)
        request["input_file"] = str(Path(__file__).resolve().parents[1] / "agent_contract" / "data" / "01_simple_department_sales.xlsx")
        request["output_file"] = str(self.root / "unsized-result.xlsx")
        request["source"] = {"sheet": "销售明细", "header_row": 1}
        request["filters"] = []
        request["dimensions"] = [{"id": "department", "field": "部门", "output_name": "部门"}]
        request["metrics"] = [{"id": "total_sales", "function": "sum", "field": "销售额", "output_name": "销售总额"}]
        request["output"] = {"sheet": "部门销售汇总", "anchor": "A1", "sort": [{"by": "total_sales", "direction": "desc"}]}
        request["acceptance"]["required_filters"] = []
        request["acceptance"]["required_dimensions"] = ["department"]
        request["acceptance"]["required_metrics"] = ["total_sales"]
        request["acceptance"]["required_sort"] = [{"by": "total_sales", "direction": "desc"}]
        result = self.runtime.run(request)
        self.assertEqual(result["status"], "RUNTIME_PASS")

    def test_binding_amendment_executes_by_candidate_id(self):
        source = self.root / "amend.xlsx"
        workbook = Workbook(); sheet = workbook.active; sheet.title = "数据"
        sheet.append(["净销售额"]); sheet.append([100]); sheet.append([200]); workbook.save(source); workbook.close()
        request = copy.deepcopy(MINIMAL_EXAMPLE); request.update({"input_file": str(source), "output_file": str(self.root / "amend-result.xlsx"), "source": {"sheet": "数据", "header_row": 1}, "filters": [], "dimensions": [{"id": "net", "field": "净销售收入", "output_name": "净销售收入"}], "metrics": [{"id": "total", "function": "sum", "field": "净销售收入", "output_name": "合计"}], "output": {"sheet": "汇总", "anchor": "A1", "sort": []}})
        request["acceptance"].update({"required_filters": [], "required_dimensions": ["net"], "required_metrics": ["total"], "required_sort": []})
        pending = self.runtime.run(request); self.assertEqual(pending["status"], "NEEDS_BINDING")
        slot = pending["binding_slots"][0]; candidate = next(x["id"] for x in pending["field_inventory"] if x["header"] == "净销售额")
        final = self.runtime.run({"task_id": pending["task_id"], "base_revision": 1, "amendments": [{"op": "add", "path": f"/bindings/{slot['id']}", "value": {"candidate_id": candidate}}]})
        self.assertEqual(final["status"], "RUNTIME_PASS"); self.assertEqual(final["request_revision"], 2)


if __name__ == "__main__": unittest.main()
